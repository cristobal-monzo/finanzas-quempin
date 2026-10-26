# -*- coding: utf-8 -*-
"""Lo que se escribe en la Planilla de Ingreso: N° fijos y proyectos nuevos del Formulador.

Siempre sobre una planilla armada en tmp_path, y sin Excel: «escribir» se reemplaza por una versión
con openpyxl que aplica las mismas operaciones (la real, planilla_excel.ps1, se probó a mano sobre una
copia de la planilla el 2026-10-26)."""
import json

import openpyxl
import pytest

import intercambio as ic
import planilla_escritura as pe
import requerimientos as rq

ENCABEZADOS = ["N°", "Estado", "Canal", "Capt.", "Eval.", "Título", "Ubicación", "ID o Referencia", "Apertura",
               "Cierre", "Visita", "Eº visita", "Visitador", "Presupuesto (IVA inc.)", "Valor ofertado (IVA inc)",
               "Valor adjudicado (IVA inc)", "Entrega (días)", "Plazo Oferta (días)"]


def _planilla(tmp_path, filas, vacias=3):
    """filas: (N° fijo o None = fórmula, título). La fórmula guarda su valor como lo haría Excel."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = rq.HOJA
    hoja.append(ENCABEZADOS)
    for i, (numero, titulo) in enumerate(filas, start=2):
        hoja.append([numero if numero is not None else f'=IF(B{i}<>"",ROW()-1, "")', "Evaluación", "Otro", "NB", "NB",
                     titulo, "Valdivia", "-"])
    for j in range(vacias):
        hoja.append([f'=IF(B{len(filas) + 2 + j}<>"",ROW()-1, "")'])
    ruta = tmp_path / "Planilla.xlsx"
    libro.save(ruta)
    return ruta


def _filas(filas):
    """Lo que leer_filas devolvería (openpyxl no calcula fórmulas: el N° de una fórmula es fila-1)."""
    return [{"fila": i, "formula": n is None, "numero": n if n is not None else i - 1, "titulo": t, "referencia": "-"}
            for i, (n, t) in enumerate(filas, start=2)]


def _pub(*pares):
    return [{"numero": n, "titulo": t, "referencia": "-"} for n, t in pares]


# ── plan_fijar ───────────────────────────────────────────────────────────────

def test_fija_cada_formula_con_el_n_publicado_aunque_la_tabla_se_haya_ordenado():
    # Publicado: 1 A, 2 B, 3 C. Alguien ordenó: ahora las fórmulas dicen 1 C, 2 A, 3 B.
    ops, avisos = pe.plan_fijar(_filas([(None, "C"), (None, "A"), (None, "B")]), _pub((1, "A"), (2, "B"), (3, "C")))
    assert [(o["titulo"], o["numero"]) for o in ops] == [("C", 3), ("A", 1), ("B", 2)]
    assert avisos == []


def test_las_filas_ya_fijas_no_se_tocan_y_una_nueva_queda_con_su_n():
    filas = _filas([(1, "A"), (None, "B"), (None, "Nueva")])
    ops, avisos = pe.plan_fijar(filas, _pub((1, "A"), (2, "B")))
    assert [(o["fila"], o["numero"]) for o in ops] == [(3, 2), (4, 3)]
    assert avisos == []


def test_una_fila_nueva_cuyo_n_ya_existe_no_se_renumera_sola():
    filas = _filas([(2, "A"), (None, "Repetida")])   # la fórmula de la fila 3 dice 2: el de A
    ops, avisos = pe.plan_fijar(filas, _pub((2, "A")))
    assert ops == []
    assert "Repetida" in avisos[0] and "una persona" in avisos[0]


def test_titulos_repetidos_se_reparten_en_orden():
    ops, _ = pe.plan_fijar(_filas([(None, "Mantención"), (None, "Mantención")]), _pub((7, "Mantención"), (9, "Mantención")))
    assert [o["numero"] for o in ops] == [7, 9]


# ── fila_de_mensaje ──────────────────────────────────────────────────────────

def _mensaje(numero=4, **req):
    r = {"numero": numero, "titulo": "Caldera gimnasio"}
    r.update(req)
    return {"esquema": ic.ESQUEMA, "id": f"nuevo{numero:04d}", "tipo": pe.TIPO_NUEVO, "destino": rq.DESTINO,
            "origen": {"herramienta": "formulador", "usuario": "NB", "enviado": "2026-10-26T16:00:00-03:00"},
            "requerimiento": r, "fuente": {"herramienta": "formulador", "uid": "f1e2d3c4b5a6", "codigo": f"{numero}. Caldera gimnasio"}}


def test_fila_nueva_con_el_formato_de_las_demas():
    op = pe.fila_de_mensaje(_mensaje(captador="CM", evaluador="NB"), None, hoy=__import__("datetime").date(2026, 10, 26))
    assert op == {"op": "nueva", "numero": 4,
                  "valores": {"B": "Evaluación", "C": "Otro", "D": "CM", "E": "NB", "F": "Caldera gimnasio", "H": "-"},
                  "fechas": {"I": "2026-10-26"}}


def test_lo_del_proyecto_al_escribir_manda_y_un_id_de_licitacion_es_mercado_publico():
    formulacion = {"datos": {"titulo": "Caldera gimnasio municipal", "ubicacion": "Ciudad Ficticia", "idLicitacion": "1234-5-LE26"}}
    op = pe.fila_de_mensaje(_mensaje(), formulacion)
    assert op["valores"]["F"] == "Caldera gimnasio municipal"
    assert op["valores"]["G"] == "Ciudad Ficticia"
    assert op["valores"]["H"] == "1234-5-LE26"
    assert op["valores"]["C"] == "Mercado Público"


def test_un_titulo_que_parece_formula_va_con_apostrofo():
    assert pe.fila_de_mensaje(_mensaje(titulo="=HOLA"), None)["valores"]["F"] == "'=HOLA"


# ── la vuelta ────────────────────────────────────────────────────────────────

@pytest.fixture
def entorno(tmp_path, monkeypatch):
    raiz = ic.asegurar_carpeta(tmp_path / "Intercambio")
    monkeypatch.setattr(pe, "ESTADO", tmp_path / "estado.json")
    escritas = []

    def falso_escribir(ruta, operaciones, limite=600):
        escritas.append(operaciones)
        libro = openpyxl.load_workbook(ruta)
        hoja = libro[rq.HOJA]
        resultados = []
        for i, o in enumerate(operaciones):
            if o["op"] == "fijar":
                hoja.cell(o["fila"], 1).value = o["numero"]
                resultados.append({"i": i, "ok": True, "fila": o["fila"]})
            else:
                fila = max(r for r in range(2, hoja.max_row + 1) if hoja.cell(r, 2).value or r == 2) + 1
                hoja.cell(fila, 1).value = o["numero"]
                for letra, v in o["valores"].items():
                    hoja[f"{letra}{fila}"] = v
                resultados.append({"i": i, "ok": True, "fila": fila})
        libro.save(ruta)
        return {"estado": "ok", "detalle": "", "resultados": resultados}

    monkeypatch.setattr(pe, "escribir", falso_escribir)
    # openpyxl no calcula: el N° de una fórmula se lee como lo dejaría Excel (fila - 1)
    original = pe.leer_filas

    def leer(ruta):
        filas = original(ruta)
        for f in filas:
            if f["formula"]:
                f["numero"] = f["fila"] - 1
        return filas

    monkeypatch.setattr(pe, "leer_filas", leer)
    return raiz, escritas


def test_registra_el_proyecto_nuevo_y_republica_la_planilla(tmp_path, entorno):
    raiz, escritas = entorno
    ruta = _planilla(tmp_path, [(1, "A"), (2, "B"), (3, "C")])
    rq.publicar(raiz, ruta)
    ic.enviar(raiz, _mensaje(4))
    r = pe.vuelta(raiz, ruta)
    assert r["detalle"] == "0 N° fijados, 1 requerimiento(s) registrado(s)"
    assert ic.leer_buzon(raiz, destino=rq.DESTINO, tipo=pe.TIPO_NUEVO)[0] == []
    publicados = ic.leer_publicacion(raiz, "requerimientos")["datos"]["requerimientos"]
    assert [(x["numero"], x["titulo"]) for x in publicados][-1] == (4, "Caldera gimnasio")
    # La vuelta siguiente no abre Excel: nada cambió
    assert pe.vuelta(raiz, ruta)["detalle"] == "sin cambios"
    assert len(escritas) == 1


def test_un_n_que_ya_tomo_otro_no_se_escribe(tmp_path, entorno):
    raiz, escritas = entorno
    ruta = _planilla(tmp_path, [(1, "A"), (2, "B")])
    rq.publicar(raiz, ruta)
    ic.enviar(raiz, _mensaje(2))
    r = pe.vuelta(raiz, ruta)
    assert escritas == []
    assert "ya es «B»" in r["avisos"][0]
    assert ic.leer_buzon(raiz, destino=rq.DESTINO, tipo=pe.TIPO_NUEVO)[0] == []
    archivado = [json.loads(p.read_text(encoding="utf-8")) for p in (raiz / "procesado").rglob("*.json")]
    assert archivado[0]["resultado"]["estado"] == "rechazado"


def test_fija_las_formulas_de_filas_que_alguien_agrego_a_mano(tmp_path, entorno):
    raiz, escritas = entorno
    ruta = _planilla(tmp_path, [(1, "A"), (None, "B")])
    rq.publicar(raiz, ruta)
    r = pe.vuelta(raiz, ruta)
    assert r["detalle"].startswith("1 N° fijados")
    assert escritas[0] == [{"op": "fijar", "fila": 3, "titulo": "B", "numero": 2}]
