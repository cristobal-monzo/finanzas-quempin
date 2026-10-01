# -*- coding: utf-8 -*-
"""
requerimientos.py -- La Planilla de Ingreso de Requerimientos, vista por las
demás herramientas.

La planilla es de las personas: la lleva a mano el equipo comercial en la
raíz de la biblioteca «Formulación de proyectos - Documentos». Su N° es lo
primero que se asigna a un proyecto y ya prefija las carpetas de SharePoint
(«280. UMAG - …»), así que es la **clave común de proyecto** del intercambio
(plan de integración, 2026-09-30, §5.2).

Este módulo hace dos cosas, y ninguna escribe la planilla:

1. **Publica** una copia de solo lectura en ``publicado/requerimientos.json``
   (``publicar``), para que el Formulador cree proyectos desde un
   requerimiento y el registro de proyectos cruce claves.
2. **Cierra las sugerencias** que le mandan otras herramientas (mensaje
   ``actualizar-requerimiento``: estado, valor ofertado, valor adjudicado)
   cuando una persona ya las pasó a la planilla (``revisar_sugerencias``).

POR QUÉ NO SE ESCRIBE LA PLANILLA (medido el 2026-10-01 sobre una copia):
guardarla con openpyxl conserva valores, validaciones, formato condicional y
la tabla «Proyectos», pero **borra** ``xl/metadata.xml`` y ``xl/richData/*``
--lo que Excel usa para las fórmulas de matriz dinámica y las imágenes dentro
de celdas-- además de ``calcChain.xml``; el archivo baja de 147 KB a 87 KB.
Es la planilla que el equipo abre todos los días y la pérdida sería
silenciosa. Por eso una sugerencia se aplica **a mano** y se archiva sola
cuando la planilla ya muestra esos valores: la evidencia se busca donde vive
el dato, no en una confirmación aparte.
"""

import math
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

import esquemas
import intercambio

PUBLICACION = "requerimientos"
HERRAMIENTA = "planilla-requerimientos"
DESTINO = "planilla-requerimientos"
TIPO_SUGERENCIA = "actualizar-requerimiento"
HOJA = "Listado Requerimientos"
FUENTE = "Planilla de Ingreso de Requerimientos"

# Encabezado normalizado (sin tildes, minúsculas, sin puntos ni °) -> clave.
# Por nombre y no por posición: si alguien inserta una columna, nada se corre.
COLUMNAS = {
    "n": "numero",
    "estado": "estado",
    "canal": "canal",
    "capt": "captador",
    "eval": "evaluador",
    "titulo": "titulo",
    "ubicacion": "ubicacion",
    "id o referencia": "referencia",
    "apertura": "apertura",
    "cierre": "cierre",
    "visita": "visita",
    "visitador": "visitador",
    "presupuesto (iva inc)": "presupuesto",
    "valor ofertado (iva inc)": "valorOfertado",
    "valor adjudicado (iva inc)": "valorAdjudicado",
    "entrega (dias)": "entregaDias",
    "plazo oferta (dias)": "plazoOfertaDias",
}
NUMERICAS = {"presupuesto", "valorOfertado", "valorAdjudicado", "entregaDias", "plazoOfertaDias"}
# Lo que la planilla y una sugerencia comparten (clave de la sugerencia ->
# clave de la planilla publicada).
CAMPOS_SUGERIBLES = {"estado": "estado", "valorOfertado": "valorOfertado", "valorAdjudicado": "valorAdjudicado"}
ENCABEZADO_DE = {v: k for k, v in COLUMNAS.items()}
TITULO_COLUMNA = {"estado": "Estado", "valorOfertado": "Valor ofertado (IVA inc)",
                  "valorAdjudicado": "Valor adjudicado (IVA inc)"}

_ERROR_EXCEL = re.compile(r"^#(VALUE|REF|N/A|NAME|DIV/0|NUM|NULL|SPILL|CALC)[!?]?$", re.IGNORECASE)


def normalizar_encabezado(texto) -> str:
    plano = "".join(c for c in unicodedata.normalize("NFD", str(texto or "")) if unicodedata.category(c) != "Mn")
    plano = plano.lower().replace("°", "").replace("º", "").replace(".", "")
    return re.sub(r"\s+", " ", plano).strip()


def _valor(valor, clave: str):
    if valor is None:
        return None
    if isinstance(valor, str):
        texto = valor.strip()
        if not texto or _ERROR_EXCEL.match(texto):
            return None
        if clave in NUMERICAS:
            try:
                return float(texto.replace(".", "").replace(",", "."))
            except ValueError:
                return None
        return texto
    if isinstance(valor, bool):
        return valor
    if isinstance(valor, datetime):
        return valor.date().isoformat() if valor.time() == datetime.min.time() else valor.isoformat(timespec="minutes")
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, (int, float)):
        if not math.isfinite(valor):
            return None
        if clave in NUMERICAS:
            return int(valor) if float(valor).is_integer() else float(valor)
        return str(int(valor)) if float(valor).is_integer() else str(valor)
    return str(valor)


def _numero(valor) -> int | None:
    if isinstance(valor, bool) or valor is None:
        return None
    if isinstance(valor, (int, float)) and math.isfinite(valor) and float(valor).is_integer() and valor >= 1:
        return int(valor)
    if isinstance(valor, str) and valor.strip().isdigit():
        return int(valor.strip())
    return None


# ── LECTURA ──────────────────────────────────────────────────────────────────

def leer_planilla(ruta: Path, avisos: list | None = None) -> list[dict]:
    """Una fila por requerimiento con N°. Solo lectura (read_only + valores).
    Un N° repetido se queda con la primera fila y lo dice en 'avisos'."""
    import openpyxl

    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        hoja = libro[HOJA] if HOJA in libro.sheetnames else libro.worksheets[0]
        filas = hoja.iter_rows(values_only=True)
        encabezados = next(filas, ())
        indices = {}
        for i, texto in enumerate(encabezados):
            clave = COLUMNAS.get(normalizar_encabezado(texto))
            if clave and clave not in indices:
                indices[clave] = i
        if "numero" not in indices:
            raise ValueError(f"la hoja «{hoja.title}» no tiene la columna N°")
        salida, vistos = [], set()
        for fila in filas:
            numero = _numero(fila[indices["numero"]] if indices["numero"] < len(fila) else None)
            if numero is None:
                continue
            if numero in vistos:
                if avisos is not None:
                    avisos.append(f"El N° {numero} aparece más de una vez en la planilla: se usa la primera fila.")
                continue
            vistos.add(numero)
            registro = {"numero": numero}
            for clave, i in indices.items():
                if clave != "numero":
                    registro[clave] = _valor(fila[i] if i < len(fila) else None, clave)
            for clave in ("estado", "titulo"):
                registro.setdefault(clave, None)
            salida.append(registro)
        return salida
    finally:
        libro.close()


def resumen(requerimientos: list[dict]) -> dict:
    por_estado: dict[str, int] = {}
    for r in requerimientos:
        estado = r.get("estado") or "(sin estado)"
        por_estado[estado] = por_estado.get(estado, 0) + 1
    adjudicados = por_estado.get("Adjudicado", 0)
    perdidos = por_estado.get("No adjudicado", 0)
    return {
        "total": len(requerimientos),
        "porEstado": por_estado,
        # De lo que se ofertó y ya se resolvió, cuánto se ganó. Lo usa Flujo
        # de Caja para ponderar lo que sigue "Ofertado".
        "tasaAdjudicacion": round(adjudicados / (adjudicados + perdidos), 4) if adjudicados + perdidos else None,
    }


def datos_publicacion(ruta_planilla: Path) -> dict:
    avisos: list[str] = []
    reqs = leer_planilla(ruta_planilla, avisos)
    modificado = datetime.fromtimestamp(Path(ruta_planilla).stat().st_mtime).astimezone()
    return {
        "fuente": FUENTE,
        "actualizado": modificado.isoformat(timespec="seconds"),
        "requerimientos": reqs,
        "resumen": resumen(reqs),
        "avisos": avisos,
    }


def publicar(raiz_intercambio: Path, ruta_planilla: Path) -> Path:
    datos = datos_publicacion(ruta_planilla)
    sobre = {"esquema": intercambio.ESQUEMA, "herramienta": HERRAMIENTA, "generado": intercambio.ahora_iso(), "datos": datos}
    errores = esquemas.validar_publicacion(PUBLICACION, sobre)
    if errores:
        raise ValueError("la publicación de requerimientos no cumple su esquema: " + "; ".join(errores[:5]))
    return intercambio.publicar(raiz_intercambio, PUBLICACION, HERRAMIENTA, datos)


def por_numero(requerimientos: list[dict]) -> dict[int, dict]:
    return {r["numero"]: r for r in requerimientos}


# ── SUGERENCIAS (mensaje actualizar-requerimiento) ───────────────────────────

def _igual(a, b) -> bool:
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(float(a) - float(b)) < 0.5
    if isinstance(a, str) and isinstance(b, str):
        return normalizar_encabezado(a) == normalizar_encabezado(b)
    return a is None and b is None


def decidir_sugerencia(mensaje: dict, requerimiento: dict | None) -> dict:
    """{'accion': 'aplicado'|'pendiente'|'conflicto'|'rechazado', 'detalle': [...], 'falta': {...}}.
    'aplicado' = la planilla ya muestra lo sugerido (se archiva).
    'pendiente' = falta pasarlo a mano. 'conflicto' = la planilla cambió
    después de que quien sugirió la miró (reemplaza no calza)."""
    errores = esquemas.validar_mensaje(mensaje)
    if errores:
        return {"accion": "rechazado", "detalle": errores, "falta": {}}
    numero = mensaje["requerimiento"]["numero"]
    if requerimiento is None:
        return {"accion": "pendiente", "falta": dict(mensaje["cambios"]),
                "detalle": [f"El requerimiento N° {numero} no está en la planilla (¿N° mal escrito?)."]}
    reemplaza = mensaje.get("reemplaza") or {}
    falta, conflictos = {}, []
    for clave, nuevo in mensaje["cambios"].items():
        actual = requerimiento.get(CAMPOS_SUGERIBLES[clave])
        if _igual(actual, nuevo):
            continue
        falta[clave] = nuevo
        if clave in reemplaza and not _igual(actual, reemplaza[clave]):
            conflictos.append(f"{TITULO_COLUMNA[clave]}: la planilla dice {_texto(actual)}; "
                              f"quien sugirió vio {_texto(reemplaza[clave])}.")
    if not falta:
        return {"accion": "aplicado", "falta": {}, "detalle": ["La planilla ya muestra lo sugerido."]}
    detalle = [f"N° {numero} · {TITULO_COLUMNA[c]}: escribir {_texto(v)} (hoy {_texto(requerimiento.get(CAMPOS_SUGERIBLES[c]))})"
               for c, v in falta.items()]
    if conflictos:
        return {"accion": "conflicto", "falta": falta, "detalle": conflictos + detalle}
    return {"accion": "pendiente", "falta": falta, "detalle": detalle}


def _texto(valor) -> str:
    if valor is None:
        return "(vacío)"
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return "$" + f"{round(valor):,}".replace(",", ".")
    return f"«{valor}»"


def revisar_sugerencias(raiz_intercambio: Path, requerimientos: list[dict], archivar: bool = True) -> list[dict]:
    """Decide cada sugerencia del buzón contra la planilla actual. Con
    'archivar', mueve a procesado/ las que ya están aplicadas ('aplicado')
    y las inválidas ('rechazado'); las demás siguen en el buzón."""
    mensajes, invalidos = intercambio.leer_buzon(raiz_intercambio, destino=DESTINO, tipo=TIPO_SUGERENCIA)
    indice = por_numero(requerimientos)
    salida = []
    for m in mensajes:
        decision = decidir_sugerencia(m, indice.get((m.get("requerimiento") or {}).get("numero")))
        decision["mensaje"] = m
        if archivar and decision["accion"] in ("aplicado", "rechazado"):
            intercambio.archivar(raiz_intercambio, m, decision["accion"], decision["detalle"],
                                 extra={"requerimiento": (m.get("requerimiento") or {}).get("numero")})
        salida.append(decision)
    return salida


def descartar_sugerencia(raiz_intercambio: Path, id_mensaje: str, motivo: str = "") -> bool:
    mensajes, _ = intercambio.leer_buzon(raiz_intercambio, destino=DESTINO, tipo=TIPO_SUGERENCIA)
    for m in mensajes:
        if m["id"] == id_mensaje:
            intercambio.archivar(raiz_intercambio, m, "descartado",
                                 [motivo or "Descartada por el usuario: la planilla queda como está."])
            return True
    return False
