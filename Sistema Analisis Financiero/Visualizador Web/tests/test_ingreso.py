# -*- coding: utf-8 -*-
"""Pestaña «Ingresar datos» del tablero (2026-10-02).

POR QUÉ ESTE TEST EXISTE

Lo que se teclea en la pestaña lo arma ingreso.js (en el navegador) y lo
aplica presupuestos_formulador.py (en Python) sobre el Excel real. Si los dos
lados entienden distinto un campo -- el avance en % o en fracción, una fecha,
qué cuenta como «el mismo valor» -- el envío queda pendiente o, peor, escribe
otra cosa. Por eso estos tests arman los mensajes con el ingreso.js de verdad
(en Node) a partir de un snapshot de verdad, y los pasan por el catálogo y por
el plan de Python.
"""
import importlib.util
import json
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import pytest

import analisis_financiero as af

pf = af.pf
_RAIZ = Path(__file__).resolve().parent.parent
_spec = importlib.util.spec_from_file_location("build_visualizador_af", _RAIZ / "build_visualizador.py")
bv = sys.modules.get("build_visualizador_af")
if bv is None:
    bv = importlib.util.module_from_spec(_spec)
    sys.modules["build_visualizador_af"] = bv
    _spec.loader.exec_module(bv)
import esquemas  # noqa: E402  (Sistema Intercambio, en sys.path desde analisis_financiero)

COL = {nombre: idx for idx, nombre in enumerate(af.HEADERS_PROYECTOS, start=1)}
hay_node = pytest.mark.skipif(shutil.which("node") is None, reason="node no está instalado en esta máquina")


def _libro(tmp_path, filas):
    """filas: [(tag, nombre, {columna: valor})]"""
    ruta = tmp_path / "Análisis de Proyectos.xlsx"
    wb = af.asegurar_estructura_workbook(ruta)
    ws = wb[af.HOJA_PROYECTOS]
    for i, (tag, nombre, valores) in enumerate(filas, start=2):
        ws.cell(row=i, column=COL["TAG proyecto"], value=tag)
        ws.cell(row=i, column=COL["Nombre del proyecto"], value=nombre)
        for columna, valor in valores.items():
            ws.cell(row=i, column=COL[columna], value=valor)
    wb.save(ruta)
    return ruta


COMPLETO = {
    "% Avance": 0.35, "Fecha de inicio": datetime(2026, 9, 1), "Monto de Venta (sin IVA)": 5_000_000,
    "Costos Materiales Proyectados": 1_000_000, "Costos Equipos Proyectados": 200_000,
    "Mano de Obra Proyectada": 450_000, "Otros Costos Proyectados": 150_000, "Mano de Obra Real": 80_000.4,
    "N° Requerimiento": 280,
}


@pytest.fixture
def libro(tmp_path):
    return _libro(tmp_path, [
        ("DEMO", "Proyecto de prueba", COMPLETO),
        ("MEDIO", "A medio cargar", {"Monto de Venta (sin IVA)": 3_000_000, "Fecha de cierre": "a confirmar"}),
        ("GGEN", "Gastos Generales", {"Categoría": af.CATEGORIA_GASTOS_GENERALES}),
    ])


# ── SNAPSHOT ─────────────────────────────────────────────────────────────────

def test_los_campos_de_la_pestana_son_los_del_canal_que_los_aplica(libro):
    ingreso = bv.extraer_datos_saneados(libro)["ingreso"]
    assert [c["clave"] for c in ingreso["campos"]] == [k for k in pf.CAMPOS_TABLERO if k not in bv.FUERA_DE_LA_PESTANA]
    assert pf.COLUMNA_REQ not in [c["clave"] for c in ingreso["campos"]]   # pedido 2026-10-05: no aporta
    for c in ingreso["campos"]:
        assert (c["columna"], c["tipo"]) == pf.CAMPOS_TABLERO[c["clave"]]
        assert c["requerido"] == (c["columna"] in af.CAMPOS_MANUALES_REQUERIDOS)
        assert c["etiqueta"] and c["grupo"]
    assert [c["clave"] for c in ingreso["campos"] if c["positivo"]] == [pf.CLAVE_VENTA]
    assert (ingreso["tipo"], ingreso["destino"], ingreso["herramienta"]) == (pf.TIPO_DATOS, pf.DESTINO, pf.HERRAMIENTA_TABLERO)
    assert ingreso["esquemas"]["mensajes"] == {pf.TIPO_DATOS: esquemas.paquete()["mensajes"][pf.TIPO_DATOS]}


def test_cada_proyecto_trae_sus_valores_como_los_devuelve_el_tablero(libro):
    ingreso = bv.extraer_datos_saneados(libro)["ingreso"]
    json.dumps(ingreso)   # nada de datetime crudo
    demo, medio, ggen = ingreso["proyectos"]
    assert demo["valores"] == {
        "% Avance": 0.35, "Fecha de inicio": "2026-09-01", "Fecha de cierre": None, "Monto de Venta": 5_000_000,
        "Materiales": 1_000_000, "Equipos": 200_000, "Mano de Obra": 450_000, "Otros": 150_000,
        "Mano de Obra Real": 80_000.4,
    }
    assert demo["faltan"] == []
    assert medio["valores"]["Fecha de cierre"] == "a confirmar"      # texto escrito a mano: tal cual
    assert medio["faltan"] == ["% Avance", "Materiales", "Equipos", "Mano de Obra", "Otros", "Mano de Obra Real"]
    assert ggen["gastos_generales"] is True and ggen["faltan"] == []
    assert ingreso["tags"] == ["DEMO", "MEDIO", "GGEN"]


def test_peru_no_tiene_pestana_de_ingreso(libro):
    assert bv.extraer_datos_saneados(libro, pais="PE")["ingreso"] is None


def test_el_build_inserta_la_logica_y_el_validador_originales(libro, tmp_path, monkeypatch):
    monkeypatch.setattr(bv, "RUTA_EXCEL", libro)
    monkeypatch.setattr(bv, "RUTA_DATA_JSON", tmp_path / "data.json")
    monkeypatch.setattr(bv, "RUTA_BUILD_HTML", tmp_path / "index.html")
    assert bv.build() == 0
    html = (tmp_path / "index.html").read_text(encoding="utf-8")
    assert "__AF_" not in html
    assert bv.RUTA_INGRESO_JS.read_text(encoding="utf-8") in html
    assert bv.RUTA_ESQUEMAS_JS.read_text(encoding="utf-8") in html


def test_template_tiene_la_pestana_y_sus_ganchos():
    template = bv.RUTA_TEMPLATE.read_text(encoding="utf-8")
    for gancho in ("tabIngreso", "tabBtnIngreso", "ingTabla", "ingBarra", "ingDialogo", "ingNuevoDialogo",
                   "ingEnvios", "ingCarpeta", "ingArchivo", "pendientesCompletar", "af-esquemas-js", "af-ingreso-js"):
        assert 'id="' + gancho + '"' in template, gancho
    for marcador in ("__AF_ESQUEMAS_JS__", "__AF_INGRESO_JS__"):
        assert template.count(marcador) == 1, marcador
    # Lo que depende de los datos llega en el snapshot: el template no repite las claves.
    for clave in ("Costos Materiales Proyectados", "Mano de Obra Real\"", "N° Requerimiento"):
        assert clave not in template, clave


# ── ingreso.js (Node) contra Python ─────────────────────────────────────────

_SCRIPT_NODE = r"""
const fs = require('fs');
const Q = require(process.argv[2]);
const V = require(process.argv[3]);
const e = JSON.parse(fs.readFileSync(0, 'utf8'));
const campo = (clave) => e.ingreso ? e.ingreso.campos.find((c) => c.clave === clave) : e.campos[clave];
const salida = {};
if (e.parsear) salida.parsear = e.parsear.map(([clave, texto]) => Q.parsear(campo(clave), texto));
if (e.iguales) salida.iguales = e.iguales.map(([clave, a, b]) => Q.iguales(campo(clave), a, b));
if (e.mostrar) salida.mostrar = e.mostrar.map(([clave, v]) => Q.mostrar(campo(clave), v, 'es-CL'));
if (e.borrador) {
  let n = 0;
  salida.mensajes = Q.armarMensajes(e.ingreso, e.borrador, {
    usuario: 'Persona de prueba', ahora: new Date(2026, 9, 2, 10, 0, 0), nuevoId: () => 'idprueba' + (n++)
  });
  salida.errores = salida.mensajes.map((m) => V.validarMensaje(m, e.ingreso.esquemas));
  salida.nombres = salida.mensajes.map(Q.nombreArchivo);
}
if (e.nuevos) salida.nuevos = e.nuevos.map(([tag, nombre]) => Q.validarNuevo(e.ingreso, tag, nombre, []));
process.stdout.write(JSON.stringify(salida));
"""


def _node(tmp_path, entrada):
    script = tmp_path / "ingreso_prueba.js"
    script.write_text(_SCRIPT_NODE, encoding="utf-8")
    proc = subprocess.run(["node", str(script), str(bv.RUTA_INGRESO_JS), str(bv.RUTA_ESQUEMAS_JS)],
                          input=json.dumps(entrada, ensure_ascii=False).encode("utf-8"), capture_output=True, check=True)
    return json.loads(proc.stdout.decode("utf-8"))


CAMPOS = {clave: {"clave": clave, "tipo": tipo, "positivo": clave == pf.CLAVE_VENTA}
          for clave, (_, tipo) in pf.CAMPOS_TABLERO.items()}


@hay_node
def test_lo_tecleado_se_lee_como_lo_guarda_el_excel(tmp_path):
    casos = [
        ("% Avance", "35"), ("% Avance", "35,5"), ("% Avance", "100"), ("% Avance", "0"), ("% Avance", "120"),
        ("% Avance", "35%"), ("Monto de Venta", "12.500.000"), ("Monto de Venta", "$ 12.500.000"),
        ("Monto de Venta", "0"), ("Materiales", "0"), ("Materiales", "1500,6"), ("Materiales", "-5"),
        ("Materiales", "mucho"), ("Materiales", ""), ("Fecha de cierre", "2026-12-15"),
        ("Fecha de cierre", "2026-02-30"), ("N° Requerimiento", "280"), ("N° Requerimiento", "28,5"),
    ]
    r = _node(tmp_path, {"campos": CAMPOS, "parsear": casos})["parsear"]
    leidos = {f"{c}={t}": x.get("valor", "ERROR" if "error" in x else None) for (c, t), x in zip(casos, r)}
    assert leidos == {
        "% Avance=35": 0.35, "% Avance=35,5": 0.355, "% Avance=100": 1, "% Avance=0": 0, "% Avance=120": "ERROR",
        "% Avance=35%": 0.35, "Monto de Venta=12.500.000": 12_500_000, "Monto de Venta=$ 12.500.000": 12_500_000,
        "Monto de Venta=0": "ERROR", "Materiales=0": 0, "Materiales=1500,6": 1501, "Materiales=-5": "ERROR",
        "Materiales=mucho": "ERROR", "Materiales=": None, "Fecha de cierre=2026-12-15": "2026-12-15",
        "Fecha de cierre=2026-02-30": "ERROR", "N° Requerimiento=280": 280, "N° Requerimiento=28,5": "ERROR",
    }


@hay_node
def test_mismo_valor_en_javascript_y_en_python(tmp_path):
    casos = [
        ("Materiales", 1000000, 1000000.4), ("Materiales", 1000000, 1000001), ("Materiales", None, 0),
        ("Materiales", "", None), ("Materiales", "a confirmar", "a confirmar"), ("Materiales", "1000000", 1000000),
        ("% Avance", 0.35, 0.3500000001), ("% Avance", 0.35, 0.351), ("% Avance", 1, 1.0),
        ("Fecha de cierre", "2026-12-15", "2026-12-15"), ("Fecha de cierre", "2026-12-15", None),
        ("N° Requerimiento", 280, "280"),
    ]
    js = _node(tmp_path, {"campos": CAMPOS, "iguales": casos})["iguales"]
    py = [pf._mismo_valor(c, a, b) for c, a, b in casos]
    assert js == py, list(zip(casos, js, py))


@hay_node
def test_lo_que_arma_la_pestana_lo_acepta_y_lo_aplica_python(libro, tmp_path):
    """El recorrido completo sin navegador: snapshot real -> ingreso.js arma los
    mensajes -> catálogo (JS y Python) -> reglas y plan de Python."""
    ingreso = bv.extraer_datos_saneados(libro)["ingreso"]
    borrador = {
        "DEMO": {"valores": {"% Avance": 0.6, "Fecha de cierre": "2026-12-15", "Mano de Obra Real": 120000}},
        "MEDIO": {"valores": {"Materiales": 900000, "Fecha de cierre": None}},
        "NUEVO": {"valores": {"Monto de Venta": 8000000, "Fecha de inicio": "2026-11-02"}, "nuevo": {"nombre": "Obra nueva"}},
        "VACIO": {"valores": {}},
    }
    salida = _node(tmp_path, {"ingreso": ingreso, "borrador": borrador})
    mensajes = salida["mensajes"]
    assert [m["proyecto"]["tag"] for m in mensajes] == ["DEMO", "MEDIO", "NUEVO"]   # sin valores, no viaja
    assert salida["errores"] == [[], [], []]
    for m in mensajes:
        assert esquemas.validar_mensaje(m) == [] and pf.validar_contenido(m) == []
    demo, medio, nuevo = mensajes
    assert demo["reemplaza"] == {"% Avance": 0.35, "Fecha de cierre": None, "Mano de Obra Real": 80000.4}
    assert medio["reemplaza"] == {"Materiales": None, "Fecha de cierre": "a confirmar"}
    assert nuevo["proyecto"] == {"tag": "NUEVO", "nombre": "Obra nueva", "crear": True}
    assert nuevo["reemplaza"] == {"Fecha de inicio": None, "Monto de Venta": None}
    assert salida["nombres"][0] == "20261002-100000_tablero-af_datos-proyecto_idprueba0.json"

    import openpyxl
    ws = openpyxl.load_workbook(libro)[af.HOJA_PROYECTOS]
    filas = {"DEMO": 2, "MEDIO": 3, "GGEN": 4}
    actuales = {tag: pf.valores_actuales(ws, fila, {n: i for i, n in enumerate(af.HEADERS_PROYECTOS, start=1)})
                for tag, fila in filas.items()}
    decisiones = pf.planificar(mensajes, filas, actuales, {"valores": {}, "confirmados": []})
    assert [d["accion"] for d in decisiones] == ["aplicar", "aplicar", "crear"]


@hay_node
def test_lo_que_se_ve_en_la_pestana_es_lo_que_hay_en_el_excel(libro, tmp_path):
    ingreso = bv.extraer_datos_saneados(libro)["ingreso"]
    demo = ingreso["proyectos"][0]["valores"]
    casos = [["% Avance", demo["% Avance"]], ["Monto de Venta", demo["Monto de Venta"]],
             ["Mano de Obra Real", demo["Mano de Obra Real"]],
             ["Fecha de cierre", None], ["Fecha de cierre", "a confirmar"]]
    assert _node(tmp_path, {"ingreso": ingreso, "mostrar": casos})["mostrar"] == [
        "35", "5.000.000", "80.000", "", "a confirmar"]


@hay_node
def test_proyecto_nuevo_valida_tag_y_nombre(libro, tmp_path):
    ingreso = bv.extraer_datos_saneados(libro)["ingreso"]
    r = _node(tmp_path, {"ingreso": ingreso, "nuevos": [["OBRA", "Obra"], ["DEMO", "Otra"], ["obra 1", "X"], ["OBRA", ""]]})
    assert [len(e) for e in r["nuevos"]] == [0, 1, 1, 1]
    assert "Ya hay un proyecto" in r["nuevos"][1][0]
