"""Canal de entrada desde el Formulador de proyectos (Intercambio): costos
proyectados que llegan por el buzón y se escriben en "Proyectos" solo cuando
es seguro. Todo en carpetas temporales: nunca la carpeta real."""

import json

import openpyxl
import pytest

import analisis_financiero as af
import presupuestos_formulador as pf

ic = pf.intercambio
COL = {nombre: idx for idx, nombre in enumerate(af.HEADERS_PROYECTOS, start=1)}


@pytest.fixture(autouse=True)
def _sin_dashboard_real(monkeypatch):
    """ejecutar() regenera el dashboard al final: en estas pruebas no hace
    falta y no debe tocar el build real."""
    monkeypatch.setattr(af, "actualizar_visualizador_af", lambda pais="CL": True)


# ── ARMADO ───────────────────────────────────────────────────────────────────

def _excel_cc(tmp_path, filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Detalle"
    encabezados = [
        "N° Ref.", "Proyecto", "Tipo de Proyecto", "N° Documento", "Nombre Ítem",
        "Descripción", "Categoría Ítem", "Cantidad", "P. Unitario sin IVA",
        "Total sin IVA (CLP)", "Total con IVA (CLP)",
    ]
    for col, encabezado in enumerate(encabezados, start=1):
        ws.cell(row=1, column=col, value=encabezado)
    for fila, (n_ref, categoria, total) in enumerate(filas, start=2):
        ws.cell(row=fila, column=1, value=n_ref)
        ws.cell(row=fila, column=7, value=categoria)
        ws.cell(row=fila, column=10, value=total)
    ruta = tmp_path / "Centro de Costos.xlsx"
    wb.save(ruta)
    return ruta


def _excel_af(tmp_path, proyectos):
    """proyectos: [(tag, nombre, {columna: valor})]"""
    ruta = tmp_path / "Análisis de Proyectos.xlsx"
    wb = af.asegurar_estructura_workbook(ruta)
    ws = wb[af.HOJA_PROYECTOS]
    for fila, (tag, nombre, valores) in enumerate(proyectos, start=2):
        ws.cell(row=fila, column=1, value=tag)
        ws.cell(row=fila, column=2, value=nombre)
        for columna, valor in valores.items():
            ws.cell(row=fila, column=COL[columna], value=valor)
    wb.save(ruta)
    return ruta


def _mensaje(id_="envio001", tag="DEMO", costos=None, enviado="2026-09-30T10:00:00-03:00", **extra):
    m = {
        "esquema": ic.ESQUEMA, "id": id_, "tipo": pf.TIPO, "destino": pf.DESTINO,
        "origen": {"herramienta": "formulador", "enviado": enviado, "usuario": "Persona de prueba"},
        "proyecto": {"tag": tag},
        "fuente": {"codigo": "100", "version": 2, "titulo": "Oferta de prueba"},
        "costos": costos if costos is not None else {
            "Materiales": 1000000, "Equipos": 200000, "Mano de Obra": 450000, "Otros": 150000,
        },
    }
    m.update(extra)
    return m


def _enviar(raiz, mensaje):
    ic.asegurar_carpeta(raiz)
    ruta = raiz / ic.CARPETA_BUZON / f"{mensaje['origen']['enviado'][:19].replace(':', '')}_{mensaje['id']}.json"
    ruta.write_text(json.dumps(mensaje), encoding="utf-8")
    return ruta


@pytest.fixture
def entorno(tmp_path):
    """Libro AF con un proyecto DEMO sin costos proyectados, un Centro de
    Costos con una factura de DEMO, y carpetas temporales para todo."""
    return {
        "af": _excel_af(tmp_path, [("DEMO", "Proyecto de prueba", {"Monto de Venta (sin IVA)": 5000000})]),
        "cc": _excel_cc(tmp_path, [("DEMO-001", "Materiales", 300000.0)]),
        "facturas": tmp_path / "Facturas y Boletas",
        "respaldos": tmp_path / "Respaldos",
        "raiz": tmp_path / "Intercambio",
        "estado": tmp_path / "presupuestos_formulador.json",
        "tmp": tmp_path,
    }


def _correr(e, dry_run=False):
    return af.ejecutar(
        e["af"], e["cc"], e["facturas"], e["respaldos"],
        ruta_clientes_pendientes=e["tmp"] / "clientes_pendientes.json",
        dry_run=dry_run, raiz_intercambio=e["raiz"], ruta_estado_intercambio=e["estado"],
    )


def _proyectados(ruta, fila=2):
    ws = openpyxl.load_workbook(ruta)[af.HOJA_PROYECTOS]
    return {c: ws.cell(row=fila, column=COL[col]).value for c, col in pf.COLUMNA_POR_CATEGORIA.items()}


def _notas(ruta, fila=2):
    ws = openpyxl.load_workbook(ruta)[af.HOJA_PROYECTOS]
    return {c: ws.cell(row=fila, column=COL[col]).comment for c, col in pf.COLUMNA_POR_CATEGORIA.items()}


def _procesados(raiz):
    return {json.loads(p.read_text(encoding="utf-8"))["id"]: json.loads(p.read_text(encoding="utf-8"))
            for p in (raiz / ic.CARPETA_PROCESADO).rglob("*.json")}


def _catalogo(raiz):
    return ic.leer_publicacion(raiz, pf.PUBLICACION)["datos"]


# ── PLAN (función pura) ──────────────────────────────────────────────────────

def _plan(mensajes, actuales=None, estado=None, filas=None):
    filas = filas if filas is not None else {"DEMO": 2}
    actuales = actuales if actuales is not None else {"DEMO": {c: None for c in pf.CATEGORIAS}}
    estado = estado or {"valores": {}, "confirmados": []}
    return pf.planificar(mensajes, filas, actuales, estado)


def test_celdas_vacias_se_aplican():
    (d,) = _plan([_mensaje()])
    assert d["accion"] == "aplicar"
    assert sorted(c for c, _, _ in d["cambios"]) == sorted(pf.CATEGORIAS)


def test_valores_iguales_no_cambian_nada():
    actuales = {"DEMO": {"Materiales": 1000000, "Equipos": 200000.4, "Mano de Obra": 450000, "Otros": 150000}}
    (d,) = _plan([_mensaje()], actuales)
    assert d["accion"] == "sin-cambios" and d["cambios"] == []


def test_valor_escrito_a_mano_queda_pendiente_y_no_se_aplica_a_medias():
    actuales = {"DEMO": {"Materiales": 900000, "Equipos": None, "Mano de Obra": None, "Otros": None}}
    (d,) = _plan([_mensaje()], actuales)
    assert d["accion"] == "pendiente"
    assert d["conflictos"] == [("Materiales", 900000.0, 1000000)]
    assert any("confirmar envio001" in linea for linea in d["detalle"])


def test_valor_que_escribio_el_formulador_se_actualiza():
    actuales = {"DEMO": {"Materiales": 900000, "Equipos": None, "Mano de Obra": None, "Otros": None}}
    estado = {"valores": {"DEMO": {"Materiales": {"valor": 900000}}}, "confirmados": []}
    (d,) = _plan([_mensaje()], actuales, estado)
    assert d["accion"] == "aplicar"


def test_reemplaza_autoriza_solo_el_valor_que_se_vio():
    actuales = {"DEMO": {"Materiales": 900000, "Equipos": 50000, "Mano de Obra": None, "Otros": None}}
    visto = _mensaje(reemplaza={"Materiales": 900000, "Equipos": 50000, "Mano de Obra": None, "Otros": None})
    (d,) = _plan([visto], actuales)
    assert d["accion"] == "aplicar"
    # Si el Excel cambió después de que se vio, vuelve a ser un conflicto.
    viejo = _mensaje(reemplaza={"Materiales": 800000, "Equipos": 50000})
    (d,) = _plan([viejo], actuales)
    assert d["accion"] == "pendiente" and [c for c, _, _ in d["conflictos"]] == ["Materiales"]


def test_confirmado_por_consola_reemplaza_valores_manuales():
    actuales = {"DEMO": {"Materiales": 900000, "Equipos": None, "Mano de Obra": None, "Otros": None}}
    (d,) = _plan([_mensaje()], actuales, {"valores": {}, "confirmados": ["envio001"]})
    assert d["accion"] == "aplicar"


def test_tag_inexistente_espera_y_proyecto_nuevo_se_crea():
    (d,) = _plan([_mensaje(tag="NUEVO")])
    assert d["accion"] == "pendiente" and "no existe" in d["detalle"][0]
    (d,) = _plan([_mensaje(tag="nuevo", proyecto={"tag": "nuevo", "crear": True, "nombre": "Obra nueva"})])
    assert d["accion"] == "crear" and d["tag"] == "NUEVO" and d["nombre"] == "Obra nueva"


@pytest.mark.parametrize("cambio", [
    {"costos": {}},
    {"costos": {"Materiales": -1}},
    {"costos": {"Materiales": "mucho"}},
    {"costos": {"Gastos generales": 10}},
    {"proyecto": {"tag": ""}},
    {"proyecto": {"tag": "X Y!", "crear": True, "nombre": "Obra"}},
    {"proyecto": {"tag": "OBRA", "crear": True}},
])
def test_contenido_invalido_se_rechaza(cambio):
    (d,) = _plan([_mensaje(**cambio)])
    assert d["accion"] == "rechazado" and d["detalle"]


def test_varios_envios_al_mismo_tag_gana_el_ultimo():
    viejo = _mensaje("envio001", enviado="2026-09-30T09:00:00-03:00")
    nuevo = _mensaje("envio002", enviado="2026-09-30T11:00:00-03:00", costos={"Materiales": 5})
    decisiones = _plan([viejo, nuevo])
    por_id = {d["mensaje"]["id"]: d["accion"] for d in decisiones}
    assert por_id == {"envio001": "reemplazado", "envio002": "aplicar"}


# ── CORRIDA COMPLETA (ejecutar) ──────────────────────────────────────────────

def test_run_aplica_anota_archiva_registra_y_publica(entorno):
    _enviar(entorno["raiz"], _mensaje())
    resumen = _correr(entorno)

    assert resumen["error"] is None
    assert resumen["intercambio"]["aplicar"]
    assert _proyectados(entorno["af"]) == {
        "Materiales": 1000000, "Equipos": 200000, "Mano de Obra": 450000, "Otros": 150000,
    }
    notas = _notas(entorno["af"])
    assert all(n is not None and n.author == pf.AUTOR_NOTA and "Formulación 100 v2" in n.text for n in notas.values())
    # el Excel se respaldó antes de escribir
    assert len(list(entorno["respaldos"].rglob("*.xlsx"))) == 1

    assert not list((entorno["raiz"] / "buzon").glob("*.json"))
    procesado = _procesados(entorno["raiz"])["envio001"]
    assert procesado["resultado"]["estado"] == "aplicado"
    assert procesado["resultado"]["cambios"]["Materiales"] == {"antes": None, "despues": 1000000}

    estado = json.loads(entorno["estado"].read_text(encoding="utf-8"))
    assert estado["valores"]["DEMO"]["Materiales"]["valor"] == 1000000

    catalogo = _catalogo(entorno["raiz"])
    (proyecto,) = catalogo["proyectos"]
    assert proyecto["tag"] == "DEMO"
    assert proyecto["proyectados"]["Materiales"] == 1000000
    assert proyecto["origen"]["Materiales"]["mensaje"] == "envio001"
    assert proyecto["reales"]["Materiales"] == 300000.0
    assert catalogo["mensajes"]["envio001"]["estado"] == "aplicado"


def test_run_usa_los_costos_en_la_misma_corrida(entorno):
    _enviar(entorno["raiz"], _mensaje())
    _correr(entorno)
    ws = openpyxl.load_workbook(entorno["af"])[af.HOJA_PROYECTOS]
    # "Costos Totales Proyectado" es fórmula sobre las 4 columnas recién escritas
    assert str(ws.cell(row=2, column=COL["Total Proyectado"]).value).startswith("=")


def test_valor_manual_no_se_pisa_y_el_envio_queda_en_el_buzon(entorno):
    wb = openpyxl.load_workbook(entorno["af"])
    wb[af.HOJA_PROYECTOS].cell(row=2, column=COL["Costos Materiales Proyectados"], value=900000)
    wb.save(entorno["af"])
    _enviar(entorno["raiz"], _mensaje())

    resumen = _correr(entorno)

    assert resumen["intercambio"]["pendiente"]
    assert _proyectados(entorno["af"])["Materiales"] == 900000
    assert _proyectados(entorno["af"])["Equipos"] is None      # nunca a medias
    assert list((entorno["raiz"] / "buzon").glob("*.json"))
    assert _catalogo(entorno["raiz"])["mensajes"]["envio001"]["estado"] == "pendiente"


def test_confirmar_y_descartar(entorno):
    wb = openpyxl.load_workbook(entorno["af"])
    wb[af.HOJA_PROYECTOS].cell(row=2, column=COL["Costos Materiales Proyectados"], value=900000)
    wb.save(entorno["af"])
    _enviar(entorno["raiz"], _mensaje())
    _correr(entorno)

    assert pf.confirmar("envio001", entorno["raiz"], entorno["estado"])
    assert not pf.confirmar("no-existe", entorno["raiz"], entorno["estado"])
    _correr(entorno)
    assert _proyectados(entorno["af"])["Materiales"] == 1000000
    assert json.loads(entorno["estado"].read_text(encoding="utf-8"))["confirmados"] == []

    _enviar(entorno["raiz"], _mensaje("envio002", costos={"Materiales": 1}, enviado="2026-10-01T10:00:00-03:00"))
    assert pf.descartar(entorno["raiz"], "envio002")
    assert _procesados(entorno["raiz"])["envio002"]["resultado"]["estado"] == "descartado"
    _correr(entorno)
    assert _proyectados(entorno["af"])["Materiales"] == 1000000


def test_cambio_a_mano_despues_de_aplicar_quita_la_procedencia(entorno):
    _enviar(entorno["raiz"], _mensaje())
    _correr(entorno)
    wb = openpyxl.load_workbook(entorno["af"])
    wb[af.HOJA_PROYECTOS].cell(row=2, column=COL["Costos Equipos Proyectados"], value=333)
    wb.save(entorno["af"])

    resumen = _correr(entorno)

    assert any("se cambió a mano" in a for a in resumen["avisos"])
    notas = _notas(entorno["af"])
    assert notas["Equipos"] is None and notas["Materiales"] is not None
    estado = json.loads(entorno["estado"].read_text(encoding="utf-8"))
    assert "Equipos" not in estado["valores"]["DEMO"]
    assert _catalogo(entorno["raiz"])["proyectos"][0]["origen"]["Equipos"] is None

    # Un envío nuevo ya no puede pisar ese valor sin haberlo visto...
    _enviar(entorno["raiz"], _mensaje("envio002", enviado="2026-10-01T10:00:00-03:00"))
    _correr(entorno)
    assert _proyectados(entorno["af"])["Equipos"] == 333
    # ...pero sí si quien envía lo vio y decidió reemplazarlo.
    _enviar(entorno["raiz"], _mensaje("envio003", enviado="2026-10-01T11:00:00-03:00",
                                      reemplaza={"Equipos": 333}))
    _correr(entorno)
    assert _proyectados(entorno["af"])["Equipos"] == 200000
    assert _procesados(entorno["raiz"])["envio002"]["resultado"]["estado"] == "reemplazado"


def test_si_no_se_puede_guardar_el_excel_nada_se_archiva(entorno, monkeypatch):
    _enviar(entorno["raiz"], _mensaje())

    def guardar_falla(self, ruta):
        raise PermissionError("abierto en Excel")

    monkeypatch.setattr(openpyxl.Workbook, "save", guardar_falla)
    resumen = _correr(entorno)
    monkeypatch.undo()

    assert resumen["error"]
    assert list((entorno["raiz"] / "buzon").glob("*.json"))
    assert not entorno["estado"].exists()
    assert _proyectados(entorno["af"])["Materiales"] is None


def test_dry_run_no_escribe_nada(entorno):
    _enviar(entorno["raiz"], _mensaje())
    antes = entorno["af"].read_bytes()

    resumen = _correr(entorno, dry_run=True)

    assert resumen["intercambio"]["aplicar"]
    assert entorno["af"].read_bytes() == antes
    assert list((entorno["raiz"] / "buzon").glob("*.json"))
    assert not (entorno["raiz"] / "publicado" / "analisis-financiero.json").exists()
    assert not entorno["estado"].exists()


def test_proyecto_nuevo_crea_su_fila(entorno):
    _enviar(entorno["raiz"], _mensaje(tag="OBRA", proyecto={"tag": "OBRA", "crear": True, "nombre": "Obra nueva"}))
    resumen = _correr(entorno)

    assert "Obra nueva" in resumen["proyectos_nuevos"]
    ws = openpyxl.load_workbook(entorno["af"])[af.HOJA_PROYECTOS]
    assert ws.cell(row=3, column=1).value == "OBRA"
    assert ws.cell(row=3, column=2).value == "Obra nueva"
    assert ws.cell(row=3, column=COL["Costos Materiales Proyectados"]).value == 1000000
    assert str(ws.cell(row=3, column=COL["Total Proyectado"]).value).startswith("=")


def test_fila_nueva_no_cae_sobre_una_fila_a_medio_cargar(entorno):
    wb = openpyxl.load_workbook(entorno["af"])
    wb[af.HOJA_PROYECTOS].cell(row=3, column=1, value="MEDIO")   # TAG sin Nombre: no es fila válida
    wb.save(entorno["af"])
    _enviar(entorno["raiz"], _mensaje(tag="OBRA", proyecto={"tag": "OBRA", "crear": True, "nombre": "Obra nueva"}))
    _correr(entorno)
    ws = openpyxl.load_workbook(entorno["af"])[af.HOJA_PROYECTOS]
    assert ws.cell(row=3, column=1).value == "MEDIO"
    assert ws.cell(row=4, column=1).value == "OBRA"


def test_sin_carpeta_de_intercambio_explicita_un_excel_temporal_no_la_usa(entorno, monkeypatch):
    """Un test (o cualquiera que pase su propio Excel) nunca toca la carpeta
    real, aunque exista: el intercambio corre solo contra el libro del país."""
    monkeypatch.setitem(af.PAISES["CL"], "raiz_intercambio", entorno["tmp"] / "NO-TOCAR")
    resumen = af.ejecutar(entorno["af"], entorno["cc"], entorno["facturas"], entorno["respaldos"],
                          ruta_clientes_pendientes=entorno["tmp"] / "clientes_pendientes.json")
    assert resumen["intercambio"] == {}
    assert not (entorno["tmp"] / "NO-TOCAR").exists()


def test_catalogo_excluye_gastos_generales(entorno, monkeypatch):
    wb = openpyxl.load_workbook(entorno["af"])
    ws = wb[af.HOJA_PROYECTOS]
    ws.cell(row=3, column=1, value="GGEN")
    ws.cell(row=3, column=2, value="Gastos Generales")
    wb.save(entorno["af"])

    categorias = {"DEMO": "Mantenimiento", "GGEN": af.CATEGORIA_GASTOS_GENERALES}
    original = af.asegurar_categoria_proyectos
    monkeypatch.setattr(
        af, "asegurar_categoria_proyectos",
        lambda ws_p, filas, _categorias, columna: original(ws_p, filas, categorias, columna),
    )
    _correr(entorno)
    assert [p["tag"] for p in _catalogo(entorno["raiz"])["proyectos"]] == ["DEMO"]


def test_buzon_con_archivo_ilegible_avisa_y_lo_deja(entorno):
    ic.asegurar_carpeta(entorno["raiz"])
    (entorno["raiz"] / "buzon" / "roto.json").write_text("{no", encoding="utf-8")
    resumen = _correr(entorno)
    assert any("roto.json" in a for a in resumen["avisos"])
    assert (entorno["raiz"] / "buzon" / "roto.json").exists()


def test_publicar_sin_run_no_escribe_el_excel_ni_aplica(entorno, monkeypatch):
    _enviar(entorno["raiz"], _mensaje())
    monkeypatch.setitem(af.PAISES, "ZZ", dict(
        af.PAISES["CL"], ruta_excel_af=entorno["af"], ruta_excel_cc=entorno["cc"],
        raiz_intercambio=entorno["raiz"], ruta_estado_intercambio=entorno["estado"],
    ))
    antes = entorno["af"].read_bytes()

    resultado = af.publicar_intercambio(pais="ZZ")

    assert resultado["error"] is None and resultado["proyectos"] == 1
    assert entorno["af"].read_bytes() == antes
    assert list((entorno["raiz"] / "buzon").glob("*.json"))          # no aplicó nada
    catalogo = _catalogo(entorno["raiz"])
    assert catalogo["proyectos"][0]["tag"] == "DEMO"
    assert catalogo["proyectos"][0]["reales"]["Materiales"] == 300000.0
    assert "envio001" not in catalogo["mensajes"]                    # se aplicaría: no es pendiente
    assert not entorno["estado"].exists()


# ── UBICACIÓN DE LA CARPETA (2026-09-30: dentro de la biblioteca de Formulación) ──

def test_ubicar_intercambio_dentro_de_la_biblioteca_sin_distinguir_mayusculas(monkeypatch, tmp_path):
    herramientas = tmp_path / pf.BIBLIOTECA_FORMULACION / ".HERRAMIENTAS FORMULACIÓN"
    herramientas.mkdir(parents=True)
    monkeypatch.delenv("QUEMPIN_INTERCAMBIO", raising=False)
    monkeypatch.setattr(pf, "RAIZ_FINANZAS", tmp_path / "Escritorio" / "Finanzas QUEMPIN")
    assert pf.ubicar_intercambio() == herramientas / "Intercambio"


def test_ubicar_intercambio_sin_biblioteca_queda_inactivo(monkeypatch, tmp_path):
    monkeypatch.delenv("QUEMPIN_INTERCAMBIO", raising=False)
    monkeypatch.setattr(pf, "RAIZ_FINANZAS", tmp_path / "Escritorio" / "Finanzas QUEMPIN")
    assert pf.ubicar_intercambio() is None


def test_ubicar_intercambio_respeta_la_variable(monkeypatch, tmp_path):
    monkeypatch.setenv("QUEMPIN_INTERCAMBIO", str(tmp_path / "otra"))
    assert pf.ubicar_intercambio() == tmp_path / "otra"


def test_sin_carpeta_de_intercambio_el_status_lo_avisa(monkeypatch, tmp_path):
    monkeypatch.setitem(af.PAISES["CL"], "raiz_intercambio", None)
    monkeypatch.setitem(af.PAISES["CL"], "aviso_sin_intercambio", pf.AVISO_SIN_INTERCAMBIO)
    monkeypatch.setitem(af.PAISES["CL"], "ruta_excel_af", tmp_path / "af.xlsx")
    monkeypatch.setitem(af.PAISES["CL"], "ruta_excel_cc", tmp_path / "no-existe.xlsx")
    resumen = af.ejecutar(dry_run=True)
    assert pf.AVISO_SIN_INTERCAMBIO in resumen["avisos"]


# ── PRESUPUESTOS ADJUDICADOS DEL REPOSITORIO (con Claude, 2026-09-30) ─────────

COSTOS_FZ = {"Materiales": 178820, "Equipos": 560000, "Mano de Obra": 225000, "Otros": 150000}


def _presupuesto(raiz, uid="a1b2c3d4e5f6", estado="Adjudicada", titulo="Proyecto de prueba",
                 costos=COSTOS_FZ, **datos_extra):
    fz = pf.formulaciones
    datos = {"uid": uid, "codigo": "QPN-2026-001", "version": 1, "titulo": titulo, "cliente": "",
             "estado": estado, "responsable": "Persona", "modificado": "2026-09-30T15:00:00.000Z"}
    datos.update(datos_extra)
    resumen = {"costoDirecto": sum(costos.values()), "precioNeto": 1856404, "costosAF": costos} if costos else {}
    carpeta = fz.carpeta_repositorio(raiz)
    carpeta.mkdir(parents=True, exist_ok=True)
    (carpeta / fz.nombre_archivo(uid)).write_text(json.dumps({
        "esquema": ic.ESQUEMA, "herramienta": "formulador", "tipo": "proyecto", "autor": "Persona",
        "historia": [], "resumen": resumen, "datos": datos,
    }), encoding="utf-8")
    return fz.leer_repositorio(raiz)[-1] if len(fz.leer_repositorio(raiz)) == 1 else next(
        i for i in fz.leer_repositorio(raiz) if i["uid"] == uid)


def test_adjudicado_sin_tag_pide_elegirlo_con_sugerencia(entorno):
    _correr(entorno)                       # publica la lista con DEMO «Proyecto de prueba»
    _presupuesto(entorno["raiz"])
    _presupuesto(entorno["raiz"], uid="borrador0001", estado="Borrador")
    _presupuesto(entorno["raiz"], uid="papelera0001", eliminado=True)
    (f,) = pf.adjudicados_para_af(entorno["raiz"], entorno["estado"])
    assert f["accion"] == "elegir-tag" and f["sugerido"] == "DEMO" and f["costos"] == COSTOS_FZ


def test_cargar_un_adjudicado_lo_aplica_en_el_run_y_despues_queda_cargado(entorno):
    _correr(entorno)
    item = _presupuesto(entorno["raiz"])
    catalogo = {p["tag"]: p for p in _catalogo(entorno["raiz"])["proyectos"]}
    m = pf.mensaje_desde_formulacion(item, "demo", visto=catalogo["DEMO"]["proyectados"], usuario="Persona (cargado con Claude)")
    ic.enviar(entorno["raiz"], m)
    (f,) = pf.adjudicados_para_af(entorno["raiz"], entorno["estado"])
    assert f["accion"] == "en-buzon"

    _correr(entorno)
    assert _proyectados(entorno["af"]) == COSTOS_FZ
    registro = json.loads(entorno["estado"].read_text(encoding="utf-8"))
    assert registro["valores"]["DEMO"]["Materiales"]["uid"] == "a1b2c3d4e5f6"
    (f,) = pf.adjudicados_para_af(entorno["raiz"], entorno["estado"])
    assert f["accion"] == "cargado" and f["tag"] == "DEMO"


def test_si_el_presupuesto_cambia_despues_de_cargarlo_se_propone_de_nuevo(entorno):
    _correr(entorno)
    item = _presupuesto(entorno["raiz"])
    ic.enviar(entorno["raiz"], pf.mensaje_desde_formulacion(item, "DEMO"))
    _correr(entorno)
    _presupuesto(entorno["raiz"], costos=dict(COSTOS_FZ, Equipos=600000))
    (f,) = pf.adjudicados_para_af(entorno["raiz"], entorno["estado"])
    assert f["accion"] == "cargar" and f["tag"] == "DEMO" and f["proyectados"]["Equipos"] == 560000


def test_el_vinculo_del_formulador_da_el_tag(entorno):
    _correr(entorno)
    _presupuesto(entorno["raiz"], vinculos={"analisisFinanciero": {"tag": "demo"}})
    (f,) = pf.adjudicados_para_af(entorno["raiz"], entorno["estado"])
    assert f["accion"] == "cargar" and f["tag"] == "DEMO"


def test_sin_resumen_no_se_puede_cargar(entorno):
    item = _presupuesto(entorno["raiz"], costos=None)
    (f,) = pf.adjudicados_para_af(entorno["raiz"], entorno["estado"])
    assert f["accion"] == "sin-resumen"
    with pytest.raises(ValueError):
        pf.mensaje_desde_formulacion(item, "DEMO")


def test_cargar_en_un_tag_nuevo_crea_el_proyecto(entorno):
    item = _presupuesto(entorno["raiz"], titulo="Obra nueva")
    ic.enviar(entorno["raiz"], pf.mensaje_desde_formulacion(item, "OBRA", nombre="Obra nueva"))
    _correr(entorno)
    ws = openpyxl.load_workbook(entorno["af"])[af.HOJA_PROYECTOS]
    assert ws.cell(row=3, column=1).value == "OBRA" and ws.cell(row=3, column=2).value == "Obra nueva"


# ── VENTA, N° DE REQUERIMIENTO Y SESGO (plan de integración, 2026-10-01) ─────

def _venta(id_="venta001", tag="DEMO", monto=7000000, enviado="2026-10-01T10:00:00-03:00", **extra):
    m = {
        "esquema": ic.ESQUEMA, "id": id_, "tipo": pf.TIPO_VENTA, "destino": pf.DESTINO,
        "origen": {"herramienta": "formulador", "enviado": enviado, "usuario": "Persona de prueba"},
        "proyecto": {"tag": tag},
        "venta": {"montoSinIva": monto, "moneda": "CLP"},
        "fuente": {"herramienta": "sistema-quempin", "folio": "602695", "pais": "Chile", "titulo": "Oferta"},
    }
    m.update(extra)
    return m


def _celda(ruta, columna, fila=2):
    return openpyxl.load_workbook(ruta)[af.HOJA_PROYECTOS].cell(row=fila, column=COL[columna])


def test_venta_y_presupuesto_del_mismo_tag_son_independientes():
    decisiones = _plan([_mensaje(), _venta()],
                       actuales={"DEMO": {**{c: None for c in pf.CATEGORIAS}, pf.CLAVE_VENTA: None}})
    assert sorted(d["accion"] for d in decisiones) == ["aplicar", "aplicar"]


def test_venta_sobre_un_monto_escrito_a_mano_queda_pendiente():
    (d,) = _plan([_venta()], actuales={"DEMO": {pf.CLAVE_VENTA: 5000000}})
    assert d["accion"] == "pendiente"
    assert "Formulador envía $7.000.000" in d["detalle"][0]


def test_venta_con_reemplaza_del_valor_visto_se_aplica():
    (d,) = _plan([_venta(reemplaza={"montoSinIva": 5000000})], actuales={"DEMO": {pf.CLAVE_VENTA: 5000000}})
    assert d["accion"] == "aplicar" and d["cambios"] == [(pf.CLAVE_VENTA, 5000000, 7000000)]


@pytest.mark.parametrize("cambio", [{"venta": {"montoSinIva": 0}}, {"venta": {"montoSinIva": 10, "moneda": "USD"}},
                                    {"proyecto": {}}])
def test_venta_invalida_se_rechaza(cambio):
    (d,) = _plan([_venta(**cambio)])
    assert d["accion"] == "rechazado"


def test_venta_a_tag_inexistente_no_crea_el_proyecto():
    (d,) = _plan([_venta(tag="NUEVO", proyecto={"tag": "NUEVO", "crear": True, "nombre": "X"})])
    assert d["accion"] == "pendiente"


def test_run_aplica_la_venta_con_nota_de_la_cotizacion(entorno):
    wb = openpyxl.load_workbook(entorno["af"])
    wb[af.HOJA_PROYECTOS].cell(row=2, column=COL["Monto de Venta (sin IVA)"]).value = None
    wb.save(entorno["af"])
    _enviar(entorno["raiz"], _venta())
    resumen = _correr(entorno)
    assert resumen["error"] is None
    celda = _celda(entorno["af"], "Monto de Venta (sin IVA)")
    assert celda.value == 7000000
    assert celda.comment is not None and "Cotización 602695" in celda.comment.text
    assert "Monto de venta enviado desde Sistema QUEMPIN" not in celda.comment.text  # lo envió el Formulador
    proyecto = _catalogo(entorno["raiz"])["proyectos"][0]
    assert proyecto["venta"]["cargada"] is True and proyecto["venta"]["origen"]["mensaje"] == "venta001"
    assert "7000000" not in json.dumps(_catalogo(entorno["raiz"]))  # el monto nunca se publica


def test_el_req_del_envio_completa_la_columna_vacia_y_no_pisa_una_escrita(entorno):
    _enviar(entorno["raiz"], _mensaje(proyecto={"tag": "DEMO", "req": "280"}))
    _correr(entorno)
    assert _celda(entorno["af"], "N° Requerimiento").value == 280
    assert _catalogo(entorno["raiz"])["proyectos"][0]["req"] == "280"

    _enviar(entorno["raiz"], _venta(id_="venta002", proyecto={"tag": "DEMO", "req": "999"}))
    _correr(entorno)
    assert _celda(entorno["af"], "N° Requerimiento").value == 280


def test_la_publicacion_trae_el_sesgo_y_cumple_su_esquema(entorno):
    import esquemas  # Sistema Intercambio, ya en sys.path por presupuestos_formulador
    _enviar(entorno["raiz"], _mensaje())
    _correr(entorno)
    sobre = ic.leer_publicacion(entorno["raiz"], pf.PUBLICACION)
    assert esquemas.validar_publicacion("analisis-financiero", sobre) == []
    sesgo = sobre["datos"]["sesgo"]
    assert set(sesgo["porCategoria"]) == set(pf.CATEGORIAS)
    assert sesgo["proyectosTerminados"] == 0          # DEMO no está terminado ni completo


def test_por_ejecutar_es_lo_proyectado_que_falta_segun_el_avance():
    valores = {"% Avance": 0.4, "Costos Materiales Proyectados": 100000, "Costos Equipos Proyectados": None,
               "Mano de Obra Proyectada": 50000, "Otros Costos Proyectados": True}
    assert af.por_ejecutar(valores) == {"Materiales": 60000, "Equipos": None, "Mano de Obra": 30000, "Otros": None}
    assert af.por_ejecutar({"% Avance": 1.3, "Costos Materiales Proyectados": 100000})["Materiales"] == 0  # acotado
    assert af.por_ejecutar({"% Avance": None, "Costos Materiales Proyectados": 100000}) is None


def test_la_publicacion_trae_cierre_y_por_ejecutar_para_flujo_de_caja(entorno, tmp_path):
    """Flujo de Caja proyecta los egresos con estos dos campos (2026-10-02)."""
    from datetime import datetime
    import esquemas
    entorno["af"] = _excel_af(tmp_path, [("DEMO", "Proyecto de prueba", {
        "Monto de Venta (sin IVA)": 5000000, "% Avance": 0.25, "Fecha de cierre": datetime(2026, 12, 15)})])
    _enviar(entorno["raiz"], _mensaje())
    _correr(entorno)
    sobre = ic.leer_publicacion(entorno["raiz"], pf.PUBLICACION)
    proyecto = sobre["datos"]["proyectos"][0]
    assert proyecto["cierre"] == "2026-12-15"
    assert proyecto["porEjecutar"] == {"Materiales": 750000, "Equipos": 150000, "Mano de Obra": 337500, "Otros": 112500}
    assert esquemas.validar_publicacion("analisis-financiero", sobre) == []


def test_mensajes_de_otro_destino_no_aparecen_en_la_publicacion(entorno):
    ic.asegurar_carpeta(entorno["raiz"])
    otro = _mensaje(id_="otrodest1")
    otro["destino"] = "sistema-quempin"
    otro["tipo"] = "borrador-cotizacion"
    ic.enviar(entorno["raiz"], otro)
    for m in ic.leer_buzon(entorno["raiz"], destino="sistema-quempin")[0]:
        ic.archivar(entorno["raiz"], m, "aplicado")
    _correr(entorno)
    assert "otrodest1" not in _catalogo(entorno["raiz"])["mensajes"]


def test_cliente_vacio_toma_la_razon_social_de_la_cotizacion_emitida(entorno):
    ic.asegurar_carpeta(entorno["raiz"])
    ic.publicar(entorno["raiz"], "documentos-comerciales", "sistema-quempin", {"documentos": [
        {"tipo": "60", "folio": "602695", "pais": "Chile", "fecha": "2026-10-01", "moneda": "CLP", "total": 1,
         "contraparte": {"razon_social": "Universidad de Prueba"}, "proyecto": {"req": None, "tag": "DEMO"}},
        {"tipo": "61", "folio": "612609", "pais": "Chile", "fecha": "2026-10-01", "moneda": "CLP", "total": 1,
         "contraparte": {"razon_social": "Proveedor X"}, "proyecto": {"tag": "DEMO"}},
    ]})
    _correr(entorno)
    assert _celda(entorno["af"], "Cliente").value == "Universidad de Prueba"


def test_tag_cotizado_a_dos_clientes_no_se_adivina():
    import openpyxl as ox
    wb = ox.Workbook()
    ws = wb.active
    for col, h in enumerate(af.HEADERS_PROYECTOS, start=1):
        ws.cell(row=1, column=col, value=h)
    ws.cell(row=2, column=COL["N° Requerimiento"], value=280)
    filas = [{"fila": 2, "tag": "OBRA", "nombre": "Obra"}]
    docs = [{"tipo": "60", "contraparte": {"razon_social": "A"}, "proyecto": {"req": "280"}},
            {"tipo": "60", "contraparte": {"razon_social": "B"}, "proyecto": {"tag": "OBRA"}}]
    assert af.clientes_desde_cotizaciones(ws, filas, docs) == {}
    assert af.clientes_desde_cotizaciones(ws, filas, docs[:1]) == {"OBRA": "A"}


# ── INGRESO MANUAL DESDE EL TABLERO (datos-proyecto, 2026-10-02) ─────────────

def _datos(id_="tablero01", tag="DEMO", valores=None, reemplaza=None, enviado="2026-10-02T10:00:00-03:00",
           proyecto=None, usuario="Persona de prueba"):
    valores = valores if valores is not None else {"% Avance": 0.4, "Fecha de cierre": "2026-12-15"}
    return {
        "esquema": ic.ESQUEMA, "id": id_, "tipo": pf.TIPO_DATOS, "destino": pf.DESTINO,
        "origen": {"herramienta": pf.HERRAMIENTA_TABLERO, "enviado": enviado, "usuario": usuario},
        "proyecto": proyecto or {"tag": tag},
        "valores": valores,
        "reemplaza": reemplaza if reemplaza is not None else {c: None for c in valores},
    }


def test_los_ejemplos_del_catalogo_son_validos_para_este_modulo():
    import esquemas
    carpeta = esquemas.CARPETA / "ejemplos" / "mensajes" / pf.TIPO_DATOS
    for ruta in carpeta.glob("valido*.json"):
        assert pf.validar_contenido(json.loads(ruta.read_text(encoding="utf-8"))) == [], ruta.name


def test_los_campos_del_tablero_son_columnas_manuales_de_proyectos():
    """Lo que escribe el tablero tiene que ser una columna manual (amarilla)
    de "Proyectos" -- nunca una fórmula ni Cliente/Categoría."""
    assert set(pf.COLUMNAS_ESCRIBIBLES.values()) <= set(af.NOMBRES_COLUMNAS_MANUALES_PROYECTOS)
    assert set(af.NOMBRES_COLUMNAS_MANUALES_PROYECTOS) - set(pf.COLUMNAS_ESCRIBIBLES.values()) == {
        "TAG proyecto", "Nombre del proyecto"}


def test_el_procesador_atiende_los_mismos_tipos_que_este_modulo():
    import procesar
    assert set(procesar.TIPOS_AF) == set(pf.TIPOS)


def test_datos_en_celdas_vacias_se_aplican():
    (d,) = _plan([_datos()], actuales={"DEMO": {"% Avance": None, "Fecha de cierre": None}})
    assert d["accion"] == "aplicar"
    assert d["cambios"] == [("% Avance", None, 0.4), ("Fecha de cierre", None, "2026-12-15")]


def test_datos_sobre_el_valor_visto_se_aplican_y_sobre_otro_quedan_pendientes():
    actuales = {"DEMO": {"% Avance": 0.25, "Fecha de cierre": None}}
    (d,) = _plan([_datos(reemplaza={"% Avance": 0.25, "Fecha de cierre": None})], actuales)
    assert d["accion"] == "aplicar"
    (d,) = _plan([_datos(reemplaza={"% Avance": 0.10, "Fecha de cierre": None})], actuales)
    assert d["accion"] == "pendiente" and [c for c, _, _ in d["conflictos"]] == ["% Avance"]
    assert "se vio 10,0 %" in d["detalle"][0] and "ya tiene 25,0 %" in d["detalle"][0]


def test_fecha_del_excel_se_compara_con_la_del_tablero():
    from datetime import datetime
    actuales = {"DEMO": {"Fecha de cierre": datetime(2026, 11, 30)}}
    (d,) = _plan([_datos(valores={"Fecha de cierre": "2026-12-15"},
                         reemplaza={"Fecha de cierre": "2026-11-30"})], actuales)
    assert d["accion"] == "aplicar" and d["cambios"] == [("Fecha de cierre", "2026-11-30", "2026-12-15")]
    (d,) = _plan([_datos(valores={"Fecha de cierre": "2026-11-30"})], actuales)
    assert d["accion"] == "sin-cambios"


def test_valor_que_escribio_el_formulador_no_se_pisa_sin_haberlo_visto():
    actuales = {"DEMO": {"Materiales": 900000}}
    estado = {"valores": {"DEMO": {"Materiales": {"valor": 900000}}}, "confirmados": []}
    (d,) = _plan([_datos(valores={"Materiales": 1200000})], actuales, estado)
    assert d["accion"] == "pendiente"
    (d,) = _plan([_datos(valores={"Materiales": 1200000}, reemplaza={"Materiales": 900000})], actuales, estado)
    assert d["accion"] == "aplicar"


def test_varios_ingresos_al_mismo_proyecto_se_aplican_todos_en_orden():
    """«Gana el último» perdería el avance del primero: son ediciones sueltas."""
    primero = _datos("tablero01", valores={"% Avance": 0.4}, reemplaza={"% Avance": 0.2})
    segundo = _datos("tablero02", valores={"% Avance": 0.6, "Mano de Obra Real": 50000},
                     reemplaza={"% Avance": 0.4, "Mano de Obra Real": None}, enviado="2026-10-02T11:00:00-03:00")
    otro = _datos("tablero03", valores={"% Avance": 0.9}, reemplaza={"% Avance": 0.2},
                  enviado="2026-10-02T12:00:00-03:00")
    decisiones = _plan([primero, segundo, otro], {"DEMO": {"% Avance": 0.2, "Mano de Obra Real": None}})
    assert [d["accion"] for d in decisiones] == ["aplicar", "aplicar", "pendiente"]


def test_ingreso_y_formulador_en_la_misma_corrida_se_ven_entre_si():
    manual = _datos(valores={"Materiales": 1200000}, enviado="2026-10-02T09:00:00-03:00")
    formulador = _mensaje(enviado="2026-10-02T10:00:00-03:00")  # no vio el 1.200.000 recién ingresado
    decisiones = _plan([manual, formulador])
    assert [d["accion"] for d in decisiones] == ["aplicar", "pendiente"]


def test_proyecto_nuevo_desde_el_tablero_y_luego_otro_ingreso_al_mismo():
    crear = _datos(proyecto={"tag": "NUEVO", "nombre": "Obra nueva", "crear": True},
                   valores={"Monto de Venta": 8000000, "% Avance": 0})
    despues = _datos("tablero02", tag="NUEVO", valores={"% Avance": 0.1}, reemplaza={"% Avance": 0},
                     enviado="2026-10-02T11:00:00-03:00")
    decisiones = _plan([crear, despues])
    assert [d["accion"] for d in decisiones] == ["crear", "aplicar"]
    (d,) = _plan([_datos(tag="NADA")])
    assert d["accion"] == "pendiente" and "no existe" in d["detalle"][0]


@pytest.mark.parametrize("cambio", [
    {"valores": {}},
    {"valores": {"% Avance": 40}},
    {"valores": {"Fecha de cierre": "2026-02-30"}},
    {"valores": {"Cliente": "Otro"}},
    {"valores": {"Monto de Venta": 0}},
    {"valores": {"N° Requerimiento": 280.5}},
    {"valores": {"Materiales": -5}},
    {"proyecto": {"tag": "obra 1"}},
    {"proyecto": {"tag": "OBRA", "crear": True}},
])
def test_ingreso_invalido_se_rechaza(cambio):
    m = _datos()
    m.update(cambio)
    (d,) = _plan([m])
    assert d["accion"] == "rechazado" and d["detalle"]


def test_run_escribe_el_ingreso_con_tipos_y_formatos_de_excel_y_sin_nota(entorno):
    from datetime import datetime
    _enviar(entorno["raiz"], _datos(valores={
        "% Avance": 0.35, "Fecha de inicio": "2026-09-01", "Fecha de cierre": "2026-12-15",
        "Mano de Obra Real": 420000.4, "N° Requerimiento": 280, "Materiales": 1000000,
    }))
    resumen = _correr(entorno)
    assert resumen["error"] is None and resumen["intercambio"]["aplicar"]
    assert _celda(entorno["af"], "% Avance").value == 0.35
    assert _celda(entorno["af"], "% Avance").number_format == "0.0%"
    assert _celda(entorno["af"], "Fecha de inicio").value == datetime(2026, 9, 1)
    assert _celda(entorno["af"], "Fecha de cierre").number_format == af.FORMATO_FECHA
    assert _celda(entorno["af"], "Mano de Obra Real").value == 420000
    assert _celda(entorno["af"], "Mano de Obra Real").number_format == af.FORMATO_MONEDA
    assert _celda(entorno["af"], "N° Requerimiento").value == 280
    assert _celda(entorno["af"], "Costos Materiales Proyectados").comment is None   # es manual: sin nota
    estado = json.loads(entorno["estado"].read_text(encoding="utf-8"))
    assert "DEMO" not in estado["valores"]                                      # ni registro del Formulador
    procesado = _procesados(entorno["raiz"])["tablero01"]
    assert procesado["resultado"]["estado"] == "aplicado"
    assert procesado["resultado"]["cambios"]["% Avance"] == {"antes": None, "despues": 0.35}
    assert _catalogo(entorno["raiz"])["mensajes"]["tablero01"]["estado"] == "aplicado"


def test_run_respeta_el_formato_que_la_celda_ya_tenia(entorno):
    wb = openpyxl.load_workbook(entorno["af"])
    celda = wb[af.HOJA_PROYECTOS].cell(row=2, column=COL["Monto de Venta (sin IVA)"])
    celda.number_format = '#,##0 "pesos"'
    wb.save(entorno["af"])
    _enviar(entorno["raiz"], _datos(valores={"Monto de Venta": 6000000}, reemplaza={"Monto de Venta": 5000000}))
    _correr(entorno)
    celda = _celda(entorno["af"], "Monto de Venta (sin IVA)")
    assert celda.value == 6000000 and celda.number_format == '#,##0 "pesos"'


def test_run_ingreso_sobre_un_valor_del_formulador_le_quita_la_nota_y_el_registro(entorno):
    _enviar(entorno["raiz"], _mensaje())
    _correr(entorno)
    assert _celda(entorno["af"], "Costos Materiales Proyectados").comment is not None
    _enviar(entorno["raiz"], _datos(valores={"Materiales": 1100000}, reemplaza={"Materiales": 1000000}))
    _correr(entorno)
    celda = _celda(entorno["af"], "Costos Materiales Proyectados")
    assert celda.value == 1100000 and celda.comment is None
    assert _celda(entorno["af"], "Costos Equipos Proyectados").comment is not None   # las demás siguen del Formulador
    estado = json.loads(entorno["estado"].read_text(encoding="utf-8"))
    assert "Materiales" not in estado["valores"]["DEMO"] and "Equipos" in estado["valores"]["DEMO"]
    assert _catalogo(entorno["raiz"])["proyectos"][0]["origen"]["Materiales"] is None


def test_run_puede_dejar_una_celda_vacia(entorno):
    _enviar(entorno["raiz"], _datos(valores={"Monto de Venta": None}, reemplaza={"Monto de Venta": 5000000}))
    _correr(entorno)
    assert _celda(entorno["af"], "Monto de Venta (sin IVA)").value is None


def test_run_crea_el_proyecto_con_todos_sus_datos(entorno):
    _enviar(entorno["raiz"], _datos(proyecto={"tag": "NUEVO", "nombre": "Obra nueva", "crear": True},
                                    valores={"Monto de Venta": 8000000, "% Avance": 0, "Otros": 300000}))
    resumen = _correr(entorno)
    assert "Obra nueva" in resumen["proyectos_nuevos"]
    ws = openpyxl.load_workbook(entorno["af"])[af.HOJA_PROYECTOS]
    assert ws.cell(row=3, column=1).value == "NUEVO" and ws.cell(row=3, column=2).value == "Obra nueva"
    assert ws.cell(row=3, column=COL["Monto de Venta (sin IVA)"]).value == 8000000
    assert ws.cell(row=3, column=COL["% Avance"]).value == 0
    assert ws.cell(row=3, column=COL["Otros Costos Proyectados"]).value == 300000


def test_run_ingreso_en_conflicto_no_escribe_nada_y_se_confirma(entorno):
    _enviar(entorno["raiz"], _datos(valores={"Monto de Venta": 6000000, "% Avance": 0.5},
                                    reemplaza={"Monto de Venta": 4000000, "% Avance": None}))
    resumen = _correr(entorno)
    assert resumen["intercambio"]["pendiente"]
    assert _celda(entorno["af"], "Monto de Venta (sin IVA)").value == 5000000
    assert _celda(entorno["af"], "% Avance").value is None                     # nunca a medias
    assert pf.confirmar("tablero01", entorno["raiz"], entorno["estado"])
    _correr(entorno)
    assert _celda(entorno["af"], "Monto de Venta (sin IVA)").value == 6000000
    assert _celda(entorno["af"], "% Avance").value == 0.5


def test_cargar_archivo_deja_el_envio_en_el_buzon_una_sola_vez(entorno, tmp_path):
    archivo = tmp_path / "descargado.json"
    archivo.write_text(json.dumps(_datos()), encoding="utf-8")
    ok, texto = pf.cargar_archivo(entorno["raiz"], archivo)
    assert ok and "DEMO" in texto
    assert len(list((entorno["raiz"] / "buzon").glob("*.json"))) == 1
    ok, texto = pf.cargar_archivo(entorno["raiz"], archivo)
    assert not ok and "ya está en el buzón" in texto
    _correr(entorno)
    ok, texto = pf.cargar_archivo(entorno["raiz"], archivo)
    assert not ok and "ya se atendió" in texto
    malo = tmp_path / "malo.json"
    malo.write_text(json.dumps(_datos(valores={"% Avance": 40})), encoding="utf-8")
    assert pf.cargar_archivo(entorno["raiz"], malo)[0] is False
