# -*- coding: utf-8 -*-
"""
presupuestos_formulador.py -- Costos proyectados que llegan desde el
Formulador de proyectos por la carpeta de Intercambio (protocolo común en
``Sistema Intercambio/intercambio.py``).

Pedido del usuario (2026-09-30): que las herramientas compartan información
entre sí y, en concreto, que el Formulador pueda actualizar los costos
estimados del Análisis Financiero. Se implementa como un **canal de entrada
auditado** a las 4 columnas manuales "... Proyectado(s)" de la hoja
"Proyectos" -- la única excepción a la regla de oro "las columnas manuales
nunca se tocan", y con estas garantías:

- Solo escribe cuando es seguro: la celda está vacía, o tiene el valor que
  el propio Formulador escribió antes (registro de procedencia), o la
  persona que envió vio ese mismo valor en el Formulador y decidió
  reemplazarlo ("reemplaza" del mensaje, como un If-Match). Cualquier otro
  caso -- un valor escrito a mano que quien envió no vio -- queda
  **pendiente** en el buzón hasta que el usuario lo confirme o lo descarte
  (``driver.py intercambio confirmar|descartar <id>``).
- Un mensaje se aplica completo o no se aplica: nunca la mitad de sus
  categorías.
- Cada celda escrita lleva una nota de Excel con su procedencia (oferta,
  versión, quién la envió y cuándo). Si después alguien la cambia a mano,
  la corrida siguiente quita la nota y la deja como manual.
- Los mensajes se archivan (``procesado/``) y el registro de procedencia se
  guarda **solo después de que el Excel se guardó bien**: si el libro está
  abierto y no se puede guardar, todo queda en el buzón para la próxima.
- Nunca toca Centro de Costos ni su carpeta de facturas (el ingreso desde
  SharePoint sigue igual): solo la hoja "Proyectos" de este módulo.

**Venta (2026-10-01, plan de integración).** El mismo canal, con las mismas
garantías, atiende el mensaje ``venta-proyecto``: el monto de venta sin IVA
de un proyecto adjudicado (de la cotización emitida en Sistema QUEMPIN, o
del precio neto del Formulador si no hay cotización) va a «Monto de Venta
(sin IVA)». No es un segundo canal: es otra clave del mismo registro de
procedencia (``CLAVE_VENTA``), para que las dos reglas no diverjan. Si un
envío trae ``proyecto.req`` y la fila no tiene «N° Requerimiento», se
completa (nunca se pisa).

**Ingreso manual desde el tablero (2026-10-02).** Pedido del usuario: una
pestaña del tablero de Análisis Financiero donde se ingresen los valores
manuales de «Proyectos» (% Avance, fechas, venta, los 4 proyectados, Mano de
Obra Real y N° Requerimiento) en vez de escribirlos en el Excel. La pestaña
deja un mensaje ``datos-proyecto`` por proyecto en el mismo buzón, y se
aplica aquí con la misma regla: la celda vacía se escribe, y una con valor
solo si todavía tiene lo que la persona vio en el tablero (``reemplaza``);
si no, el envío queda pendiente. Tres diferencias con el Formulador, porque
esto es ingreso manual y no una copia del estado de otra herramienta:

- Varios envíos al mismo proyecto se aplican **todos, en orden** (cada uno
  sobre lo que dejó el anterior): son ediciones sueltas, no el estado
  completo de un presupuesto, así que «gana el último» perdería datos.
- No lleva nota de procedencia ni entra al registro de lo que escribió el
  Formulador: queda como valor manual. Si pisa un valor que había escrito el
  Formulador, le quita su nota y lo saca del registro (desde ahí es manual).
- Puede dejar una celda vacía (``null``) y crear un proyecto con todos sus
  datos.

Al final de cada corrida real publica ``publicado/analisis-financiero.json``
con los proyectos (TAG, nombre, cliente, avance, costos proyectados y
reales por categoría, y de dónde viene cada proyectado) y el estado de los
mensajes recibidos: con eso el Formulador ofrece a qué proyecto enviar,
muestra qué va a cambiar y sabe si su envío ya se aplicó.

No importa analisis_financiero (evita el import circular): quien llama le
pasa la hoja, las filas y un dict {encabezado: número de columna}.
"""

import json
import math
import os
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl.comments import Comment

RAIZ = Path(__file__).resolve().parent                  # Sistema Analisis Financiero/Sistema
RAIZ_FINANZAS = RAIZ.parents[1]                          # Finanzas QUEMPIN
RAIZ_SISTEMA_INTERCAMBIO = RAIZ_FINANZAS / "Sistema Intercambio"
if str(RAIZ_SISTEMA_INTERCAMBIO) not in sys.path:
    sys.path.insert(0, str(RAIZ_SISTEMA_INTERCAMBIO))
import intercambio  # noqa: E402

# Carpeta de intercambio, compartida con el equipo. Desde el 2026-09-30
# (pedido explícito del usuario) vive en «.Herramientas formulación/Intercambio»,
# dentro de la biblioteca de SharePoint «Formulación de proyectos», que los
# colegas ya tienen sincronizada: ahí vive también el repositorio de las
# formulaciones (publicado/formulador/). Antes estaba en Finanzas
# QUEMPIN/Intercambio. Se busca subiendo desde este repo hasta la carpeta de
# OneDrive que contiene la biblioteca; la variable QUEMPIN_INTERCAMBIO la fija
# a mano. Si la biblioteca no está sincronizada en este equipo, el intercambio
# queda inactivo con un aviso: nunca se crea una biblioteca falsa ni se vuelve
# a la carpeta antigua (sería un buzón que nadie más lee).
BIBLIOTECA_FORMULACION = "Formulación de proyectos - Documentos"
CARPETA_HERRAMIENTAS = ".Herramientas formulación"
CARPETA_INTERCAMBIO = "Intercambio"


def _hija(padre: Path, nombre: str) -> Path:
    """La subcarpeta con ese nombre sin distinguir mayúsculas ni la forma de
    las tildes (en el disco está como «.HERRAMIENTAS FORMULACIÓN»)."""
    clave = unicodedata.normalize("NFC", nombre).casefold()
    try:
        for hija in padre.iterdir():
            if hija.is_dir() and unicodedata.normalize("NFC", hija.name).casefold() == clave:
                return hija
    except OSError:
        pass
    return padre / nombre


def ubicar_intercambio() -> Path | None:
    fija = os.environ.get("QUEMPIN_INTERCAMBIO")
    if fija:
        return Path(fija)
    for base in RAIZ_FINANZAS.parents:
        biblioteca = base / BIBLIOTECA_FORMULACION
        if biblioteca.is_dir():
            return _hija(_hija(biblioteca, CARPETA_HERRAMIENTAS), CARPETA_INTERCAMBIO)
    return None


RAIZ_INTERCAMBIO = ubicar_intercambio()
AVISO_SIN_INTERCAMBIO = (
    f"Intercambio con el Formulador inactivo: no se encontró la biblioteca «{BIBLIOTECA_FORMULACION}» "
    "sincronizada en este equipo (o fija la carpeta con la variable QUEMPIN_INTERCAMBIO)."
)
# Registro de procedencia (lleva montos: ignorado por git).
RUTA_ESTADO = RAIZ / "presupuestos_formulador.json"

TIPO = "presupuesto-proyecto"
TIPO_VENTA = "venta-proyecto"
# Ingreso manual desde la pestaña «Ingresar datos» del tablero (2026-10-02).
TIPO_DATOS = "datos-proyecto"
TIPOS = (TIPO, TIPO_VENTA, TIPO_DATOS)
# Los del Formulador traen el estado completo de un presupuesto o una venta:
# de varios al mismo TAG cuenta el último. Los del tablero se aplican todos.
TIPOS_FORMULADOR = (TIPO, TIPO_VENTA)
HERRAMIENTA_TABLERO = "tablero-af"
DESTINO = "analisis-financiero"
HERRAMIENTA = "analisis-financiero"
PUBLICACION = "analisis-financiero"
AUTOR_NOTA = "Formulador (Intercambio)"

# Categoría del mensaje -> columna manual de "Proyectos". Los gastos
# generales e imprevistos del Formulador no viajan: este módulo compara
# contra costos directos del proyecto (los gastos generales de la empresa
# tienen su propio proyecto, "Gastos Generales").
COLUMNA_POR_CATEGORIA = {
    "Materiales": "Costos Materiales Proyectados",
    "Equipos": "Costos Equipos Proyectados",
    "Mano de Obra": "Mano de Obra Proyectada",
    "Otros": "Otros Costos Proyectados",
}
CATEGORIAS = tuple(COLUMNA_POR_CATEGORIA)
# La venta es una clave más del mismo registro de procedencia.
CLAVE_VENTA = "Monto de Venta"
COLUMNA_VENTA = "Monto de Venta (sin IVA)"
COLUMNA_REQ = "N° Requerimiento"
COLUMNAS_INTERCAMBIO = {**COLUMNA_POR_CATEGORIA, CLAVE_VENTA: COLUMNA_VENTA}
# Lo que escribe el tablero: clave del mensaje -> (columna de "Proyectos",
# tipo de valor). Costos y venta usan la misma clave que el Formulador: así
# los dos hablan de la misma celda con el mismo nombre, en el mismo plan.
# El orden es el de la pestaña. TAG y Nombre no están: identifican la fila
# (Nombre solo se escribe al crear un proyecto).
CAMPOS_TABLERO = {
    "% Avance": ("% Avance", "fraccion"),
    "Fecha de inicio": ("Fecha de inicio", "fecha"),
    "Fecha de cierre": ("Fecha de cierre", "fecha"),
    CLAVE_VENTA: (COLUMNA_VENTA, "pesos"),
    **{c: (col, "pesos") for c, col in COLUMNA_POR_CATEGORIA.items()},
    "Mano de Obra Real": ("Mano de Obra Real", "pesos"),
    COLUMNA_REQ: (COLUMNA_REQ, "entero"),
}
# Toda columna que puede escribir este canal, venga de donde venga el envío.
COLUMNAS_ESCRIBIBLES = {clave: columna for clave, (columna, _) in CAMPOS_TABLERO.items()}
TIPO_DE_CAMPO = {clave: tipo for clave, (_, tipo) in CAMPOS_TABLERO.items()}
# Mismo formato que el prefijo del N° Ref de Centro de Costos (UMAG, FCH1...).
PATRON_TAG = re.compile(r"^[A-Z0-9]{2,10}$")

ESTADO_FINAL_POR_ACCION = {
    "aplicar": "aplicado", "crear": "aplicado", "sin-cambios": "sin-cambios",
    "reemplazado": "reemplazado", "rechazado": "rechazado",
}


# ── REGISTRO DE PROCEDENCIA ──────────────────────────────────────────────────
# {"valores": {TAG: {categoría: {valor, mensaje, fuente, enviado_por,
#   aplicado}}}, "confirmados": [id, ...]}

def leer_estado(ruta: Path = RUTA_ESTADO) -> dict:
    try:
        with open(ruta, encoding="utf-8") as f:
            estado = json.load(f)
    except (OSError, ValueError):
        estado = {}
    if not isinstance(estado, dict):
        estado = {}
    estado.setdefault("valores", {})
    estado.setdefault("confirmados", [])
    return estado


def guardar_estado(estado: dict, ruta: Path = RUTA_ESTADO) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    temporal = ruta.with_name(f".{ruta.name}.tmp")
    with open(temporal, "w", encoding="utf-8") as f:
        json.dump(estado, f, ensure_ascii=False, indent=2)
    os.replace(temporal, ruta)


# ── NÚMEROS Y TEXTOS ─────────────────────────────────────────────────────────

def _numero(valor):
    """Número de una celda o de un JSON; None si está vacía. Un texto que no
    es número se devuelve tal cual (cuenta como valor escrito a mano)."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, bool):
        return valor
    if isinstance(valor, (int, float)):
        return float(valor) if math.isfinite(valor) else None
    return valor


def _iguales(a, b) -> bool:
    return (
        isinstance(a, float) and isinstance(b, (int, float)) and not isinstance(b, bool)
        and abs(a - float(b)) < 0.5
    )


def _pesos(valor) -> str:
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return "$" + f"{round(valor):,}".replace(",", ".")
    return "(vacío)" if valor is None else f"«{valor}»"


def normalizar_tag(tag) -> str:
    return str(tag or "").strip().upper()


def _legible(clave: str, valor) -> str:
    """Un valor de este canal como se lee en la consola y en el archivo:
    pesos, porcentaje, fecha DD-MM-AAAA o número, según el campo."""
    tipo = TIPO_DE_CAMPO.get(clave, "pesos")
    valor = _normalizar(clave, valor)
    if valor is None:
        return "(vacío)"
    if isinstance(valor, float) and tipo == "fraccion":
        return f"{valor * 100:.1f}".replace(".", ",") + " %"
    if isinstance(valor, float) and tipo == "entero":
        return str(int(valor)) if valor.is_integer() else str(valor)
    if tipo == "fecha" and isinstance(valor, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", valor):
        return "-".join(reversed(valor.split("-")))
    return _pesos(valor)


def _normalizar(clave: str, valor):
    """El valor de una celda o de un mensaje en la forma en que se compara:
    número (float) o 'AAAA-MM-DD' para las fechas. Vacío -> None. Un texto
    que no es número ni fecha se devuelve tal cual: cuenta como valor
    escrito a mano."""
    if TIPO_DE_CAMPO.get(clave) == "fecha":
        if valor is None or (isinstance(valor, str) and not valor.strip()):
            return None
        if isinstance(valor, (datetime, date)):
            return (valor.date() if isinstance(valor, datetime) else valor).isoformat()
        return str(valor).strip()
    numero = _numero(valor)
    if isinstance(numero, str):
        texto = numero.strip()
        try:
            return float(texto) if re.fullmatch(r"-?\d+(\.\d+)?", texto) else texto
        except ValueError:
            return texto
    return numero


def _mismo_valor(clave: str, a, b) -> bool:
    """¿Son el mismo dato? Pesos con tolerancia de medio peso (el Excel
    guarda decimales que el tablero muestra redondeados) y el avance con
    tolerancia de una millonésima."""
    a, b = _normalizar(clave, a), _normalizar(clave, b)
    if a is None or b is None:
        return a is None and b is None
    if isinstance(a, float) and isinstance(b, float):
        return abs(a - b) < (1e-6 if TIPO_DE_CAMPO.get(clave) == "fraccion" else 0.5)
    return a == b


def _valor_celda(clave: str, valor):
    """Lo que se escribe en la celda: fechas como fecha de Excel, pesos y
    N° de requerimiento como enteros, el avance como fracción."""
    tipo = TIPO_DE_CAMPO.get(clave, "pesos")
    if valor is None:
        return None
    if tipo == "fecha":
        return datetime.strptime(valor, "%Y-%m-%d")
    if tipo == "fraccion":
        return float(valor)
    return int(round(valor))


def fuente_legible(mensaje: dict) -> str:
    """'Formulación 261 v2 · Título' (o 'Cotización 602695 · Título' para una
    venta que viene de Sistema QUEMPIN, o 'Tablero · Persona' para un ingreso
    manual) -- lo que se muestra en la nota y en la consola."""
    if mensaje.get("tipo") == TIPO_DATOS:
        quien = str((mensaje.get("origen") or {}).get("usuario") or "").strip()
        return "Ingreso en el tablero" + (f" · {quien}" if quien else "")
    f = mensaje.get("fuente") or {}
    if f.get("folio"):
        titulo = str(f.get("titulo") or "").strip()
        return f"Cotización {f['folio']}" + (f" · {titulo}" if titulo else "")
    codigo = str(f.get("codigo") or "").strip() or "sin código"
    texto = f"Formulación {codigo} v{f.get('version') or 1}"
    titulo = str(f.get("titulo") or "").strip()
    return f"{texto} · {titulo}" if titulo else texto


def _fecha_corta(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso).strftime("%d-%m-%Y")
    except (TypeError, ValueError):
        return str(iso or "")


# ── VALIDACIÓN DEL CONTENIDO ─────────────────────────────────────────────────

def validar_presupuesto(mensaje: dict) -> list[str]:
    """Errores del contenido de un mensaje 'presupuesto-proyecto' (el sobre
    ya lo validó intercambio.validar_mensaje)."""
    errores = []
    proyecto = mensaje.get("proyecto")
    if not isinstance(proyecto, dict) or not normalizar_tag(proyecto.get("tag")):
        errores.append("falta el TAG del proyecto de destino")
        proyecto = {}
    costos = mensaje.get("costos")
    if not isinstance(costos, dict) or not costos:
        errores.append("no trae costos")
    else:
        for categoria, valor in costos.items():
            if categoria not in COLUMNA_POR_CATEGORIA:
                errores.append(f"categoría desconocida: {categoria}")
            elif (not isinstance(valor, (int, float)) or isinstance(valor, bool)
                  or not math.isfinite(valor) or valor < 0):
                errores.append(f"{categoria}: el costo debe ser un número mayor o igual a 0")
    if proyecto.get("crear"):
        if not PATRON_TAG.match(normalizar_tag(proyecto.get("tag"))):
            errores.append("el TAG de un proyecto nuevo debe tener 2 a 10 letras o números (ej. UMAG)")
        if not str(proyecto.get("nombre") or "").strip():
            errores.append("falta el nombre del proyecto nuevo")
    return errores


def validar_venta(mensaje: dict) -> list[str]:
    """Errores del contenido de un mensaje 'venta-proyecto'."""
    errores = []
    proyecto = mensaje.get("proyecto")
    if not isinstance(proyecto, dict) or not normalizar_tag(proyecto.get("tag")):
        errores.append("falta el TAG del proyecto de destino")
    venta = mensaje.get("venta") if isinstance(mensaje.get("venta"), dict) else {}
    monto = venta.get("montoSinIva")
    if (not isinstance(monto, (int, float)) or isinstance(monto, bool)
            or not math.isfinite(monto) or monto <= 0):
        errores.append("el monto de venta sin IVA debe ser un número mayor que 0")
    if venta.get("moneda", "CLP") != "CLP":
        errores.append("el monto de venta tiene que venir en pesos (CLP)")
    if not isinstance(mensaje.get("fuente"), dict):
        errores.append("falta la fuente del monto (cotización o formulación)")
    return errores


def _es_numero(valor) -> bool:
    return isinstance(valor, (int, float)) and not isinstance(valor, bool) and math.isfinite(valor)


def validar_datos(mensaje: dict) -> list[str]:
    """Errores del contenido de un mensaje 'datos-proyecto' (ingreso manual
    desde el tablero). Las mismas reglas que su esquema, más las que el
    esquema no puede decir (al menos un valor; fechas que existen)."""
    errores = []
    proyecto = mensaje.get("proyecto")
    if not isinstance(proyecto, dict) or not normalizar_tag(proyecto.get("tag")):
        errores.append("falta el TAG del proyecto")
        proyecto = {}
    elif not PATRON_TAG.match(normalizar_tag(proyecto.get("tag"))):
        errores.append("el TAG debe tener 2 a 10 letras o números (ej. UMAG)")
    if proyecto.get("crear") and not str(proyecto.get("nombre") or "").strip():
        errores.append("falta el nombre del proyecto nuevo")
    valores = mensaje.get("valores")
    if not isinstance(valores, dict) or not valores:
        return errores + ["no trae ningún valor"]
    if not isinstance(mensaje.get("reemplaza"), dict):
        errores.append("falta 'reemplaza' (lo que se vio en el tablero)")
    for clave, valor in valores.items():
        tipo = TIPO_DE_CAMPO.get(clave)
        if tipo is None:
            errores.append(f"campo desconocido: {clave}")
        elif valor is None:
            continue
        elif tipo == "fecha":
            try:
                datetime.strptime(str(valor), "%Y-%m-%d")
            except ValueError:
                errores.append(f"{clave}: la fecha debe venir como AAAA-MM-DD")
        elif not _es_numero(valor):
            errores.append(f"{clave}: debe ser un número")
        elif tipo == "fraccion" and not 0 <= valor <= 1:
            errores.append(f"{clave}: debe estar entre 0 y 1 (0,35 = 35 %)")
        elif tipo == "entero" and (not float(valor).is_integer() or valor < 1):
            errores.append(f"{clave}: debe ser un número entero mayor que 0")
        elif clave == CLAVE_VENTA and valor <= 0:
            errores.append(f"{clave}: debe ser mayor que 0")
        elif valor < 0:
            errores.append(f"{clave}: no puede ser negativo")
    return errores


def validar_contenido(mensaje: dict) -> list[str]:
    if mensaje.get("tipo") == TIPO_DATOS:
        return validar_datos(mensaje)
    return validar_venta(mensaje) if mensaje.get("tipo") == TIPO_VENTA else validar_presupuesto(mensaje)


def valores_mensaje(mensaje: dict) -> dict:
    """{clave: valor} que el mensaje quiere escribir: pesos enteros (y, en
    el ingreso manual, el avance como fracción, fechas 'AAAA-MM-DD', el N° de
    requerimiento como entero y None para dejar la celda vacía)."""
    if mensaje.get("tipo") == TIPO_VENTA:
        return {CLAVE_VENTA: int(round(mensaje["venta"]["montoSinIva"]))}
    if mensaje.get("tipo") == TIPO_DATOS:
        salida = {}
        for clave, valor in mensaje["valores"].items():
            tipo = TIPO_DE_CAMPO[clave]
            if valor is None or tipo == "fecha":
                salida[clave] = valor
            elif tipo == "fraccion":
                salida[clave] = float(valor)
            else:
                salida[clave] = int(round(valor))
        return salida
    return {c: int(round(v)) for c, v in mensaje["costos"].items()}


def reemplaza_mensaje(mensaje: dict) -> dict:
    """El If-Match del mensaje, con las mismas claves que valores_mensaje."""
    reemplaza = mensaje.get("reemplaza") if isinstance(mensaje.get("reemplaza"), dict) else {}
    if mensaje.get("tipo") == TIPO_VENTA:
        return {CLAVE_VENTA: reemplaza["montoSinIva"]} if "montoSinIva" in reemplaza else {}
    return reemplaza


def _quien(mensaje: dict) -> str:
    herramienta = (mensaje.get("origen") or {}).get("herramienta")
    if mensaje.get("tipo") == TIPO_DATOS:
        return "el tablero"
    return "Sistema QUEMPIN" if herramienta == "sistema-quempin" else "el Formulador"


# ── PLAN ─────────────────────────────────────────────────────────────────────

def leer_buzon(raiz_intercambio: Path):
    """(mensajes para este módulo, avisos por archivos ilegibles del buzón).
    Los tipos que atiende: presupuesto-proyecto y venta-proyecto (del
    Formulador) y datos-proyecto (ingreso manual desde el tablero)."""
    validos, invalidos = intercambio.leer_buzon(raiz_intercambio, destino=DESTINO)
    validos = [m for m in validos if m["tipo"] in TIPOS]
    avisos = [
        f"Intercambio: el archivo '{m['_archivo'].name}' del buzón no es un mensaje válido "
        f"({'; '.join(m['errores'])}) -- se deja donde está."
        for m in invalidos
    ]
    return validos, avisos


def valores_actuales(ws, fila: int, columnas: dict) -> dict:
    return {c: ws.cell(row=fila, column=columnas[col]).value
            for c, col in COLUMNAS_ESCRIBIBLES.items() if col in columnas}


def planificar(mensajes: list[dict], filas_por_tag: dict, actuales_por_tag: dict, estado: dict) -> list[dict]:
    """Una decisión por mensaje ('aplicar', 'crear', 'sin-cambios',
    'pendiente', 'reemplazado' o 'rechazado'). Función pura: no escribe
    nada. 'mensajes' viene ordenado por fecha de envío y se decide en ese
    orden, cada uno sobre los valores que van dejando los anteriores (un
    ingreso del tablero ve lo que acaba de aplicar otro). De varios del
    Formulador del mismo tipo para el mismo TAG solo cuenta el último (un
    presupuesto y una venta del mismo TAG son independientes: escriben
    columnas distintas); los del tablero se aplican todos."""
    decisiones = []
    vigentes = []
    ultimo: dict[tuple[str, str], dict] = {}
    for mensaje in mensajes:
        errores = validar_contenido(mensaje)
        if errores:
            decisiones.append({"mensaje": mensaje, "accion": "rechazado", "tag": None, "costos": {},
                               "detalle": errores})
            continue
        vigentes.append(mensaje)
        if mensaje.get("tipo", TIPO) in TIPOS_FORMULADOR:
            ultimo[(normalizar_tag(mensaje["proyecto"]["tag"]), mensaje.get("tipo", TIPO))] = mensaje

    # Copias: lo que deja cada decisión es lo que ve la siguiente, sin tocar
    # lo que pasó quien llama. Un proyecto que se crea en esta misma corrida
    # existe para los envíos que vienen después (su fila la pone aplicar()).
    filas = dict(filas_por_tag)
    actuales = {tag: dict(valores) for tag, valores in actuales_por_tag.items()}
    for mensaje in vigentes:
        tag, tipo = normalizar_tag(mensaje["proyecto"]["tag"]), mensaje.get("tipo", TIPO)
        if tipo in TIPOS_FORMULADOR and ultimo[(tag, tipo)] is not mensaje:
            decisiones.append({
                "mensaje": mensaje, "accion": "reemplazado", "tag": tag, "costos": {},
                "detalle": [f"Llegó un envío más reciente para {tag} ({fuente_legible(ultimo[(tag, tipo)])})."],
            })
            continue
        decidir = _decidir_datos if tipo == TIPO_DATOS else _decidir
        d = decidir(mensaje, tag, filas, actuales, estado)
        decisiones.append(d)
        if d["accion"] in ("aplicar", "crear"):
            filas.setdefault(tag, None)
            for clave, _antes, despues in d["cambios"]:
                actuales.setdefault(tag, {})[clave] = despues
    return decisiones


def _decidir_datos(mensaje, tag, filas_por_tag, actuales_por_tag, estado) -> dict:
    """La decisión de un ingreso manual del tablero. Cada valor se escribe si
    la celda está vacía, si ya tiene ese valor (nada que hacer) o si todavía
    tiene lo que la persona vio ('reemplaza'). Lo que escribió el Formulador
    no se pisa por ser suyo: quien no lo vio no lo reemplaza. Todo o nada."""
    nuevos = valores_mensaje(mensaje)
    base = {"mensaje": mensaje, "tag": tag, "costos": nuevos}
    if tag not in filas_por_tag:
        if mensaje["proyecto"].get("crear"):
            nombre = str(mensaje["proyecto"]["nombre"]).strip()
            cambios = [(c, None, v) for c, v in nuevos.items() if v is not None]
            return dict(base, accion="crear", nombre=nombre, cambios=cambios,
                        detalle=[f"Proyecto nuevo {tag} «{nombre}»."]
                        + [f"{c}: {_legible(c, v)}" for c, _, v in cambios])
        return dict(base, accion="pendiente", detalle=[
            f"El TAG {tag} no existe en la hoja Proyectos. Se aplicará cuando exista; si el TAG "
            f"está mal, descártalo (driver.py intercambio descartar {mensaje['id']}) y vuelve a "
            "ingresarlo en el tablero."
        ])

    actuales = actuales_por_tag.get(tag, {})
    reemplaza = reemplaza_mensaje(mensaje)
    confirmado = mensaje["id"] in estado.get("confirmados", [])
    cambios, conflictos = [], []
    for clave, nuevo in nuevos.items():
        actual = actuales.get(clave)
        if _mismo_valor(clave, actual, nuevo):
            continue
        if (_normalizar(clave, actual) is None
                or (clave in reemplaza and _mismo_valor(clave, actual, reemplaza[clave]))
                or confirmado):
            cambios.append((clave, _normalizar(clave, actual), nuevo))
        else:
            conflictos.append((clave, _normalizar(clave, actual), nuevo))

    if conflictos:
        return dict(base, accion="pendiente", cambios=cambios, conflictos=conflictos, detalle=[
            f"{c}: en el tablero se vio {_legible(c, reemplaza.get(c))}, pero el Excel ya tiene "
            f"{_legible(c, a)}; el envío trae {_legible(c, n)}."
            for c, a, n in conflictos
        ] + [
            f"Para escribirlo igual: driver.py intercambio confirmar {mensaje['id']} -- "
            f"para dejar el Excel como está: driver.py intercambio descartar {mensaje['id']}."
        ])
    if not cambios:
        return dict(base, accion="sin-cambios", cambios=[],
                    detalle=["El Excel ya tenía esos valores."])
    return dict(base, accion="aplicar", cambios=cambios,
                detalle=[f"{c}: {_legible(c, a)} → {_legible(c, n)}" for c, a, n in cambios])


def _decidir(mensaje, tag, filas_por_tag, actuales_por_tag, estado) -> dict:
    costos = valores_mensaje(mensaje)
    base = {"mensaje": mensaje, "tag": tag, "costos": costos}
    if tag not in filas_por_tag:
        if mensaje["proyecto"].get("crear") and mensaje.get("tipo", TIPO) == TIPO:
            nombre = str(mensaje["proyecto"]["nombre"]).strip()
            return dict(base, accion="crear", nombre=nombre,
                        cambios=[(c, None, v) for c, v in costos.items()],
                        detalle=[f"Proyecto nuevo {tag} «{nombre}»."]
                        + [f"{c}: {_pesos(v)}" for c, v in costos.items()])
        return dict(base, accion="pendiente", detalle=[
            f"El TAG {tag} no existe en la hoja Proyectos. Se aplicará cuando exista "
            "(Centro de Costos lo crea con la primera factura del proyecto); si el TAG "
            "está mal, descarta este envío y vuelve a enviarlo desde el Formulador."
        ])

    actuales = actuales_por_tag.get(tag, {})
    propios = estado["valores"].get(tag, {})
    reemplaza = reemplaza_mensaje(mensaje)
    confirmado = mensaje["id"] in estado.get("confirmados", [])
    cambios, conflictos = [], []
    for categoria, nuevo in costos.items():
        actual = _numero(actuales.get(categoria))
        if actual is None:
            cambios.append((categoria, None, nuevo))
        elif _iguales(actual, nuevo):
            continue
        elif _iguales(actual, (propios.get(categoria) or {}).get("valor")):
            cambios.append((categoria, actual, nuevo))      # lo había escrito el Formulador
        elif categoria in reemplaza and (
            _iguales(actual, _numero(reemplaza[categoria])) or actual == reemplaza[categoria]
        ):
            cambios.append((categoria, actual, nuevo))      # quien envió vio este valor
        elif confirmado:
            cambios.append((categoria, actual, nuevo))      # confirmado por consola
        else:
            conflictos.append((categoria, actual, nuevo))

    if conflictos:
        return dict(base, accion="pendiente", cambios=cambios, conflictos=conflictos, detalle=[
            f"{c}: el Excel tiene {_pesos(a)}, escrito a mano; {_quien(mensaje)} envía {_pesos(n)}."
            for c, a, n in conflictos
        ] + [
            f"Para reemplazarlos: driver.py intercambio confirmar {mensaje['id']} -- "
            f"para dejar el Excel como está: driver.py intercambio descartar {mensaje['id']}."
        ])
    if not cambios:
        return dict(base, accion="sin-cambios", cambios=[],
                    detalle=["El Excel ya tenía esos valores."])
    return dict(base, accion="aplicar", cambios=cambios,
                detalle=[f"{c}: {_pesos(a)} → {_pesos(n)}" for c, a, n in cambios])


# ── ESCRITURA EN LA HOJA (en memoria; el guardado lo hace ejecutar()) ────────

def _nota(mensaje: dict) -> Comment:
    origen = mensaje.get("origen") or {}
    quien = str(origen.get("usuario") or "").strip()
    que = ("Monto de venta enviado desde " + ("Sistema QUEMPIN" if _quien(mensaje) == "Sistema QUEMPIN"
                                              else "el Formulador de proyectos")
           if mensaje.get("tipo") == TIPO_VENTA else "Costo proyectado enviado desde el Formulador de proyectos")
    texto = (
        f"{que}\n"
        f"{fuente_legible(mensaje)}\n"
        f"Enviado{' por ' + quien if quien else ''} el {_fecha_corta(origen.get('enviado'))}; "
        f"aplicado el {datetime.now().strftime('%d-%m-%Y')}.\n"
        "Si lo cambias a mano, esta nota se quita en la próxima corrida."
    )
    nota = Comment(texto, AUTOR_NOTA)
    nota.width = 320
    nota.height = 110
    return nota


def sincronizar_procedencia(ws, filas_por_tag: dict, columnas: dict, estado: dict) -> list[str]:
    """Olvida (registro + nota de la celda) los valores que alguien cambió a
    mano después de que el Formulador los escribió: desde ahí son manuales."""
    avisos = []
    for tag in list(estado["valores"]):
        fila = filas_por_tag.get(tag)
        for categoria in list(estado["valores"][tag]):
            registro = estado["valores"][tag][categoria]
            if categoria not in COLUMNAS_INTERCAMBIO:
                continue
            celda = ws.cell(row=fila, column=columnas[COLUMNAS_INTERCAMBIO[categoria]]) if fila else None
            if celda is not None and _iguales(_numero(celda.value), registro.get("valor")):
                continue
            if celda is not None and celda.comment is not None and celda.comment.author == AUTOR_NOTA:
                celda.comment = None
            del estado["valores"][tag][categoria]
            if fila:
                avisos.append(
                    f"Intercambio: {tag} {categoria} se cambió a mano después de venir del "
                    "Formulador; desde ahora se trata como valor manual."
                )
        if not estado["valores"][tag]:
            del estado["valores"][tag]
    return avisos


def aplicar(ws, decisiones: list[dict], filas_por_tag: dict, columnas: dict, siguiente_fila: int,
            formatos: dict | None = None) -> list[dict]:
    """Escribe en la hoja las decisiones 'aplicar'/'crear' y pone la nota de
    procedencia (también en 'sin-cambios': el valor ya coincidía y desde ahora
    viene del Formulador). Un ingreso del tablero no lleva nota: queda como
    valor manual, y si pisa uno del Formulador le quita la suya. 'formatos'
    ({columna: formato numérico}) se pone solo en celdas sin formato propio:
    una celda que la persona ya formateó a mano lo conserva. Devuelve las
    filas creadas, con el formato de leer_filas_proyectos, para sumarlas a
    filas_validas."""
    formatos = formatos or {}
    nuevas = []
    for d in decisiones:
        if d["accion"] == "crear":
            fila = siguiente_fila
            siguiente_fila += 1
            ws.cell(row=fila, column=columnas["TAG proyecto"], value=d["tag"])
            ws.cell(row=fila, column=columnas["Nombre del proyecto"], value=d["nombre"])
            filas_por_tag[d["tag"]] = fila
            nuevas.append({"fila": fila, "tag": d["tag"], "nombre": d["nombre"]})
        if d["accion"] not in ("crear", "aplicar", "sin-cambios"):
            continue
        fila = filas_por_tag[d["tag"]]
        del_tablero = d["mensaje"].get("tipo") == TIPO_DATOS
        for clave, _antes, despues in d["cambios"]:
            columna = COLUMNAS_ESCRIBIBLES[clave]
            celda = ws.cell(row=fila, column=columnas[columna])
            celda.value = _valor_celda(clave, despues)
            if despues is not None and columna in formatos and celda.number_format == "General":
                celda.number_format = formatos[columna]
            if del_tablero and celda.comment is not None and celda.comment.author == AUTOR_NOTA:
                celda.comment = None
        if del_tablero:
            continue
        for categoria in d["costos"]:
            ws.cell(row=fila, column=columnas[COLUMNAS_INTERCAMBIO[categoria]]).comment = _nota(d["mensaje"])
        completar_requerimiento(ws, fila, columnas, d["mensaje"])
    return nuevas


def completar_requerimiento(ws, fila: int, columnas: dict, mensaje: dict) -> bool:
    """Escribe el N° de requerimiento del envío si la fila no tiene uno:
    nunca pisa uno escrito (a mano o por otro envío)."""
    req = str((mensaje.get("proyecto") or {}).get("req") or "").strip()
    if not req.isdigit() or COLUMNA_REQ not in columnas:
        return False
    celda = ws.cell(row=fila, column=columnas[COLUMNA_REQ])
    if celda.value not in (None, ""):
        return False
    celda.value = int(req)
    return True


# ── CIERRE (solo después de guardar el Excel) ────────────────────────────────

def finalizar(raiz_intercambio: Path, decisiones: list[dict], estado: dict, ruta_estado: Path) -> list[str]:
    """Registra la procedencia de lo aplicado, guarda el registro y mueve los
    mensajes atendidos a procesado/. Los pendientes quedan en el buzón."""
    avisos = []
    ahora = intercambio.ahora_iso()
    for d in decisiones:
        if d["accion"] not in ("aplicar", "crear", "sin-cambios"):
            continue
        if d["mensaje"].get("tipo") == TIPO_DATOS:
            # Ingreso manual: no entra al registro del Formulador, y lo que
            # pisó de él deja de ser suyo (desde ahora es un valor manual).
            propios = estado["valores"].get(d["tag"]) or {}
            for clave, _antes, _despues in d.get("cambios", []):
                propios.pop(clave, None)
            if d["tag"] in estado["valores"] and not propios:
                del estado["valores"][d["tag"]]
            if d["mensaje"]["id"] in estado["confirmados"]:
                estado["confirmados"].remove(d["mensaje"]["id"])
            continue
        origen = d["mensaje"].get("origen") or {}
        for categoria, valor in d["costos"].items():
            estado["valores"].setdefault(d["tag"], {})[categoria] = {
                "valor": valor, "mensaje": d["mensaje"]["id"], "fuente": fuente_legible(d["mensaje"]),
                "enviado_por": origen.get("usuario") or "", "aplicado": ahora,
                # uid del presupuesto: así se reconoce qué presupuesto quedó cargado en qué TAG
                "uid": (d["mensaje"].get("fuente") or {}).get("uid") or "",
            }
        if d["mensaje"]["id"] in estado["confirmados"]:
            estado["confirmados"].remove(d["mensaje"]["id"])
    guardar_estado(estado, ruta_estado)

    for d in decisiones:
        final = ESTADO_FINAL_POR_ACCION.get(d["accion"])
        if not final:
            continue
        try:
            intercambio.archivar(raiz_intercambio, d["mensaje"], final, d["detalle"], extra={
                "tag": d["tag"],
                "cambios": {c: {"antes": a, "despues": n} for c, a, n in d.get("cambios", [])},
            })
        except OSError as exc:
            avisos.append(f"Intercambio: no se pudo archivar {d['mensaje']['_archivo'].name} ({exc}); "
                          "se vuelve a revisar en la próxima corrida.")
    return avisos


# ── PUBLICACIÓN PARA LAS DEMÁS HERRAMIENTAS ──────────────────────────────────

def _json_simple(valor):
    if isinstance(valor, bool) or valor is None or isinstance(valor, str):
        return valor
    if isinstance(valor, (int, float)):
        return valor if math.isfinite(valor) else None
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if hasattr(valor, "isoformat"):
        return valor.isoformat()
    return str(valor)


def _origen(propios: dict, clave: str, valor_actual):
    if clave in propios and _iguales(_numero(valor_actual), propios[clave]["valor"]):
        return {k: propios[clave][k] for k in ("mensaje", "fuente", "enviado_por", "aplicado")}
    return None


def _req_publicable(valor):
    numero = _numero(valor)
    if isinstance(numero, float) and numero.is_integer() and numero > 0:
        return str(int(numero))
    if isinstance(numero, str) and numero.strip().isdigit():
        return numero.strip()
    return None


def publicar_catalogo(raiz_intercambio: Path, proyectos_af: list[dict], estado: dict,
                      decisiones: list[dict], sesgo: dict | None = None) -> Path:
    """publicado/analisis-financiero.json. 'proyectos_af': [{tag, nombre,
    cliente, categoria, avance, req, venta, proyectados{cat}, reales{cat}}] ya
    armado por quien llama. Los proyectados llevan su origen si vienen del
    Formulador. De la venta se publica solo si está cargada y de dónde vino,
    nunca el monto (la carpeta la ve todo el que entra a la biblioteca).
    'sesgo': el de analisis_financiero.sesgo_cartera()."""
    proyectos = []
    for p in proyectos_af:
        propios = estado["valores"].get(p["tag"], {})
        venta = _numero(p.get("venta"))
        proyectos.append({
            "tag": p["tag"], "nombre": _json_simple(p.get("nombre")),
            "cliente": _json_simple(p.get("cliente")), "categoria": _json_simple(p.get("categoria")),
            "avance": _json_simple(p.get("avance")),
            "proyectados": {c: _json_simple(_numero(p["proyectados"].get(c))) for c in CATEGORIAS},
            "origen": {c: _origen(propios, c, p["proyectados"].get(c)) for c in CATEGORIAS},
            "reales": {c: _json_simple(_numero(p["reales"].get(c))) for c in CATEGORIAS},
            "req": _req_publicable(p.get("req")),
            "venta": {"cargada": venta not in (None, 0), "origen": _origen(propios, CLAVE_VENTA, p.get("venta"))},
            "cierre": _json_simple(p.get("cierre")),
            "porEjecutar": p.get("porEjecutar"),
        })
    mensajes = {
        id_: {k: v for k, v in r.items() if k in ("estado", "fecha", "detalle", "tag")}
        for id_, r in intercambio.resultados_recientes(raiz_intercambio, destino=DESTINO).items()
    }
    for d in decisiones:
        if d["accion"] == "pendiente":
            mensajes[d["mensaje"]["id"]] = {
                "estado": "pendiente", "fecha": intercambio.ahora_iso(), "detalle": d["detalle"], "tag": d["tag"],
            }
    datos = {
        "moneda": "CLP",
        "categorias": list(CATEGORIAS),
        "proyectos": proyectos,
        "mensajes": mensajes,
    }
    if sesgo is not None:
        datos["sesgo"] = sesgo
    return intercambio.publicar(raiz_intercambio, PUBLICACION, HERRAMIENTA, datos)


# ── CONSOLA ──────────────────────────────────────────────────────────────────

def resumen_decisiones(decisiones: list[dict]) -> dict[str, list[str]]:
    """Textos para la consola de run/status, agrupados por acción."""
    salida: dict[str, list[str]] = {}
    for d in decisiones:
        clave = {"crear": "aplicar"}.get(d["accion"], d["accion"])
        cabeza = f"{d['tag'] or '(sin TAG)'} <- {fuente_legible(d['mensaje'])} [{d['mensaje']['id']}]"
        salida.setdefault(clave, []).append(cabeza + "".join(f"\n      {linea}" for linea in d["detalle"]))
    return salida


def descartar(raiz_intercambio: Path, id_mensaje: str) -> bool:
    """Archiva un mensaje del buzón como 'descartado' sin tocar el Excel."""
    mensajes, _ = leer_buzon(raiz_intercambio)
    for m in mensajes:
        if m["id"] == id_mensaje:
            intercambio.archivar(raiz_intercambio, m, "descartado",
                                 ["Descartado por el usuario: el Excel queda como estaba."],
                                 extra={"tag": normalizar_tag((m.get("proyecto") or {}).get("tag"))})
            return True
    return False


def cargar_archivo(raiz_intercambio: Path, ruta: Path) -> tuple[bool, str]:
    """Deja en el buzón un envío que llegó como archivo (lo descargado en la
    pestaña «Ingresar datos» desde un navegador sin acceso a la carpeta, o
    mandado por correo). El mismo mensaje, con su id: se valida contra el
    catálogo y las reglas de este módulo, y nunca entra dos veces (ni si ya
    está en el buzón ni si ya se atendió). (ok, texto para la consola)."""
    import esquemas  # Sistema Intercambio, ya en sys.path
    try:
        with open(ruta, encoding="utf-8-sig") as f:
            mensaje = json.load(f)
    except (OSError, ValueError) as exc:
        return False, f"{Path(ruta).name}: no se pudo leer ({exc})"
    if not isinstance(mensaje, dict) or mensaje.get("destino") != DESTINO or mensaje.get("tipo") not in TIPOS:
        return False, f"{Path(ruta).name}: no es un envío para el Análisis Financiero"
    errores = esquemas.validar_mensaje(mensaje) + validar_contenido(mensaje)
    if errores:
        return False, f"{Path(ruta).name}: " + "; ".join(errores)
    en_buzon, _ = leer_buzon(raiz_intercambio)
    if any(m["id"] == mensaje["id"] for m in en_buzon):
        return False, f"{Path(ruta).name}: ya está en el buzón (id {mensaje['id']})"
    atendido = intercambio.resultados_recientes(raiz_intercambio, dias=3650, destino=DESTINO).get(mensaje["id"])
    if atendido:
        return False, f"{Path(ruta).name}: ya se atendió ({atendido.get('estado')}, id {mensaje['id']})"
    intercambio.asegurar_carpeta(raiz_intercambio)
    destino = intercambio.enviar(raiz_intercambio, mensaje)
    tag = normalizar_tag((mensaje.get("proyecto") or {}).get("tag"))
    return True, f"{Path(ruta).name}: {tag} <- {fuente_legible(mensaje)} queda en el buzón como {destino.name}"


def confirmar(id_mensaje: str, raiz_intercambio: Path, ruta_estado: Path = RUTA_ESTADO) -> bool:
    """Autoriza a la próxima corrida a reemplazar los valores manuales que
    ese mensaje tenía en conflicto. Devuelve False si no está en el buzón."""
    mensajes, _ = leer_buzon(raiz_intercambio)
    if not any(m["id"] == id_mensaje for m in mensajes):
        return False
    estado = leer_estado(ruta_estado)
    if id_mensaje not in estado["confirmados"]:
        estado["confirmados"].append(id_mensaje)
    guardar_estado(estado, ruta_estado)
    return True


# ── PRESUPUESTOS ADJUDICADOS POR CARGAR (con Claude) ─────────────────────────
# Pedido del usuario (2026-09-30): cuando un presupuesto del repositorio del
# Formulador queda «Adjudicada», Claude propone cargar sus costos como
# proyectados sin que nadie tenga que apretar «Enviar costos» en el
# Formulador. Claude pregunta el TAG y deja en el buzón el mismo mensaje
# «presupuesto-proyecto» que dejaría el Formulador: se aplica en el próximo
# run, con todas las reglas de arriba (nunca pisa un valor escrito a mano que
# no se vio).

import formulaciones  # noqa: E402  (Sistema Intercambio, ya en sys.path)

_COMUNES = set((
    "mantencion mantenimiento mantenciones preventiva preventivo correctiva correctivo servicio servicios "
    "sistema sistemas instalacion instalaciones equipo equipos adquisicion reparacion suministro anual "
    "para del las los por con una uno sin que proyecto"
).split())


def _palabras(texto) -> set:
    plano = "".join(c for c in unicodedata.normalize("NFD", str(texto or "").lower()) if unicodedata.category(c) != "Mn")
    return {w for w in re.split(r"[^a-z0-9]+", plano) if len(w) >= 3 and not w.isdigit() and w not in _COMUNES}


def sugerir_tag(item: dict, proyectos_af: list[dict]) -> str:
    """TAG del Análisis Financiero más parecido (título y cliente), si hay uno
    claro. La misma idea que usa el Formulador al enviar costos."""
    d = item["datos"]
    mias = _palabras(f"{d.get('titulo')} {d.get('cliente')}")
    if not mias:
        return ""
    puntajes = sorted(
        ((len(_palabras(f"{p.get('nombre')} {p.get('cliente')}") & mias), p["tag"]) for p in proyectos_af),
        reverse=True,
    )
    if puntajes and puntajes[0][0] > 0 and (len(puntajes) == 1 or puntajes[0][0] > puntajes[1][0]):
        return puntajes[0][1]
    return ""


def _costos_af(item: dict) -> dict | None:
    costos = (item.get("resumen") or {}).get("costosAF")
    if not isinstance(costos, dict) or set(costos) != set(CATEGORIAS):
        return None
    try:
        return {c: int(round(float(costos[c]))) for c in CATEGORIAS}
    except (TypeError, ValueError):
        return None


def _tag_de(item: dict, estado: dict) -> str:
    vinculo = ((item["datos"].get("vinculos") or {}).get("analisisFinanciero") or {})
    if vinculo.get("tag"):
        return normalizar_tag(vinculo["tag"])
    for tag, categorias in estado.get("valores", {}).items():
        if any((v or {}).get("uid") == item["uid"] for v in categorias.values()):
            return tag
    return ""


def adjudicados_para_af(raiz_intercambio: Path, ruta_estado: Path = RUTA_ESTADO) -> list[dict]:
    """Presupuestos adjudicados del repositorio y en qué quedaron respecto del
    Análisis Financiero. Cada uno: {item, costos, tag, sugerido, accion,
    detalle, proyectados}. Acciones: 'cargado' (AF ya tiene esos costos),
    'en-buzon' (enviado, se aplica en el próximo run), 'cargar' (hay TAG y los
    costos difieren), 'elegir-tag' (falta decidir el proyecto de AF),
    'sin-resumen' (lo guardó una versión del Formulador sin resumen: abrirlo y
    guardarlo de nuevo)."""
    estado = leer_estado(ruta_estado)
    sobre = intercambio.leer_publicacion(raiz_intercambio, PUBLICACION) or {}
    proyectos_af = ((sobre.get("datos") if isinstance(sobre, dict) else None) or {}).get("proyectos") or []
    por_tag = {p["tag"]: p for p in proyectos_af if isinstance(p, dict) and p.get("tag")}
    mensajes, _ = leer_buzon(raiz_intercambio)
    en_buzon = {(m.get("fuente") or {}).get("uid") for m in mensajes if m["tipo"] == TIPO}
    salida = []
    for item in formulaciones.leer_repositorio(raiz_intercambio):
        d = item["datos"]
        if d.get("estado") != "Adjudicada" or d.get("eliminado"):
            continue
        costos = _costos_af(item)
        tag = _tag_de(item, estado)
        fila = {"item": item, "costos": costos, "tag": tag, "sugerido": "", "proyectados": None, "detalle": ""}
        if costos is None:
            fila.update(accion="sin-resumen", detalle="lo guardó una versión anterior del Formulador: ábrelo y guárdalo de nuevo")
        elif item["uid"] in en_buzon:
            fila.update(accion="en-buzon", detalle="ya enviado: se aplica en el próximo run")
        elif tag and tag in por_tag:
            hoy = por_tag[tag].get("proyectados") or {}
            fila["proyectados"] = {c: hoy.get(c) for c in CATEGORIAS}
            if all(_iguales(_numero(hoy.get(c)), costos[c]) for c in CATEGORIAS):
                fila.update(accion="cargado", detalle=f"el Análisis Financiero ya tiene estos costos en {tag}")
            else:
                fila.update(accion="cargar", detalle=f"los costos de {tag} en el Análisis Financiero son otros")
        elif tag:
            fila.update(accion="cargar", detalle=f"{tag} todavía no está en la lista del Análisis Financiero")
        else:
            fila.update(accion="elegir-tag", sugerido=sugerir_tag(item, proyectos_af),
                        detalle="falta decidir a qué proyecto del Análisis Financiero corresponde")
        salida.append(fila)
    return salida


def mensaje_desde_formulacion(item: dict, tag: str, nombre: str = "", visto: dict | None = None,
                              usuario: str = "") -> dict:
    """El mismo «presupuesto-proyecto» que arma el Formulador (js/intercambio.js).
    'visto': los proyectados que el usuario vio en el Análisis Financiero al
    decidir (If-Match); con nombre y un TAG que no existe, se crea el proyecto."""
    d, r = item["datos"], item.get("resumen") or {}
    costos = _costos_af(item)
    if costos is None:
        raise ValueError("el presupuesto no trae sus costos por categoría (ábrelo y guárdalo en el Formulador)")
    tag = normalizar_tag(tag)
    mensaje = {
        "esquema": intercambio.ESQUEMA, "id": intercambio.nuevo_id(), "tipo": TIPO, "destino": DESTINO,
        "origen": {"herramienta": "formulador", "usuario": usuario, "enviado": intercambio.ahora_iso()},
        "proyecto": {"tag": tag, "nombre": nombre.strip(), "crear": True} if nombre.strip() else {"tag": tag},
        "fuente": {
            "uid": d.get("uid"), "codigo": d.get("codigo") or "", "version": d.get("version"),
            "titulo": d.get("titulo") or "", "cliente": d.get("cliente") or "", "estado": d.get("estado") or "",
            "modificado": d.get("modificado") or "",
        },
        "costos": costos,
        "informativo": {
            "moneda": "CLP", "costoDirecto": r.get("costoDirecto"), "gastosGenerales": r.get("gastosGenerales"),
            "imprevistos": r.get("imprevistos"), "costoTotal": r.get("costoTotal"), "precioNeto": r.get("precioNeto"),
        },
    }
    if visto is not None:
        # Tal cual vienen en la lista publicada, igual que los manda el Formulador.
        mensaje["reemplaza"] = {c: visto.get(c) for c in CATEGORIAS}
    errores = intercambio.validar_mensaje(mensaje) + validar_presupuesto(mensaje)
    if errores:
        raise ValueError("; ".join(errores))
    return mensaje
