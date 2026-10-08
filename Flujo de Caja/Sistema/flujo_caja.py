# -*- coding: utf-8 -*-
"""
flujo_caja.py -- Flujo de Caja de QUEMPIN, armado solo con lo que ya producen
las demás herramientas (plan de integración 2026-09-30, fase 4).

No hay planilla de Flujo de Caja que mantener a mano: el módulo es un
**consumidor puro** (como pide el CLAUDE.md raíz de Finanzas: "consumir datos
que ya producen módulos anteriores"). Cada movimiento dice de dónde salió:

| Línea | Clase | Fuente |
|---|---|---|
| Egresos pagados | real | Centro de Costos: documentos «Pagado» (foto del tablero) |
| Facturas por pagar | comprometido | Centro de Costos: documentos «Pendiente» |
| Órdenes de compra sin factura | comprometido | Sistema QUEMPIN: OC (61) que todavía no calzan con una factura |
| Cuotas por cobrar | comprometido | Sistema QUEMPIN: cuotas de las cotizaciones de proyectos adjudicados |
| Ventas según Análisis Financiero | estimado | Excel del Análisis Financiero: «Monto de Venta (sin IVA)» + IVA, al cierre o en estados de pago mensuales |
| Ofertas por adjudicar | probable | Planilla de Ingreso: «Ofertado» × tasa histórica de adjudicación |
| Costo por ejecutar | estimado | Análisis Financiero: Proyectado × (1 − avance), menos lo ya comprometido |

Lo que NO existe en ninguna herramienta y por eso no está: los **cobros
reales** (no hay contabilidad ni banco conectados) y el **saldo de caja**
(se fija a mano en ``parametros_flujo_caja.json``; sin él, el acumulado
parte en 0 y se lee como variación, no como saldo).

Todo lo que es un supuesto (a cuántos días se paga una factura, cuándo se
cobra una cuota «contra entrega», cuándo una OC vieja se da por facturada)
vive en ``SUPUESTOS`` y se publica junto con el resultado: el número nunca
viaja sin la regla que lo produjo.

No se publica en la carpeta de intercambio: es la caja de la empresa, y esa
carpeta la ve todo el que entra a la biblioteca de Formulación. Sale a su
Excel (``Excel/Flujo de Caja.xlsx``) y a su tablero (con contraseña).
"""

import json
import math
import os
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

RAIZ = Path(__file__).resolve().parent                 # Flujo de Caja/Sistema
RAIZ_MODULO = RAIZ.parent                               # Flujo de Caja
RAIZ_FINANZAS = RAIZ_MODULO.parent
RAIZ_SISTEMA_INTERCAMBIO = RAIZ_FINANZAS / "Sistema Intercambio"
if str(RAIZ_SISTEMA_INTERCAMBIO) not in sys.path:
    sys.path.insert(0, str(RAIZ_SISTEMA_INTERCAMBIO))

import intercambio  # noqa: E402
import ubicacion  # noqa: E402

RUTA_FOTO_CC = RAIZ_FINANZAS / "Centro de Costos" / "Visualizador Web" / "data" / "centro-de-costos.json"
RUTA_EXCEL = RAIZ_MODULO / "Excel" / "Flujo de Caja.xlsx"
# La venta de cada proyecto no viaja por la carpeta de intercambio (la ve toda
# la biblioteca de Formulación): se lee del Excel del Análisis Financiero, que
# vive aquí al lado, como la foto del Centro de Costos.
RUTA_EXCEL_AF = RAIZ_FINANZAS / "Análisis Financiero" / "Análisis de Proyectos 2026.xlsx"
RUTA_PARAMETROS = RAIZ / "parametros_flujo_caja.json"   # saldo inicial y ajustes (gitignoreado)

SUPUESTOS = {
    "saldoInicial": 0,
    "saldoInicialFecha": None,
    "diasPagoProveedores": 30,        # factura pendiente u OC sin factura: se paga a N días de su fecha
    "diasCobroAdelanto": 15,          # cuota «por adelantado»: a N días de la cotización
    "diasEntregaPorDefecto": 60,      # cuota «contra entrega» sin plazo legible en la cotización
    "diasOtraCondicion": 30,          # cualquier otra condición, o una cotización sin cuotas
    "diasOCSinFactura": 120,          # una OC más antigua se da por facturada (y pagada)
    "toleranciaOC": 0.02,             # una factura calza con una OC si su total difiere a lo más 2 %
    "diasCobroDadoPorHecho": 60,      # una cuota vencida hace más de N días se da por cobrada
    "diasAdjudicacionProbable": 60,   # requerimiento ofertado: se cobraría a N días del cierre
    "diasCobroVenta": 30,             # venta del Análisis Financiero: se cobra a N días del cierre (o del fin de cada mes)
    "diasVentaEnUnPago": 62,          # un proyecto más largo que esto se cobra en estados de pago mensuales
    "ivaVentas": 0.19,                # el Monto de Venta del Análisis Financiero viene sin IVA
    "ivaCostos": 0.19,                # el costo por ejecutar viene sin IVA: se agrega a todo menos mano de obra
    "mesesHistoria": 6,
    "mesesProyeccion": 6,
}
CATEGORIAS_SIN_IVA = ("Mano de Obra",)
DESCRIPCION_SUPUESTOS = {
    "saldoInicial": "Saldo de caja al inicio del mes en curso (estimación a mano, no del banco; 0 = el acumulado es solo la variación). Se cambia con «driver.py saldo <monto>»; en el tablero se puede probar otro valor.",
    "diasPagoProveedores": "Días desde la fecha de una factura pendiente o de una OC hasta que se paga.",
    "diasCobroAdelanto": "Días desde la cotización hasta que se cobra la cuota «por adelantado».",
    "diasEntregaPorDefecto": "Días hasta la cuota «contra entrega» cuando la cotización no trae un plazo legible.",
    "diasOtraCondicion": "Días hasta el cobro de cualquier otra condición o de una cotización sin cuotas.",
    "diasOCSinFactura": "Una OC más antigua que esto se da por facturada y pagada.",
    "toleranciaOC": "Diferencia máxima entre el total de una OC y el de la factura que la cierra.",
    "diasCobroDadoPorHecho": "Una cuota que debió cobrarse hace más de esto se da por cobrada.",
    "diasAdjudicacionProbable": "Días desde el cierre de una oferta hasta su primer cobro, si se gana.",
    "diasCobroVenta": "Venta de un proyecto del Análisis Financiero: días desde su cierre (o desde el fin de cada mes, si va en estados de pago) hasta que se cobra.",
    "diasVentaEnUnPago": "Un proyecto que dura más que esto (de inicio a cierre) se cobra en estados de pago mensuales proporcionales a sus días en cada mes; uno más corto, de una vez.",
    "ivaVentas": "IVA agregado al Monto de Venta (sin IVA) del Análisis Financiero.",
    "ivaCostos": "IVA agregado al costo por ejecutar de materiales, equipos y otros.",
    "mesesHistoria": "Meses hacia atrás con egresos reales.",
    "mesesProyeccion": "Meses hacia adelante proyectados.",
}


# ── UTILIDADES ───────────────────────────────────────────────────────────────

def _fecha(valor) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    try:
        return date.fromisoformat(str(valor or "")[:10])
    except ValueError:
        return None


def _mes(d: date) -> str:
    return f"{d.year:04d}-{d.month:02d}"


def _sumar_meses(d: date, n: int) -> date:
    total = d.year * 12 + d.month - 1 + n
    return date(total // 12, total % 12 + 1, 1)


def _numero(valor) -> float | None:
    if isinstance(valor, bool) or valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor) if math.isfinite(valor) else None
    try:
        return float(str(valor).replace(".", "").replace(",", "."))
    except ValueError:
        return None


_PLAZO = re.compile(r"(\d+)\s*(d[ií]as?|semanas?|mes(?:es)?)", re.IGNORECASE)


def _miles(n) -> str:
    """12345678 -> «12.345.678» (solo el número: el texto que lo rodea no se toca)."""
    return f"{int(n):,}".replace(",", ".")


def dias_de_plazo(texto) -> int | None:
    """'90 días' -> 90; '6 semanas' -> 42; '2 meses' -> 60; ilegible -> None."""
    m = _PLAZO.search(str(texto or ""))
    if not m:
        return None
    n, unidad = int(m.group(1)), m.group(2).lower()
    return n * 7 if unidad.startswith("semana") else n * 30 if unidad.startswith("mes") else n


def _tag_de_ref(ref) -> str | None:
    texto = str(ref or "").strip().upper()
    return texto.split("-")[0] if "-" in texto else None


def _movimiento(fecha: date, sentido: str, clase: str, monto: float, concepto: str, fuente: str,
                referencia: str = "", proyecto: str | None = None, vencido: bool = False) -> dict:
    return {"fecha": fecha.isoformat(), "mes": _mes(fecha), "sentido": sentido, "clase": clase,
            "monto": int(round(monto)), "concepto": concepto, "fuente": fuente, "referencia": referencia,
            "proyecto": proyecto or "", "vencido": vencido}


def _al_presente(fecha: date, hoy: date) -> tuple[date, bool]:
    """Lo que debió ocurrir antes de hoy y sigue pendiente se mueve a hoy (vencido)."""
    return (hoy, True) if fecha < hoy else (fecha, False)


# ── LÍNEAS ───────────────────────────────────────────────────────────────────

def egresos_centro_costos(documentos: list[dict], hoy: date, sup: dict) -> tuple[list[dict], list[str]]:
    """Pagados = egreso real en su fecha; pendientes = comprometido a
    diasPagoProveedores (si eso ya pasó, vencido y al mes en curso). Las notas
    de crédito restan (monto negativo)."""
    movs, avisos = [], []
    sin_fecha = 0
    for d in documentos:
        f = _fecha(d.get("fecha"))
        monto = _numero(d.get("total_con_iva"))
        if f is None or monto is None:
            sin_fecha += 1
            continue
        tag = _tag_de_ref(d.get("ref"))
        concepto = f"{d.get('tipo_documento') or 'Documento'} {d.get('n_documento') or ''} · {d.get('proveedor_tag') or ''}".strip(" ·")
        if str(d.get("estado") or "").lower() == "pendiente":
            fp, vencido = _al_presente(f + timedelta(days=sup["diasPagoProveedores"]), hoy)
            movs.append(_movimiento(fp, "egreso", "comprometido", monto, concepto, "Centro de Costos",
                                    d.get("ref") or "", tag, vencido))
        else:
            movs.append(_movimiento(f, "egreso", "real", monto, concepto, "Centro de Costos", d.get("ref") or "", tag))
    if sin_fecha:
        avisos.append(f"{sin_fecha} documento(s) de Centro de Costos sin fecha o sin total: quedan fuera.")
    return movs, avisos


def egresos_ordenes_de_compra(docs: list[dict], documentos_cc: list[dict], hoy: date, sup: dict) -> tuple[list[dict], list[str]]:
    """OC emitidas en Sistema QUEMPIN que todavía no calzan con una factura
    del Centro de Costos (mismo total ± toleranciaOC, fechada desde la OC en
    adelante y dentro de diasOCSinFactura). Una OC más antigua se da por
    facturada. Sin RUT en Centro de Costos, el cruce es por monto y fecha."""
    movs, avisos = [], []
    facturas = [(_fecha(d.get("fecha")), _numero(d.get("total_con_iva"))) for d in documentos_cc]
    facturas = [(f, t) for f, t in facturas if f and t]
    usadas = set()
    viejas = otras_monedas = 0
    for oc in sorted((x for x in docs if x.get("tipo") == "61"), key=lambda x: str(x.get("fecha"))):
        f, total = _fecha(oc.get("fecha")), _numero(oc.get("total"))
        if f is None or not total:
            continue
        if oc.get("moneda") != "CLP":
            otras_monedas += 1
            continue
        if (hoy - f).days > sup["diasOCSinFactura"]:
            viejas += 1
            continue
        calza = next((i for i, (ff, tt) in enumerate(facturas) if i not in usadas and f <= ff <= f + timedelta(days=sup["diasOCSinFactura"])
                      and abs(tt - total) <= sup["toleranciaOC"] * total), None)
        if calza is not None:
            usadas.add(calza)
            continue
        fp, vencido = _al_presente(f + timedelta(days=sup["diasPagoProveedores"]), hoy)
        proyecto = (oc.get("proyecto") or {}).get("tag") or (oc.get("proyecto") or {}).get("req")
        movs.append(_movimiento(fp, "egreso", "comprometido", total,
                                f"OC {oc.get('folio')} · {(oc.get('contraparte') or {}).get('razon_social') or ''}".strip(" ·"),
                                "Sistema QUEMPIN", oc.get("folio") or "", proyecto, vencido))
    if viejas:
        avisos.append(f"{viejas} OC con más de {sup['diasOCSinFactura']} días se dan por facturadas.")
    if otras_monedas:
        avisos.append(f"{otras_monedas} OC en otra moneda quedan fuera (no hay tipo de cambio).")
    return movs, avisos


def _clave_proyecto(doc: dict) -> str | None:
    p = doc.get("proyecto") or {}
    return (f"req:{p['req']}" if p.get("req") else f"tag:{p['tag']}" if p.get("tag") else None)


def ingresos_cotizaciones(docs: list[dict], requerimientos: list[dict], proyectos_af: list[dict],
                          hoy: date, sup: dict) -> tuple[list[dict], list[str], set]:
    """Cuotas de la última cotización de cada proyecto adjudicado (su N° está
    «Adjudicado» en la Planilla, o su TAG existe en Análisis Financiero).
    Devuelve también los N° de requerimiento ya cubiertos, para no contarlos
    además como probables."""
    adjudicados_req = {str(r["numero"]) for r in requerimientos if str(r.get("estado") or "").lower() == "adjudicado"}
    tags_af = {str(p.get("tag")) for p in proyectos_af if p.get("tag")}
    ultimas: dict[str, dict] = {}
    for d in docs:
        if d.get("tipo") != "60":
            continue
        clave = _clave_proyecto(d)
        p = d.get("proyecto") or {}
        if not clave or not (str(p.get("req") or "") in adjudicados_req or str(p.get("tag") or "") in tags_af):
            continue
        if clave not in ultimas or str(d.get("fecha")) > str(ultimas[clave].get("fecha")):
            ultimas[clave] = d
    movs, avisos = [], []
    dadas = otras = 0
    for d in ultimas.values():
        f = _fecha(d.get("fecha"))
        if f is None:
            continue
        if d.get("moneda") != "CLP":
            otras += 1
            continue
        proyecto = (d.get("proyecto") or {}).get("tag") or (d.get("proyecto") or {}).get("req")
        cliente = (d.get("contraparte") or {}).get("razon_social") or ""
        cuotas = [c for c in d.get("cuotas") or [] if _numero(c.get("monto"))]
        if not cuotas:
            cuotas = [{"condicion": "sin cuotas", "monto": d.get("total")}]
        plazo = dias_de_plazo(d.get("plazoEntrega"))
        for c in cuotas:
            condicion = str(c.get("condicion") or "").lower()
            if "adelant" in condicion or "advance" in condicion:
                dias = sup["diasCobroAdelanto"]
            elif "entrega" in condicion or "delivery" in condicion:
                dias = plazo if plazo is not None else sup["diasEntregaPorDefecto"]
            else:
                dias = sup["diasOtraCondicion"]
            esperada = f + timedelta(days=dias)
            concepto = f"Cotización {d.get('folio')} · {c.get('condicion') or ''} · {cliente}".strip(" ·")
            if (hoy - esperada).days > sup["diasCobroDadoPorHecho"]:
                # Se da por cobrada y queda en su mes, en la historia (2026-10-06).
                dadas += 1
                cobro, vencido, concepto = esperada, False, concepto + " (cobro dado por hecho)"
            else:
                cobro, vencido = _al_presente(esperada, hoy)
            movs.append(_movimiento(cobro, "ingreso", "comprometido", _numero(c["monto"]), concepto,
                                    "Sistema QUEMPIN", d.get("folio") or "", proyecto, vencido))
    sin_proyecto = sum(1 for d in docs if d.get("tipo") == "60" and not _clave_proyecto(d))
    if sin_proyecto:
        avisos.append(f"{sin_proyecto} cotización(es) de Sistema QUEMPIN sin proyecto (N° de requerimiento o TAG): "
                      "sus cuotas no se cuentan hasta asociarlas en el campo «Proyecto».")
    if dadas:
        avisos.append(f"{dadas} cuota(s) que debieron cobrarse hace más de {sup['diasCobroDadoPorHecho']} días se dan por cobradas "
                      "y quedan en su mes, en la historia.")
    if otras:
        avisos.append(f"{otras} cotización(es) adjudicada(s) en otra moneda quedan fuera (no hay tipo de cambio).")
    # Lo ya cubierto por una cotización: sus N° de requerimiento (no se cuentan
    # además como probables) y sus TAG (no se cuentan además con la venta del AF).
    cubiertos = {k[4:] for k in ultimas if k.startswith("req:")}
    cubiertos |= {f"tag:{(d.get('proyecto') or {}).get('tag')}" for d in ultimas.values() if (d.get("proyecto") or {}).get("tag")}
    return movs, avisos, cubiertos


def _estados_de_pago(inicio: date | None, cierre: date, venta: float, sup: dict) -> list[tuple[date, int]]:
    """[(fecha de cobro, monto)]. Un proyecto que dura hasta diasVentaEnUnPago
    se cobra de una vez a diasCobroVenta del cierre; uno más largo (contratos
    de mantención de un año, 2026-10-06), en estados de pago mensuales
    proporcionales a sus días en cada mes, cada uno a diasCobroVenta del fin
    de ese mes (el último, del cierre)."""
    total = int(round(venta))
    plazo = timedelta(days=sup["diasCobroVenta"])
    if inicio is None or inicio >= cierre or (cierre - inicio).days <= sup["diasVentaEnUnPago"]:
        return [(cierre + plazo, total)]
    tramos, desde = [], inicio
    while desde <= cierre:
        hasta = min(_sumar_meses(desde, 1) - timedelta(days=1), cierre)
        tramos.append((hasta, (hasta - desde).days + 1))
        desde = hasta + timedelta(days=1)
    dias = sum(d for _, d in tramos)
    partes = [total * d // dias for _, d in tramos]
    partes[-1] += total - sum(partes)          # que sumen la venta exacta
    return [(hasta + plazo, parte) for (hasta, _), parte in zip(tramos, partes)]


def ingresos_ventas_af(proyectos: list[dict], cubiertos: set, ultimo_gasto: dict, hoy: date, sup: dict) -> tuple[list[dict], list[str]]:
    """Las ventas de los proyectos del Análisis Financiero (2026-10-06: el
    usuario pidió sacarlas de ahí, también las de los proyectos terminados).
    «Monto de Venta (sin IVA)» + IVA, en los estados de pago de
    _estados_de_pago. Sin fecha de cierre: la del último gasto del proyecto
    en Centro de Costos si está terminado, o el fin de los 3 meses que usa el
    costo por ejecutar si sigue en curso. Lo que debió cobrarse hace más de
    diasCobroDadoPorHecho se da por cobrado y queda en su mes (la historia);
    lo más reciente que sigue pendiente, vencido en el mes en curso.

    Un proyecto con cotización asociada (su TAG o su N° en `cubiertos`) no
    entra: sus cuotas reales ya están en el flujo."""
    movs, avisos = [], []
    inicio_mes = date(hoy.year, hoy.month, 1)
    entran, en_cuotas, por_gasto, sin_fecha, sin_venta, dadas = [], 0, [], [], [], 0
    for p in proyectos:
        tag = str(p.get("tag") or "").strip()
        if not tag or f"tag:{tag}" in cubiertos or (p.get("req") and str(p["req"]) in cubiertos):
            continue
        venta = _numero(p.get("venta"))
        if not venta or venta <= 0:
            if p.get("categoria") != "Gastos Generales":
                sin_venta.append(tag)
            continue
        inicio, cierre = _fecha(p.get("inicio")), _fecha(p.get("cierre"))
        if cierre is None:
            if (_numero(p.get("avance")) or 0) < 1:
                cierre = _sumar_meses(inicio_mes, 3) - timedelta(days=1)
            elif tag in ultimo_gasto:
                cierre = ultimo_gasto[tag]
                por_gasto.append(tag)
            else:
                sin_fecha.append(tag)
                continue
        pagos = _estados_de_pago(inicio, cierre, venta * (1 + sup["ivaVentas"]), sup)
        entran.append(tag)
        en_cuotas += len(pagos) > 1
        for i, (cobro, monto) in enumerate(pagos, start=1):
            concepto = f"{tag} · {p.get('nombre') or ''} · venta + IVA" + (f" (estado de pago {i}/{len(pagos)})" if len(pagos) > 1 else "")
            if (hoy - cobro).days > sup["diasCobroDadoPorHecho"]:
                dadas += 1
                vencido, concepto = False, concepto + " (cobro dado por hecho)"
            else:
                cobro, vencido = _al_presente(cobro, hoy)
            movs.append(_movimiento(cobro, "ingreso", "estimado", monto, concepto, "Análisis Financiero", tag, tag, vencido))
    if entran:
        avisos.append(f"{len(entran)} proyecto(s) entran con la venta del Análisis Financiero + IVA, cobrada a {sup['diasCobroVenta']} días "
                      f"del cierre ({en_cuotas} de ellos, por durar más de {sup['diasVentaEnUnPago']} días, en estados de pago mensuales). "
                      "Asociar su cotización en Sistema QUEMPIN (campo «Proyecto» = TAG) la reemplaza por sus cuotas reales.")
    if dadas:
        avisos.append(f"{dadas} cobro(s) de ventas que debieron entrar hace más de {sup['diasCobroDadoPorHecho']} días se dan por hechos "
                      "y quedan en su mes, en la historia: el flujo no sabe si de verdad se cobraron (no hay banco conectado).")
    if por_gasto:
        avisos.append(f"Sin fecha de cierre en el Análisis Financiero, se usa la de su último gasto en Centro de Costos: {', '.join(por_gasto)}.")
    if sin_fecha:
        avisos.append(f"Sin fecha de cierre ni gastos, su venta queda fuera: {', '.join(sin_fecha)}.")
    if sin_venta:
        avisos.append(f"Sin «Monto de Venta» en el Análisis Financiero, su venta no está en el flujo: {', '.join(sin_venta)}.")
    return movs, avisos


def _req(valor) -> str | None:
    """N° de requerimiento como texto («279», no «279.0»)."""
    if valor is None or str(valor).strip() == "":
        return None
    n = _numero(valor)
    return str(int(n)) if n is not None and n == int(n) else str(valor).strip()


def leer_proyectos_af(ruta: Path = RUTA_EXCEL_AF) -> tuple[list[dict], list[str]]:
    """Los proyectos de la hoja «Proyectos» del Excel del Análisis Financiero,
    con su venta (solo lectura: ese libro lo escribe su propio módulo)."""
    columnas = {"tag": "TAG proyecto", "nombre": "Nombre del proyecto", "categoria": "Categoría", "avance": "% Avance",
                "inicio": "Fecha de inicio", "cierre": "Fecha de cierre", "venta": "Monto de Venta (sin IVA)"}
    try:
        import openpyxl
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    except Exception as error:      # no existe, abierto con bloqueo o no es un .xlsx
        return [], [f"No se pudo leer el Excel del Análisis Financiero ({ruta.name}: {error}): sin ventas de proyectos."]
    try:
        if "Proyectos" not in libro.sheetnames:
            return [], ["El Excel del Análisis Financiero no tiene la hoja «Proyectos»: sin ventas de proyectos."]
        filas = libro["Proyectos"].iter_rows(values_only=True)
        encabezado = [str(c or "").strip() for c in next(filas, ())]
        if columnas["tag"] not in encabezado or columnas["venta"] not in encabezado:
            return [], ["El Excel del Análisis Financiero no tiene «TAG proyecto» y «Monto de Venta (sin IVA)»: sin ventas de proyectos."]
        indice = {k: encabezado.index(v) for k, v in columnas.items() if v in encabezado}
        i_req = next((i for i, h in enumerate(encabezado) if "Requerimiento" in h), None)
        proyectos = []
        for f in filas:
            def valor(i):
                return f[i] if i is not None and i < len(f) else None
            if not valor(indice["tag"]):
                continue
            proyecto = {k: valor(i) for k, i in indice.items()}
            proyecto["tag"] = str(proyecto["tag"]).strip()
            proyecto["req"] = _req(valor(i_req))
            proyectos.append(proyecto)
        return proyectos, []
    finally:
        libro.close()


def ingresos_probables(requerimientos: list[dict], tasa: float | None, cubiertos: set, hoy: date, sup: dict) -> tuple[list[dict], list[str]]:
    """Requerimientos «Ofertado» con valor ofertado (con IVA) × la tasa
    histórica de adjudicación de lo ofertado, a diasAdjudicacionProbable del
    cierre de la oferta."""
    if not tasa:
        return [], ["Sin tasa de adjudicación (la Planilla no tiene ofertas resueltas): no se proyectan ingresos probables."]
    movs, ofertado = [], {}
    for r in requerimientos:
        if str(r.get("estado") or "").lower() != "ofertado" or str(r.get("numero")) in cubiertos:
            continue
        valor = _numero(r.get("valorOfertado"))
        if not valor:
            continue
        cierre = _fecha(r.get("cierre")) or hoy
        esperada = cierre + timedelta(days=sup["diasAdjudicacionProbable"])
        if esperada < hoy:
            esperada = hoy + timedelta(days=sup["diasOtraCondicion"])
        movs.append(_movimiento(esperada, "ingreso", "probable", valor * tasa,
                                f"N° {r.get('numero')} · {r.get('titulo') or ''} ({round(tasa * 100)} % de {_miles(valor)})",
                                "Planilla de Ingreso", str(r.get("numero")), str(r.get("numero"))))
        ofertado[str(r.get("numero"))] = valor
    # Una oferta se gana o se pierde entera: si una sola domina lo probable,
    # el monto ponderado no es un ingreso que vaya a llegar (2026-10-02: una
    # sola licitación era casi todo lo probable).
    avisos = []
    total = sum(m["monto"] for m in movs)
    mayor = max(movs, key=lambda m: m["monto"]) if movs else None
    if mayor and len(movs) > 1 and mayor["monto"] >= 0.5 * total:
        avisos.append(f"La oferta N° {mayor['referencia']} es el {round(mayor['monto'] / total * 100)} % de lo probable: se gana o se "
                      f"pierde entera (${_miles(ofertado[mayor['referencia']])} o nada), no llega ponderada.")
    return movs, avisos


def egresos_por_ejecutar(proyectos_af: list[dict], comprometido_por_tag: dict, hoy: date, sup: dict) -> tuple[list[dict], list[str]]:
    """Lo que falta ejecutar de cada proyecto en curso (porEjecutar del
    Análisis Financiero, con IVA salvo la mano de obra), menos lo que ya está
    comprometido para ese TAG, repartido en partes iguales hasta su cierre
    (o en 3 meses si no tiene cierre futuro)."""
    movs, avisos = [], []
    inicio = date(hoy.year, hoy.month, 1)
    for p in proyectos_af:
        pe = p.get("porEjecutar")
        if not isinstance(pe, dict):
            continue
        total = sum((_numero(v) or 0) * (1 if cat in CATEGORIAS_SIN_IVA else 1 + sup["ivaCostos"]) for cat, v in pe.items())
        total -= comprometido_por_tag.get(str(p.get("tag")), 0)
        if total <= 0:
            continue
        cierre = _fecha(p.get("cierre"))
        meses = 3
        if cierre and cierre > hoy:
            meses = max(1, (cierre.year - hoy.year) * 12 + cierre.month - hoy.month + 1)
        # Partes enteras; la última se lleva la diferencia para que sumen el total exacto.
        total = int(round(total))
        partes = [total // meses] * meses
        partes[-1] += total - sum(partes)
        for i, parte in enumerate(partes):
            mes = _sumar_meses(inicio, i)
            movs.append(_movimiento(max(mes, hoy) if i == 0 else mes, "egreso", "estimado", parte,
                                    f"{p.get('tag')} · {p.get('nombre') or ''} (por ejecutar, {i + 1}/{meses})",
                                    "Análisis Financiero", str(p.get("tag")), str(p.get("tag"))))
    return movs, avisos


# ── RESUMEN POR MES ──────────────────────────────────────────────────────────

LINEAS = (
    ("ingreso", "comprometido", "Cuotas por cobrar"),
    ("ingreso", "estimado", "Ventas según Análisis Financiero"),
    ("ingreso", "probable", "Ofertas por adjudicar (ponderadas)"),
    ("egreso", "real", "Egresos pagados"),
    ("egreso", "comprometido", "Por pagar (facturas pendientes y OC)"),
    ("egreso", "estimado", "Costo por ejecutar de proyectos en curso"),
)


def resumen_por_mes(movimientos: list[dict], hoy: date, sup: dict) -> list[dict]:
    inicio = date(hoy.year, hoy.month, 1)
    meses = [_mes(_sumar_meses(inicio, i)) for i in range(-sup["mesesHistoria"], sup["mesesProyeccion"])]
    sumas = defaultdict(float)
    for m in movimientos:
        if m["mes"] in meses:
            sumas[(m["mes"], m["sentido"], m["clase"])] += m["monto"]
    actual = _mes(inicio)
    saldo = saldo_sin = float(sup.get("saldoInicial") or 0)
    salida = []
    for mes in meses:
        fila = {"mes": mes, "proyectado": mes >= actual}
        for sentido, clase, _ in LINEAS:
            fila[f"{sentido}_{clase}"] = int(round(sumas[(mes, sentido, clase)]))
        fila["neto"] = int(round(fila["ingreso_comprometido"] + fila["ingreso_estimado"] + fila["ingreso_probable"]
                                 - fila["egreso_comprometido"] - fila["egreso_estimado"] - fila["egreso_real"]))
        fila["netoSinProbables"] = fila["neto"] - fila["ingreso_probable"]
        if fila["proyectado"]:
            saldo += fila["neto"]
            saldo_sin += fila["netoSinProbables"]
            fila["acumulado"] = int(round(saldo))
            fila["acumuladoSinProbables"] = int(round(saldo_sin))
        else:
            fila["acumulado"] = fila["acumuladoSinProbables"] = None
        salida.append(fila)
    return salida


# ── ORQUESTACIÓN ─────────────────────────────────────────────────────────────

def leer_parametros(ruta: Path = RUTA_PARAMETROS) -> dict:
    sup = dict(SUPUESTOS)
    try:
        with open(ruta, encoding="utf-8") as f:
            propios = json.load(f)
        sup.update({k: v for k, v in propios.items() if k in SUPUESTOS})
    except (OSError, ValueError):
        pass
    return sup


def guardar_saldo(monto, fecha: date | None = None, ruta: Path = RUTA_PARAMETROS) -> dict:
    """Fija el saldo de caja inicial (estimación a mano) en
    parametros_flujo_caja.json, conservando lo demás que tenga el archivo.
    La fecha por defecto es el inicio del mes en curso, que es cuando el
    cálculo lo aplica. Acepta «1500000», «1.500.000» o «$1.500.000»."""
    valor = _numero(str(monto).replace("$", "").replace(" ", ""))
    if valor is None:
        raise ValueError(f"No es un monto: {monto!r}")
    try:
        with open(ruta, encoding="utf-8") as f:
            propios = json.load(f)
        if not isinstance(propios, dict):
            propios = {}
    except (OSError, ValueError):
        propios = {}
    propios["saldoInicial"] = int(round(valor))
    propios["saldoInicialFecha"] = (fecha or date.today().replace(day=1)).isoformat()
    ruta.parent.mkdir(parents=True, exist_ok=True)
    temporal = ruta.with_name(ruta.name + ".tmp")
    temporal.write_text(json.dumps(propios, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    os.replace(temporal, ruta)
    return propios


def _datos(raiz, nombre) -> dict:
    sobre = intercambio.leer_publicacion(raiz, nombre) if raiz else None
    return (sobre or {}).get("datos") or {}


def armar(raiz_intercambio: Path | None, ruta_foto_cc: Path = RUTA_FOTO_CC, hoy: date | None = None,
          sup: dict | None = None, ruta_excel_af: Path = RUTA_EXCEL_AF) -> dict:
    """El flujo completo, desde lo que publicaron las demás herramientas."""
    hoy = hoy or date.today()
    sup = sup or leer_parametros()
    avisos = []
    try:
        with open(ruta_foto_cc, encoding="utf-8") as f:
            documentos_cc = json.load(f).get("documentos") or []
    except (OSError, ValueError):
        documentos_cc = []
        avisos.append("No se encontró la foto del Centro de Costos (corre su visualizador): sin egresos reales ni por pagar.")
    if raiz_intercambio is None:
        avisos.append("Sin carpeta de intercambio: sin cotizaciones, OC, requerimientos ni costo por ejecutar.")
    docs = _datos(raiz_intercambio, "documentos-comerciales").get("documentos") or []
    reqs_pub = _datos(raiz_intercambio, "requerimientos")
    reqs = reqs_pub.get("requerimientos") or []
    tasa = (reqs_pub.get("resumen") or {}).get("tasaAdjudicacion")
    proyectos_af = _datos(raiz_intercambio, "analisis-financiero").get("proyectos") or []

    movs = []
    for armador in (lambda: egresos_centro_costos(documentos_cc, hoy, sup),
                    lambda: egresos_ordenes_de_compra(docs, documentos_cc, hoy, sup)):
        m, a = armador()
        movs += m
        avisos += a
    m, a, cubiertos = ingresos_cotizaciones(docs, reqs, proyectos_af, hoy, sup)
    movs += m
    avisos += a
    proyectos_venta, a = leer_proyectos_af(ruta_excel_af)
    avisos += a
    ultimo_gasto = {}
    for d in documentos_cc:
        tag, f = _tag_de_ref(d.get("ref")), _fecha(d.get("fecha"))
        if tag and f and f > ultimo_gasto.get(tag, date.min):
            ultimo_gasto[tag] = f
    m, a = ingresos_ventas_af(proyectos_venta, cubiertos, ultimo_gasto, hoy, sup)
    movs += m
    avisos += a
    # Un requerimiento cuya venta ya está en el Análisis Financiero no se
    # cuenta además como oferta probable.
    con_venta = {x["proyecto"] for x in m}
    cubiertos |= {p["req"] for p in proyectos_venta if p.get("req") and p["tag"] in con_venta}
    m, a = ingresos_probables(reqs, tasa, cubiertos, hoy, sup)
    movs += m
    avisos += a
    comprometido_por_tag = defaultdict(float)
    for x in movs:
        if x["sentido"] == "egreso" and x["clase"] == "comprometido" and x["proyecto"]:
            comprometido_por_tag[x["proyecto"]] += x["monto"]
    m, a = egresos_por_ejecutar(proyectos_af, comprometido_por_tag, hoy, sup)
    movs += m
    avisos += a
    if sup.get("saldoInicial"):
        desde = f" al {sup['saldoInicialFecha']}" if sup.get("saldoInicialFecha") else ""
        avisos.append(f"Saldo de caja inicial: ${_miles(sup['saldoInicial'])}{desde}. Es una estimación a mano "
                      "(parametros_flujo_caja.json), no un dato del banco: en el tablero se puede probar otro valor.")
    else:
        avisos.append("Sin saldo inicial de caja (parametros_flujo_caja.json): el acumulado es la variación desde hoy, no el saldo.")
    movs.sort(key=lambda x: (x["fecha"], x["sentido"], x["clase"]))
    return {
        "generado": datetime.now().astimezone().isoformat(timespec="seconds"),
        "hoy": hoy.isoformat(),
        "moneda": "CLP",
        "meses": resumen_por_mes(movs, hoy, sup),
        "movimientos": movs,
        "lineas": [{"sentido": s, "clase": c, "nombre": n} for s, c, n in LINEAS],
        "supuestos": {k: {"valor": sup[k], "descripcion": DESCRIPCION_SUPUESTOS.get(k, "")} for k in SUPUESTOS if k != "saldoInicialFecha"},
        "avisos": avisos,
    }


# ── EXCEL ────────────────────────────────────────────────────────────────────

def escribir_excel(datos: dict, ruta: Path = RUTA_EXCEL) -> Path:
    """El libro se regenera completo en cada corrida (no tiene columnas a
    mano). Guardado atómico: si está abierto, PermissionError y el anterior
    queda intacto."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = openpyxl.Workbook()
    encabezado = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="54565A")      # Cool Gray 11C del manual de marca
    moneda = '"$"#,##0;[Red]-"$"#,##0'

    def hoja(titulo, columnas, filas, anchos):
        ws = libro.create_sheet(titulo)
        ws.append(columnas)
        for i, c in enumerate(ws[1], start=1):
            c.font, c.fill, c.alignment = encabezado, relleno, Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = anchos[i - 1]
        for fila in filas:
            ws.append(fila)
        ws.freeze_panes = "B2"
        return ws

    meses = datos["meses"]
    # El saldo se calcula con fórmulas desde la celda del saldo inicial en
    # «Supuestos»: quien no está seguro del saldo (2026-10-02) lo cambia ahí
    # y ve el resultado sin correr nada (para que quede, «driver.py saldo»).
    fila_saldo = 2 + list(datos["supuestos"]).index("saldoInicial")
    fila_neto, fila_neto_sin = 2 + len(LINEAS), 3 + len(LINEAS)

    def saldos(fila_propia, fila_neto_usada):
        celdas, anterior = [], None
        for i, m in enumerate(meses):
            col = get_column_letter(i + 2)
            if not m["proyectado"]:
                celdas.append(None)
                continue
            base = f"Supuestos!$B${fila_saldo}" if anterior is None else f"{anterior}{fila_propia}"
            celdas.append(f"={base}+{col}{fila_neto_usada}")
            anterior = col
        return celdas

    ws = hoja("Resumen", ["Línea"] + [m["mes"] + ("" if m["proyectado"] else " (real)") for m in meses],
              [[n] + [m[f"{s}_{c}"] for m in meses] for s, c, n in LINEAS]
              + [["Neto del mes"] + [m["neto"] for m in meses],
                 ["Neto sin ofertas probables"] + [m["netoSinProbables"] for m in meses],
                 ["Saldo proyectado sin ofertas por adjudicar"] + saldos(4 + len(LINEAS), fila_neto_sin),
                 ["Saldo proyectado con ofertas ponderadas"] + saldos(5 + len(LINEAS), fila_neto)],
              [38] + [14] * len(meses))
    for fila in ws.iter_rows(min_row=2, min_col=2):
        for c in fila:
            c.number_format = moneda
    for fila in (4 + len(LINEAS), 5 + len(LINEAS)):
        for c in ws[fila]:
            c.font = Font(bold=True)
    ws = hoja("Movimientos", ["Fecha", "Mes", "Sentido", "Clase", "Monto", "Concepto", "Fuente", "Referencia", "Proyecto", "Vencido"],
              [[m["fecha"], m["mes"], m["sentido"], m["clase"], m["monto"], m["concepto"], m["fuente"], m["referencia"],
                m["proyecto"], "Sí" if m["vencido"] else ""] for m in datos["movimientos"]],
              [12, 9, 9, 13, 14, 60, 18, 14, 10, 8])
    for c in ws["E"][1:]:
        c.number_format = moneda
    ws.auto_filter.ref = ws.dimensions
    ws = hoja("Supuestos", ["Supuesto", "Valor", "Qué significa"],
              [[k, v["valor"], v["descripcion"]] for k, v in datos["supuestos"].items()]
              + [[], ["Avisos de esta corrida"]] + [[a] for a in datos["avisos"]]
              + [[], [f"Generado {datos['generado']}. Se regenera completo en cada corrida: para probar otro "
                      "saldo inicial, cambia B" + str(fila_saldo) + " (el Resumen se recalcula); para que quede, "
                      "«driver.py saldo <monto>»."]],
              [26, 14, 90])
    ws.cell(row=fila_saldo, column=2).number_format = moneda
    ws.cell(row=fila_saldo, column=2).font = Font(bold=True)
    del libro["Sheet"]
    ruta.parent.mkdir(parents=True, exist_ok=True)
    temporal = ruta.with_name(f"{ruta.stem}.tmp{ruta.suffix}")
    libro.save(temporal)
    try:
        os.replace(temporal, ruta)
    except OSError:
        temporal.unlink(missing_ok=True)
        raise
    return ruta


def ejecutar(raiz_intercambio: Path | None = None, ruta_excel: Path = RUTA_EXCEL, escribir: bool = True,
             hoy: date | None = None) -> dict:
    """status (escribir=False) o run. Devuelve los datos y, en run, la ruta del Excel o el error."""
    raiz = raiz_intercambio if raiz_intercambio is not None else ubicacion.ubicar_intercambio(RAIZ)
    if raiz is not None and not intercambio.es_carpeta_de_intercambio(raiz):
        raiz = None
    datos = armar(raiz, hoy=hoy)
    datos["excel"] = None
    if escribir:
        try:
            datos["excel"] = str(escribir_excel(datos, ruta_excel))
        except PermissionError as error:
            datos["avisos"].append(f"No se pudo guardar el Excel (¿abierto?): {error}")
    return datos
