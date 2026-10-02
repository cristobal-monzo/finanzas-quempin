# -*- coding: utf-8 -*-
"""Flujo de Caja: cada línea sale de una herramienta y cada supuesto se
aplica como dice. Datos sintéticos en tmp_path: nunca los reales."""
import json
from datetime import date

import openpyxl
import pytest

import flujo_caja as fc

ic = fc.intercambio
HOY = date(2026, 10, 1)
SUP = dict(fc.SUPUESTOS)


def _cc(ref, fecha, total, estado="Pagado", tipo="Factura"):
    return {"ref": ref, "fecha": fecha, "total_con_iva": total, "estado": estado, "tipo_documento": tipo,
            "n_documento": "1", "proveedor_tag": "Ferre"}


def _cot(folio, fecha, total, cuotas, req=None, tag=None, plazo="", moneda="CLP"):
    return {"tipo": "60", "folio": folio, "pais": "Chile", "fecha": fecha, "moneda": moneda, "total": total,
            "contraparte": {"razon_social": "Cliente SpA"}, "proyecto": {"req": req, "tag": tag},
            "cuotas": cuotas, "plazoEntrega": plazo}


def _oc(folio, fecha, total, tag=None):
    return {"tipo": "61", "folio": folio, "pais": "Chile", "fecha": fecha, "moneda": "CLP", "total": total,
            "contraparte": {"razon_social": "Proveedor SpA"}, "proyecto": {"req": None, "tag": tag}}


def test_plazo_legible():
    assert fc.dias_de_plazo("90 días") == 90
    assert fc.dias_de_plazo("Entrega en 6 semanas") == 42
    assert fc.dias_de_plazo("2 meses") == 60
    assert fc.dias_de_plazo("a convenir") is None


def test_centro_de_costos_pagados_reales_y_pendientes_por_pagar():
    movs, _ = fc.egresos_centro_costos([
        _cc("UMAG-001", "2026-09-10", 119000),
        _cc("UMAG-002", "2026-09-20", 50000, estado="Pendiente"),
        _cc("UMAG-003", "2026-07-01", 20000, estado="Pendiente"),
        _cc("UMAG-004", "2026-09-11", -5000, estado="Nota de Crédito", tipo="Nota de Crédito"),
    ], HOY, SUP)
    por_ref = {m["referencia"]: m for m in movs}
    assert por_ref["UMAG-001"]["clase"] == "real" and por_ref["UMAG-001"]["mes"] == "2026-09"
    assert por_ref["UMAG-002"]["clase"] == "comprometido" and por_ref["UMAG-002"]["fecha"] == "2026-10-20"
    assert por_ref["UMAG-003"]["vencido"] and por_ref["UMAG-003"]["fecha"] == HOY.isoformat()   # vencida: al mes en curso
    assert por_ref["UMAG-004"]["monto"] == -5000 and por_ref["UMAG-001"]["proyecto"] == "UMAG"


def test_orden_de_compra_sin_factura_queda_comprometida_y_con_factura_no():
    docs = [_oc("612610", "2026-09-15", 100000, tag="UMAG"), _oc("612611", "2026-09-20", 300000), _oc("612500", "2026-03-01", 99)]
    cc = [_cc("UMAG-010", "2026-09-25", 101000)]           # calza con la 612610 (1 %)
    movs, avisos = fc.egresos_ordenes_de_compra(docs, cc, HOY, SUP)
    assert [m["referencia"] for m in movs] == ["612611"]
    assert movs[0]["fecha"] == "2026-10-20"
    assert any("se dan por facturadas" in a for a in avisos)


def test_cuotas_de_la_ultima_cotizacion_de_cada_proyecto_adjudicado():
    reqs = [{"numero": 280, "estado": "Adjudicado"}, {"numero": 281, "estado": "Ofertado"}]
    docs = [
        _cot("602690", "2026-09-01", 1000, [{"porcentaje": 100, "condicion": "por adelantado", "monto": 1000}], req="280"),
        _cot("602695", "2026-09-20", 2000, [{"porcentaje": 50, "condicion": "por adelantado", "monto": 1000},
                                            {"porcentaje": 50, "condicion": "contra entrega", "monto": 1000}], req="280", plazo="90 días"),
        _cot("602696", "2026-09-20", 5000, [{"porcentaje": 100, "condicion": "por adelantado", "monto": 5000}], req="281"),
        _cot("602697", "2026-09-25", 700, [], tag="DEMO"),
    ]
    movs, _, cubiertos = fc.ingresos_cotizaciones(docs, reqs, [{"tag": "DEMO"}], HOY, SUP)
    por = sorted((m["referencia"], m["fecha"], m["monto"]) for m in movs)
    assert por == [("602695", "2026-10-05", 1000), ("602695", "2026-12-19", 1000), ("602697", "2026-10-25", 700)]
    assert cubiertos == {"280"}


def test_cuota_vencida_hace_mucho_se_da_por_cobrada():
    docs = [_cot("602600", "2026-03-01", 1000, [{"porcentaje": 100, "condicion": "por adelantado", "monto": 1000}], req="280")]
    movs, avisos, _ = fc.ingresos_cotizaciones(docs, [{"numero": 280, "estado": "Adjudicado"}], [], HOY, SUP)
    assert movs == [] and "se dan por cobradas" in avisos[0]


def test_probables_ponderados_por_la_tasa_y_sin_los_ya_cotizados():
    reqs = [{"numero": 300, "estado": "Ofertado", "titulo": "A", "valorOfertado": 1000000, "cierre": "2026-10-10"},
            {"numero": 301, "estado": "Ofertado", "titulo": "B", "valorOfertado": 500000},
            {"numero": 302, "estado": "Descartado", "valorOfertado": 9}]
    movs, _ = fc.ingresos_probables(reqs, 0.3, {"301"}, HOY, SUP)
    assert [(m["referencia"], m["monto"], m["fecha"]) for m in movs] == [("300", 300000, "2026-12-09")]
    assert fc.ingresos_probables(reqs, None, set(), HOY, SUP)[0] == []


def test_una_oferta_que_domina_lo_probable_se_avisa():
    reqs = [{"numero": 901, "estado": "Ofertado", "titulo": "Planta, etapa 2", "valorOfertado": 400000000},
            {"numero": 280, "estado": "Ofertado", "valorOfertado": 3500000}]
    movs, avisos = fc.ingresos_probables(reqs, 0.31, set(), HOY, SUP)
    assert movs[0]["concepto"] == "N° 901 · Planta, etapa 2 (31 % de 400.000.000)"     # la coma del título se respeta
    assert avisos ==["La oferta N° 901 es el 99 % de lo probable: se gana o se pierde entera ($400.000.000 o nada), no llega ponderada."]
    assert fc.ingresos_probables(reqs[:1], 0.31, set(), HOY, SUP)[1] == []      # una sola: nada que comparar


def test_cotizaciones_sin_proyecto_se_avisan():
    docs = [_cot("602700", "2026-09-25", 900, []), _cot("602701", "2026-09-26", 900, [], req="280")]
    movs, avisos, _ = fc.ingresos_cotizaciones(docs, [{"numero": 280, "estado": "Adjudicado"}], [], HOY, SUP)
    assert [m["referencia"] for m in movs] == ["602701"]
    assert any(a.startswith("1 cotización(es) de Sistema QUEMPIN sin proyecto") for a in avisos)


def test_por_ejecutar_con_iva_menos_lo_comprometido_repartido_hasta_el_cierre():
    proyectos = [{"tag": "UMAG", "nombre": "Lab", "cierre": "2026-12-15",
                  "porEjecutar": {"Materiales": 100000, "Equipos": 0, "Mano de Obra": 50000, "Otros": None}},
                 {"tag": "SINAV", "porEjecutar": None}]
    movs, _ = fc.egresos_por_ejecutar(proyectos, {"UMAG": 29000}, HOY, SUP)
    # 100.000 × 1,19 + 50.000 − 29.000 = 140.000, en 3 meses (oct, nov, dic)
    assert [m["mes"] for m in movs] == ["2026-10", "2026-11", "2026-12"]
    assert sum(m["monto"] for m in movs) == 140000


def test_resumen_por_mes_y_acumulado_con_saldo_inicial():
    sup = dict(SUP, saldoInicial=1000)
    movs = [fc._movimiento(date(2026, 9, 5), "egreso", "real", 400, "x", "CC"),
            fc._movimiento(date(2026, 10, 5), "ingreso", "comprometido", 500, "x", "SQ"),
            fc._movimiento(date(2026, 10, 9), "egreso", "comprometido", 200, "x", "CC"),
            fc._movimiento(date(2026, 11, 1), "ingreso", "probable", 100, "x", "PI")]
    meses = {m["mes"]: m for m in fc.resumen_por_mes(movs, HOY, sup)}
    assert meses["2026-09"]["egreso_real"] == 400 and meses["2026-09"]["acumulado"] is None
    assert meses["2026-10"]["neto"] == 300 and meses["2026-10"]["acumulado"] == 1300
    assert meses["2026-11"]["neto"] == 100 and meses["2026-11"]["netoSinProbables"] == 0 and meses["2026-11"]["acumulado"] == 1400
    assert len(meses) == SUP["mesesHistoria"] + SUP["mesesProyeccion"]


@pytest.fixture
def carpeta(tmp_path):
    raiz = ic.asegurar_carpeta(tmp_path / "Intercambio")
    ic.publicar(raiz, "documentos-comerciales", "sistema-quempin", {"documentos": [
        _cot("602695", "2026-09-20", 2000, [{"porcentaje": 100, "condicion": "por adelantado", "monto": 2000}], req="280"),
        _oc("612611", "2026-09-20", 300000, tag="UMAG")]})
    ic.publicar(raiz, "requerimientos", "planilla-requerimientos", {"fuente": "x", "resumen": {"tasaAdjudicacion": 0.25},
        "requerimientos": [{"numero": 280, "estado": "Adjudicado", "titulo": "Lab"},
                           {"numero": 300, "estado": "Ofertado", "titulo": "Otra", "valorOfertado": 400000, "cierre": "2026-10-20"}]})
    ic.publicar(raiz, "analisis-financiero", "analisis-financiero", {"moneda": "CLP", "categorias": [], "mensajes": {}, "proyectos": [
        {"tag": "UMAG", "nombre": "Lab", "cierre": "2026-11-30", "porEjecutar": {"Materiales": 500000, "Mano de Obra": 100000},
         "proyectados": {}, "reales": {}}]})
    foto = tmp_path / "centro-de-costos.json"
    foto.write_text(json.dumps({"documentos": [_cc("UMAG-001", "2026-09-10", 119000),
                                               _cc("UMAG-002", "2026-09-28", 60000, estado="Pendiente")]}), encoding="utf-8")
    return raiz, foto


def test_armar_desde_las_publicaciones_y_escribir_el_excel(carpeta, tmp_path):
    raiz, foto = carpeta
    datos = fc.armar(raiz, foto, hoy=HOY, sup=dict(SUP))
    clases = {(m["sentido"], m["clase"]) for m in datos["movimientos"]}
    assert clases == {("egreso", "real"), ("egreso", "comprometido"), ("ingreso", "comprometido"),
                      ("ingreso", "probable"), ("egreso", "estimado")}
    estimado = sum(m["monto"] for m in datos["movimientos"] if m["clase"] == "estimado")
    # 500.000 × 1,19 + 100.000 − (OC 300.000 + factura pendiente 60.000) = 335.000
    assert estimado == 335000
    assert "saldoInicial" in datos["supuestos"]
    ruta = fc.escribir_excel(datos, tmp_path / "Flujo de Caja.xlsx")
    libro = openpyxl.load_workbook(ruta)
    assert libro.sheetnames == ["Resumen", "Movimientos", "Supuestos"]
    assert libro["Resumen"]["A2"].value == "Cuotas por cobrar"
    assert libro["Movimientos"].max_row == len(datos["movimientos"]) + 1


def test_sin_carpeta_ni_foto_avisa_y_no_falla(tmp_path):
    datos = fc.armar(None, tmp_path / "no-existe.json", hoy=HOY, sup=dict(SUP))
    assert datos["movimientos"] == []
    assert any("Centro de Costos" in a for a in datos["avisos"]) and any("intercambio" in a for a in datos["avisos"])


def test_parametros_propios_solo_las_claves_conocidas(tmp_path):
    ruta = tmp_path / "parametros_flujo_caja.json"
    ruta.write_text(json.dumps({"saldoInicial": 5000000, "otraCosa": 1}), encoding="utf-8")
    sup = fc.leer_parametros(ruta)
    assert sup["saldoInicial"] == 5000000 and "otraCosa" not in sup
