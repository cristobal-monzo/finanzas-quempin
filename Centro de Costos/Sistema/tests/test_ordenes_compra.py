# -*- coding: utf-8 -*-
"""Ordenes de compra de Sistema QUEMPIN en el Centro de Costos.

Una OC con proyecto entra al libro como costo del proyecto en la proxima
corrida de 'run'; cuando se registra su factura, la factura la reemplaza (o la
deja por el saldo si es parcial). Proveedores, folios y montos son ficticios:
este repositorio es publico.
"""

import json

import openpyxl
import pytest

import auditor_centro_costos as acc

RUT_PROVEEDOR = "76123456-7"


def _oc(folio="900001", fecha="2026-07-10", items=((2, "1000.00"), (1, "500.00")), neto=None,
        moneda="CLP", exenta=False, req="314", tag="", pais="Chile"):
    lineas = [{"descripcion": f"Valvula {i}", "detalles": ["Acero inoxidable"],
               "cantidad": str(c), "precio_unitario": p} for i, (c, p) in enumerate(items, 1)]
    bruto = sum(int(c) * float(p) for c, p in items)
    neto = bruto if neto is None else neto
    iva = 0 if exenta else round(neto * 0.19, 2)
    return {
        "folio": folio, "tipo": "61", "fecha": fecha, "pais": pais, "moneda": moneda,
        "exento_impuesto": exenta, "autor": "XX", "referencia": "", "motivo": "",
        "items": lineas,
        "totales": {"bruto": f"{bruto:.2f}", "neto": f"{neto:.2f}", "iva": f"{iva:.2f}",
                    "total": f"{neto + iva:.2f}"},
        "proyecto": {"req": req, "tag": tag},
        "contraparte": {"rut": RUT_PROVEEDOR, "razon_social": "Proveedor Test SpA"},
        "ruta_pdf": f"salida\\OC\\Chile\\2026\\{folio}_Proveedor.pdf",
    }


def _sembrar_quempin(carpeta, *ocs):
    for oc in ocs:
        destino = carpeta / "documentos" / "61" / "2026"
        destino.mkdir(parents=True, exist_ok=True)
        (destino / f"{oc['folio']}.json").write_text(json.dumps(oc), encoding="utf-8")
        pdf = carpeta / oc["ruta_pdf"].replace("\\", "/")
        pdf.parent.mkdir(parents=True, exist_ok=True)
        pdf.write_bytes(b"%PDF-1.4 OC")


def _factura(archivo, neto, n_documento="5551", fecha="20-07-2026", **kwargs):
    base = {
        "archivo": archivo, "proyecto": "Proyecto Test", "tipo_proyecto": "Obra",
        "fecha": fecha, "n_documento": n_documento, "tipo_documento": "Factura",
        "proveedor": "Proveedor Test SpA", "rut_proveedor": RUT_PROVEEDOR,
        "categoria": "Materiales", "estado": "Pagado",
        "items": [{"nombre_item": "Valvula", "descripcion": "d", "categoria_item": "Materiales",
                   "cantidad": 1, "p_unitario_sin_iva": neto}],
        "iva": round(neto * 0.19),
    }
    base.update(kwargs)
    return base


# ── Entrada de datos desde la OC ────────────────────────────────────────────

def test_entrada_lleva_el_descuento_en_los_precios_y_suma_el_neto():
    oc = _oc(neto=2250.0)  # 10 % de descuento sobre 2500

    entrada = acc.entrada_desde_oc(oc, "Proyecto Test", "OC 900001.pdf")

    assert [i["p_unitario_sin_iva"] for i in entrada["items"]] == [900, 450]
    assert acc.total_sin_iva_items(entrada["items"]) == 2250
    assert entrada["iva"] == round(2250 * 0.19)
    assert entrada["tipo_documento"] == "Orden de Compra"
    assert entrada["estado"] == "Pendiente"
    assert entrada["fecha"] == "10-07-2026"
    assert entrada["oc_folio"] == "900001"


def test_entrada_en_dolares_se_pasa_a_pesos_con_el_cambio_dado():
    oc = _oc(items=((2, "10.00"),), moneda="USD")

    entrada = acc.entrada_desde_oc(oc, "Proyecto Test", "OC 900001.pdf", tipo_cambio=950.5,
                                   nota_cambio=", pasada a pesos")

    assert entrada["items"][0]["p_unitario_sin_iva"] == 9505
    assert "pasada a pesos" in entrada["notas"]


def test_entrada_exenta_no_lleva_impuesto():
    entrada = acc.entrada_desde_oc(_oc(exenta=True), "Proyecto Test", "OC 900001.pdf")

    assert entrada["tipo_documento"] == "Orden de Compra Exenta"
    assert entrada["iva"] == 0


def test_proveedor_ya_conocido_por_rut_conserva_su_nombre():
    datos = [_factura("A.pdf", 100, proveedor="PROVEEDOR TEST S.P.A.")]

    entrada = acc.entrada_desde_oc(_oc(), "Proyecto Test", "OC 900001.pdf", datos)

    assert entrada["proveedor"] == "PROVEEDOR TEST S.P.A."
    assert entrada["tipo_proyecto"] == "Obra"


# ── Carpeta del proyecto ────────────────────────────────────────────────────

def test_carpeta_existente_por_numero_de_requerimiento(tmp_path):
    (tmp_path / "314. Proyecto Test").mkdir()
    (tmp_path / "31. Otro").mkdir()

    carpeta, existe = acc.carpeta_para_oc(tmp_path, _oc(req="314"))

    assert (carpeta.name, existe) == ("314. Proyecto Test", True)


def test_carpeta_nueva_con_el_nombre_publicado(tmp_path):
    proyectos = [{"req": "317", "carpetas": {"facturas": [], "ejecucion": [
        "Gestion\\Proyectos\\317. Proyecto Nuevo"]}}]

    carpeta, existe = acc.carpeta_para_oc(tmp_path, _oc(req="317"), proyectos)

    assert (carpeta.name, existe) == ("317. Proyecto Nuevo", False)


def test_sin_carpeta_ni_nombre_no_hay_destino(tmp_path):
    assert acc.carpeta_para_oc(tmp_path, _oc(req="999")) == (None, False)


# ── Traer las OC ────────────────────────────────────────────────────────────

@pytest.fixture
def quempin(tmp_path):
    carpeta = tmp_path / "Carpeta compartida QUEMPIN"
    docs = tmp_path / "Facturas"
    (docs / "314. Proyecto Test").mkdir(parents=True)
    return carpeta, docs


def test_sin_aplicar_no_toca_nada(quempin):
    carpeta, docs = quempin
    _sembrar_quempin(carpeta, _oc())
    registro = {"ordenes": {}}

    traidas, avisos = acc.traer_ordenes_compra(docs, registro, [], carpeta, aplicar=False)

    assert [t["accion"] for t in traidas] == ["traer"]
    assert registro == {"ordenes": {}}
    assert list((docs / "314. Proyecto Test").iterdir()) == []


def test_aplicar_copia_el_pdf_y_anota_la_oc(quempin):
    carpeta, docs = quempin
    _sembrar_quempin(carpeta, _oc(), _oc(folio="900002", req=""))
    registro = {"ordenes": {}}

    traidas, avisos = acc.traer_ordenes_compra(docs, registro, [], carpeta)

    assert [t["folio"] for t in traidas] == ["900001"]
    assert (docs / "314. Proyecto Test" / "OC 900001.pdf").exists()
    orden = registro["ordenes"]["900001"]
    assert orden["estado"] == "en_cc"
    assert orden["entrada"]["proyecto"] == "Proyecto Test"
    assert acc.entradas_ordenes_compra(registro) == [orden["entrada"]]
    assert any("sin proyecto" in a for a in avisos)

    # Ya anotada: la segunda corrida no la vuelve a traer.
    assert acc.traer_ordenes_compra(docs, registro, [], carpeta)[0] == []


def test_otro_pais_no_entra(quempin):
    carpeta, docs = quempin
    _sembrar_quempin(carpeta, _oc(pais="Peru"))

    assert acc.traer_ordenes_compra(docs, {"ordenes": {}}, [], carpeta)[0] == []


def test_en_dolares_usa_el_cambio_del_dia_de_la_oc(quempin):
    carpeta, docs = quempin
    _sembrar_quempin(carpeta, _oc(items=((1, "100.00"),), moneda="USD"))
    consultas = []

    def cambio(moneda, fecha):
        consultas.append((moneda, fecha.isoformat()))
        return 900.0, fecha

    registro = {"ordenes": {}}
    acc.traer_ordenes_compra(docs, registro, [], carpeta, consultar_cambio=cambio)

    assert consultas == [("USD", "2026-07-10")]
    assert registro["ordenes"]["900001"]["entrada"]["items"][0]["p_unitario_sin_iva"] == 90000


def test_sin_tipo_de_cambio_queda_para_la_proxima_corrida(quempin):
    carpeta, docs = quempin
    _sembrar_quempin(carpeta, _oc(moneda="USD"))

    def sin_cambio(moneda, fecha):
        raise acc.TipoCambioNoDisponibleError("sin conexion")

    registro = {"ordenes": {}}
    traidas, avisos = acc.traer_ordenes_compra(docs, registro, [], carpeta,
                                               consultar_cambio=sin_cambio)

    assert traidas == [] and registro == {"ordenes": {}}
    assert any("proxima corrida" in a for a in avisos)


def test_ya_facturada_por_el_total_no_entra(quempin):
    carpeta, docs = quempin
    _sembrar_quempin(carpeta, _oc())
    registro = {"ordenes": {}}

    traidas, _ = acc.traer_ordenes_compra(docs, registro, [_factura("F.pdf", 2500)], carpeta)

    assert [t["accion"] for t in traidas] == ["facturada_antes"]
    assert registro["ordenes"]["900001"]["estado"] == "facturada_antes"
    assert not (docs / "314. Proyecto Test" / "OC 900001.pdf").exists()


# ── Que factura es de que OC ────────────────────────────────────────────────

def _entrada_oc():
    return acc.entrada_desde_oc(_oc(), "Proyecto Test", "OC 900001.pdf")


def test_la_factura_que_nombra_la_oc_se_vincula():
    vinculadas, dudosas = acc.facturas_de_oc(
        _entrada_oc(), [_factura("F.pdf", 1000, oc="900001"), _factura("G.pdf", 2500, n_documento="2")])

    assert [d["archivo"] for d in vinculadas] == ["F.pdf"]
    assert dudosas == []


def test_una_sola_factura_que_calza_se_vincula_sola():
    datos = [_factura("F.pdf", 2520), _factura("G.pdf", 2500, proveedor="Otro SpA",
                                                rut_proveedor="77777777-7")]

    vinculadas, dudosas = acc.facturas_de_oc(_entrada_oc(), datos)

    assert [d["archivo"] for d in vinculadas] == ["F.pdf"]


def test_dos_facturas_que_calzan_no_se_adivina():
    datos = [_factura("F.pdf", 2500), _factura("G.pdf", 2500, n_documento="5552")]

    vinculadas, dudosas = acc.facturas_de_oc(_entrada_oc(), datos)

    assert vinculadas == [] and len(dudosas) == 2


def test_factura_muy_anterior_a_la_oc_no_calza():
    assert acc.facturas_de_oc(_entrada_oc(), [_factura("F.pdf", 2500, fecha="01-01-2026")]) == ([], [])


# ── De punta a punta en main() ──────────────────────────────────────────────

@pytest.fixture
def sandbox(tmp_path, monkeypatch):
    docs = tmp_path / "Facturas" / "314. Proyecto Test"
    docs.mkdir(parents=True)
    (tmp_path / "Excel").mkdir()
    carpeta_quempin = tmp_path / "Carpeta compartida QUEMPIN"
    _sembrar_quempin(carpeta_quempin, _oc())

    cfg = dict(acc.PAISES["CL"])
    cfg.update({
        "ruta_docs": tmp_path / "Facturas",
        "ruta_excel": tmp_path / "Excel" / "Centro de Costos.xlsx",
        "ruta_backups": tmp_path / "Excel" / "Respaldos",
        "ruta_json": tmp_path / "datos_extraidos.json",
        "ruta_reconciliacion": tmp_path / "reconciliacion_archivos.json",
        "ruta_correcciones": tmp_path / "correcciones_manuales.json",
        "ruta_ordenes_compra": tmp_path / "ordenes_compra.json",
        "ruta_errores_md": tmp_path / "ERRORES.md",
        "ruta_logs": tmp_path / "logs",
        "ruta_visualizador_web": tmp_path / "Visualizador Web (inexistente)",
        "ruta_excel_sitio_comunicacion": None,
        "prefijos_proyecto": {"Proyecto Test": "PTST"},
    })
    monkeypatch.setitem(acc.PAISES, "TEST", cfg)
    monkeypatch.setattr(acc, "RAIZ_ANALISIS_FINANCIERO", tmp_path / "AF (inexistente)")
    monkeypatch.setattr(acc, "proyectos_publicados", lambda: [])
    monkeypatch.setenv("QUEMPIN_CARPETA_COMPARTIDA", str(carpeta_quempin))
    (tmp_path / "correcciones_manuales.json").write_text("[]", encoding="utf-8")
    (tmp_path / "datos_extraidos.json").write_text("[]", encoding="utf-8")
    return {"tmp": tmp_path, "docs": docs}


def _facturas(sandbox, *facturas):
    """Agrega las facturas al JSON y deja en la carpeta el archivo de las
    nuevas (las de corridas anteriores ya fueron renombradas)."""
    ruta_json = sandbox["tmp"] / "datos_extraidos.json"
    datos = json.loads(ruta_json.read_text(encoding="utf-8"))
    for f in facturas:
        (sandbox["docs"] / f["archivo"]).write_bytes(b"")
    ruta_json.write_text(json.dumps(datos + list(facturas), ensure_ascii=False), encoding="utf-8")


def _libro(sandbox):
    wb = openpyxl.load_workbook(str(sandbox["tmp"] / "Excel" / "Centro de Costos.xlsx"))
    ws_m, ws_d = wb["Master"], wb["Detalle"]
    master = {ws_m.cell(row=r, column=1).value: [ws_m.cell(row=r, column=c).value for c in range(1, 17)]
              for r in range(2, acc.ultima_fila_datos(ws_m) + 1)}
    netos = {}
    for r in range(2, acc.ultima_fila_datos(ws_d) + 1):
        n_ref = ws_d.cell(row=r, column=1).value
        netos[n_ref] = netos.get(n_ref, 0) + ws_d.cell(row=r, column=10).value
    formulas_ok = all(ws_m.cell(row=r, column=11).value == f"=SUMIF(Detalle!$A:$A,$A{r},Detalle!$J:$J)"
                      for r in range(2, acc.ultima_fila_datos(ws_m) + 1))
    return master, netos, formulas_ok


def _registro_oc(sandbox):
    return acc.cargar_ordenes_compra(sandbox["tmp"] / "ordenes_compra.json")


def test_la_oc_entra_queda_por_el_saldo_y_la_reemplaza_su_factura(sandbox):
    # 1. La OC entra al libro como costo del proyecto, pendiente de pago.
    acc.main(pais="TEST")

    master, netos, _ = _libro(sandbox)
    assert list(master) == ["PTST-001"]
    fila_oc = master["PTST-001"]
    assert (fila_oc[4], fila_oc[5], fila_oc[13]) == ("900001", "Orden de Compra", "Pendiente")
    assert netos["PTST-001"] == 2500
    assert _registro_oc(sandbox)["ordenes"]["900001"]["estado"] == "en_cc"

    # 2. Llega una factura por parte de la OC: la OC queda por el saldo.
    parcial = _factura("F1.pdf", 1000, oc="900001")
    _facturas(sandbox, parcial)
    acc.main(pais="TEST")

    master, netos, formulas_ok = _libro(sandbox)
    assert set(master) == {"PTST-001", "PTST-002"}
    assert netos == {"PTST-001": 1500, "PTST-002": 1000}
    assert master["PTST-001"][11] == round(1500 * 0.19)
    assert formulas_ok
    assert _registro_oc(sandbox)["ordenes"]["900001"]["facturado"] == 1000

    # 3. Llega el resto: la factura reemplaza a la OC y su PDF se archiva.
    _facturas(sandbox, _factura("F2.pdf", 1500, n_documento="5552", oc="900001"))
    acc.main(pais="TEST")

    master, netos, formulas_ok = _libro(sandbox)
    assert set(master) == {"PTST-002", "PTST-003"}
    assert sum(netos.values()) == 2500
    assert formulas_ok
    orden = _registro_oc(sandbox)["ordenes"]["900001"]
    assert orden["estado"] == "reemplazada"
    assert orden["facturas"] == ["PTST-002", "PTST-003"]
    assert not any(p.name.startswith("PTST-001") for p in sandbox["docs"].iterdir())
    archivados = list((sandbox["tmp"] / "Excel" / "Respaldos").rglob("*reemplazada por factura*"))
    assert len(archivados) == 1

    # 4. La OC sigue en Sistema QUEMPIN, pero no vuelve a entrar.
    acc.main(pais="TEST")

    master, _, _ = _libro(sandbox)
    assert set(master) == {"PTST-002", "PTST-003"}


def test_si_el_libro_esta_bloqueado_la_oc_no_se_cierra(sandbox, monkeypatch):
    acc.main(pais="TEST")
    _facturas(sandbox, _factura("F1.pdf", 2500, oc="900001"))
    monkeypatch.setattr(acc, "excel_esta_bloqueado", lambda ruta: True)

    acc.main(pais="TEST")

    assert _registro_oc(sandbox)["ordenes"]["900001"]["estado"] == "en_cc"
    assert any(p.name.startswith("PTST-001") for p in sandbox["docs"].iterdir())


def test_status_dice_que_traeria_y_que_cerraria_sin_tocar_nada(sandbox, capsys):
    import importlib.util
    from pathlib import Path

    ruta = (Path(__file__).resolve().parent.parent.parent
            / ".claude" / "skills" / "Registro_Centro_de_Costos" / "driver.py")
    spec = importlib.util.spec_from_file_location("driver_oc_bajo_prueba", ruta)
    driver = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(driver)

    driver.cmd_status(pais="TEST")

    salida = capsys.readouterr().out
    assert "Traería OC 900001 a 314. Proyecto Test" in salida
    assert "incluye 1 OC nueva" in salida
    assert _registro_oc(sandbox) == {"ordenes": {}}
    assert list(sandbox["docs"].iterdir()) == []

    acc.main(pais="TEST")
    _facturas(sandbox, _factura("F1.pdf", 1000, oc="900001"))
    driver.cmd_status(pais="TEST")

    assert "OC 900001: quedaría por el saldo (F1.pdf)" in capsys.readouterr().out
