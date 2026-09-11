# -*- coding: utf-8 -*-
"""
Tests del enlace hallazgo -> fila de Master sobre un corpus YA registrado.

Motivo (revision del 2026-09-10, un dia despues de crear el registro): el
registro de errores funcionaba de punta a punta en el benchmark y era
inservible sobre los datos reales. anotar_n_ref() solo corre dentro del bucle
de escritura del PASO 6, o sea solo para documentos que se registran en esa
misma corrida; con 728 filas ya en Master y 0 pendientes, ningun hallazgo
llegaba a tener N Ref y corregir_hallazgos() los rechazaba TODOS con "el
documento todavia no tiene fila en Master".

Y al reves: la validacion corre sobre datos_extraidos.json, que es la entrada
del pipeline y no se reescribe nunca, asi que un documento corregido hace
meses en el Excel volvia a producir el mismo hallazgo y el operador tenia que
arreglar de nuevo lo ya arreglado.

Estos tests fijan las dos reglas que cierran ese hueco:
  * se enlaza solo cuando la identificacion es INEQUIVOCA (un numero repetido
    no identifica una fila: se omite en vez de adjudicar cualquiera de las dos);
  * se cierra solo con la evidencia que el modulo ya trata como adjudicacion
    humana de ESA celda (azul marino, o la bitacora de correcciones).
"""

from datetime import datetime

import openpyxl
import pytest

import auditor_centro_costos as acc


def _master(filas):
    """Hoja Master minima. `filas`: lista de (n_ref, proyecto, n_documento)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master"
    for c, h in enumerate(acc.ENCABEZADOS_MASTER, 1):
        ws.cell(row=1, column=c, value=h)
    for i, (n_ref, proyecto, n_doc) in enumerate(filas, start=2):
        ws.cell(row=i, column=1, value=n_ref)
        ws.cell(row=i, column=2, value=proyecto)
        ws.cell(row=i, column=4, value=datetime(2026, 7, 15))
        ws.cell(row=i, column=5, value=n_doc)
    return ws


def _doc(proyecto, archivo, n_documento):
    return {"proyecto": proyecto, "archivo": archivo, "n_documento": n_documento}


# ── mapa_documento_a_n_ref ──────────────────────────────────────────────────

def test_enlaza_por_n_documento_unico():
    ws = _master([("BALF-001", "Alfa", "12345")])
    mapa = acc.mapa_documento_a_n_ref(ws, [_doc("Alfa", "DOC_1.jpg", "12345")],
                                      correcciones=[])
    assert mapa == {"Alfa\\DOC_1.jpg": "BALF-001"}


def test_el_renombrado_no_rompe_el_enlace():
    """El puente NO puede ser 'Archivo origen': el renombrado automatico lo
    reescribe a '<N Ref>_<Proveedor>_<Fecha>' y ninguna de las 681 claves del
    JSON real calzaba. El N Documento sobrevive al renombrado."""
    ws = _master([("BALF-001", "Alfa", "12345")])
    ws.cell(row=2, column=15, value="Alfa\\BALF-001_Proveedor_15-07-2026.pdf")
    mapa = acc.mapa_documento_a_n_ref(ws, [_doc("Alfa", "IMG_9999.HEIC", "12345")],
                                      correcciones=[])
    assert mapa == {"Alfa\\IMG_9999.HEIC": "BALF-001"}


def test_no_enlaza_un_numero_repetido_en_master():
    """Un N Documento es unico por emisor, no globalmente: si aparece en dos
    filas no identifica ninguna. Adjudicar cualquiera de las dos escribiria la
    correccion en el documento equivocado."""
    ws = _master([("BALF-001", "Alfa", "12345"), ("BBET-007", "Beta", "12345")])
    mapa = acc.mapa_documento_a_n_ref(ws, [_doc("Alfa", "DOC_1.jpg", "12345")],
                                      correcciones=[])
    assert mapa == {}


def test_no_enlaza_un_numero_repetido_en_el_json():
    """El otro lado del mismo criterio: dos documentos de origen con el mismo
    numero (el caso tipico de un duplicado) no se pueden distinguir."""
    ws = _master([("BALF-001", "Alfa", "12345")])
    datos = [_doc("Alfa", "DOC_1.jpg", "12345"), _doc("Alfa", "DOC_2.jpg", "12345")]
    assert acc.mapa_documento_a_n_ref(ws, datos, correcciones=[]) == {}


def test_enlaza_un_sin_numero_porque_su_texto_es_unico():
    """'S/N (<archivo>)' no es un numero real (es_n_documento_real lo rechaza)
    pero si es un identificador unico, y N_DOC_ILEGIBLE es justamente la clase
    con mas hallazgos abiertos. El puente se decide por unicidad, no por
    'realidad' del numero."""
    ws = _master([("BALF-001", "Alfa", "S/N (boleta_rota.jpg)")])
    mapa = acc.mapa_documento_a_n_ref(
        ws, [_doc("Alfa", "boleta_rota.jpg", "S/N (boleta_rota.jpg)")],
        correcciones=[])
    assert mapa == {"Alfa\\boleta_rota.jpg": "BALF-001"}


def test_los_marcadores_sin_numero_no_enlazan_nada():
    """172 de las entradas reales son peajes con 'N/A'. Como se repiten, el
    criterio de unicidad los descarta solo."""
    ws = _master([("BALF-001", "Alfa", "N/A"), ("BALF-002", "Alfa", "N/A")])
    datos = [_doc("Alfa", "peaje_1.jpg", "N/A"), _doc("Alfa", "peaje_2.jpg", "N/A")]
    assert acc.mapa_documento_a_n_ref(ws, datos, correcciones=[]) == {}


def test_enlaza_por_la_bitacora_cuando_el_numero_ya_fue_corregido():
    """Si alguien corrigio el N Documento, Master tiene el valor bueno y el
    JSON el viejo: el primer puente no los encuentra. correcciones_manuales
    guarda el par (valor_anterior -> N Ref) de esa misma celda, que es el
    enlace exacto que se perdio."""
    ws = _master([("BALF-001", "Alfa", "998877")])
    correcciones = [{"n_ref": "BALF-001", "columna": 5, "campo": "N° Documento",
                     "valor_anterior": "S/N (foto_borrosa.jpg)",
                     "valor_corregido": "998877", "estado": "Aplicado"}]
    mapa = acc.mapa_documento_a_n_ref(
        ws, [_doc("Alfa", "foto_borrosa.jpg", "S/N (foto_borrosa.jpg)")],
        correcciones=correcciones)
    assert mapa == {"Alfa\\foto_borrosa.jpg": "BALF-001"}


def test_la_bitacora_no_pisa_un_enlace_directo():
    ws = _master([("BALF-001", "Alfa", "12345"), ("BALF-002", "Alfa", "777")])
    correcciones = [{"n_ref": "BALF-002", "columna": 5,
                     "valor_anterior": "12345", "valor_corregido": "777"}]
    mapa = acc.mapa_documento_a_n_ref(ws, [_doc("Alfa", "DOC_1.jpg", "12345")],
                                      correcciones=correcciones)
    assert mapa == {"Alfa\\DOC_1.jpg": "BALF-001"}


def test_la_bitacora_solo_aporta_la_columna_del_n_documento():
    """Una correccion de IVA no dice nada sobre que fila es cual."""
    ws = _master([("BALF-001", "Alfa", "998877")])
    correcciones = [{"n_ref": "BALF-001", "columna": 12,
                     "valor_anterior": "11111", "valor_corregido": "19000"}]
    mapa = acc.mapa_documento_a_n_ref(ws, [_doc("Alfa", "x.jpg", "11111")],
                                      correcciones=correcciones)
    assert mapa == {}


# ── anotar_n_ref_desde_master ───────────────────────────────────────────────

def _registro_con(hallazgos):
    registro = {"version": 1, "errores": []}
    acc.fusionar_hallazgos(registro, hallazgos, hoy="2026-09-10")
    return registro


def test_anota_el_n_ref_de_un_documento_ya_registrado():
    dato = {"proyecto": "Alfa", "archivo": "DOC_1.jpg", "n_documento": "12345"}
    registro = _registro_con([acc._hallazgo("PROVEEDOR_VACIO", dato)])
    assert registro["errores"][0]["n_ref"] is None

    ws = _master([("BALF-001", "Alfa", "12345")])
    assert acc.anotar_n_ref_desde_master(registro, ws, [dato]) == 1
    assert registro["errores"][0]["n_ref"] == "BALF-001"


def test_no_pisa_un_n_ref_ya_anotado():
    """anotar_n_ref() (PASO 6) sigue siendo la fuente preferente para lo que se
    escribe en la corrida; este relleno es solo para lo que ya estaba."""
    dato = {"proyecto": "Alfa", "archivo": "DOC_1.jpg", "n_documento": "12345"}
    registro = _registro_con([acc._hallazgo("PROVEEDOR_VACIO", dato)])
    acc.anotar_n_ref(registro, "Alfa\\DOC_1.jpg", "BALF-099")

    ws = _master([("BALF-001", "Alfa", "12345")])
    assert acc.anotar_n_ref_desde_master(registro, ws, [dato]) == 0
    assert registro["errores"][0]["n_ref"] == "BALF-099"


# ── cerrar_hallazgos_ya_corregidos ──────────────────────────────────────────

def _abierto_en(ws, columna, codigo="IMPUESTO_MENOR"):
    dato = {"proyecto": "Alfa", "archivo": "DOC_1.jpg", "n_documento": "12345"}
    registro = _registro_con([acc._hallazgo(codigo, dato, valor_actual=1)])
    acc.anotar_n_ref_desde_master(registro, ws, [dato])
    entrada = registro["errores"][0]
    assert entrada["n_ref"] == "BALF-001"
    assert entrada["columna"] == columna
    return registro, entrada


def test_cierra_el_hallazgo_cuya_celda_quedo_azul_marino():
    ws = _master([("BALF-001", "Alfa", "12345")])
    ws.cell(row=2, column=12, value=19000).font = acc.AZUL_MARINO_FONT
    registro, entrada = _abierto_en(ws, 12)

    cerrados = acc.cerrar_hallazgos_ya_corregidos(registro, ws, correcciones=[],
                                                  hoy="2026-09-10")
    assert [c["id"] for c in cerrados] == [entrada["id"]]
    assert entrada["estado"] == "resuelto"
    assert "azul marino" in entrada["resolucion"]


def test_cierra_el_hallazgo_anotado_en_la_bitacora():
    ws = _master([("BALF-001", "Alfa", "12345")])
    ws.cell(row=2, column=12, value=19000)
    registro, entrada = _abierto_en(ws, 12)
    correcciones = [{"n_ref": "BALF-001", "columna": 12,
                     "valor_anterior": 1, "valor_corregido": 19000}]

    assert acc.cerrar_hallazgos_ya_corregidos(registro, ws, correcciones=correcciones,
                                              hoy="2026-09-10")
    assert entrada["estado"] == "resuelto"
    assert "correcciones_manuales.json" in entrada["resolucion"]


def test_no_cierra_lo_que_no_tiene_evidencia():
    """El freno: sin adjudicacion humana de esa celda, el hallazgo sigue
    abierto. Cerrarlo seria esconder un error para que la lista se vea corta."""
    ws = _master([("BALF-001", "Alfa", "12345")])
    ws.cell(row=2, column=12, value=1)
    registro, entrada = _abierto_en(ws, 12)

    assert acc.cerrar_hallazgos_ya_corregidos(registro, ws, correcciones=[]) == []
    assert entrada["estado"] == "abierto"


def test_la_evidencia_es_por_celda_no_por_documento():
    """Una correccion de N Documento no adjudica el descuadre de impuesto del
    mismo documento."""
    ws = _master([("BALF-001", "Alfa", "12345")])
    ws.cell(row=2, column=5).font = acc.AZUL_MARINO_FONT
    ws.cell(row=2, column=12, value=1)
    registro, entrada = _abierto_en(ws, 12)

    assert acc.cerrar_hallazgos_ya_corregidos(registro, ws, correcciones=[]) == []
    assert entrada["estado"] == "abierto"


def test_no_se_reabre_en_la_corrida_siguiente():
    """datos_extraidos.json no se reescribe, asi que el hallazgo se vuelve a
    producir identico en cada corrida. Cerrarlo sin dejar anotado el valor
    adjudicado era justamente lo que lo hacia reaparecer para siempre."""
    ws = _master([("BALF-001", "Alfa", "12345")])
    ws.cell(row=2, column=12, value=19000).font = acc.AZUL_MARINO_FONT
    dato = {"proyecto": "Alfa", "archivo": "DOC_1.jpg", "n_documento": "12345"}
    hallazgo = acc._hallazgo("IMPUESTO_MENOR", dato, valor_actual=1)

    registro = _registro_con([hallazgo])
    acc.anotar_n_ref_desde_master(registro, ws, [dato])
    acc.cerrar_hallazgos_ya_corregidos(registro, ws, correcciones=[], hoy="2026-09-10")

    acc.fusionar_hallazgos(registro, [dict(hallazgo)], hoy="2026-09-11")
    assert registro["errores"][0]["estado"] == "resuelto"
    assert registro["errores"][0]["origen_sin_corregir"] is True


def test_se_reabre_si_el_dato_de_origen_cambia():
    """El contrapeso: un valor distinto es un error que nadie adjudico."""
    ws = _master([("BALF-001", "Alfa", "12345")])
    ws.cell(row=2, column=12, value=19000).font = acc.AZUL_MARINO_FONT
    dato = {"proyecto": "Alfa", "archivo": "DOC_1.jpg", "n_documento": "12345"}

    registro = _registro_con([acc._hallazgo("IMPUESTO_MENOR", dato, valor_actual=1)])
    acc.anotar_n_ref_desde_master(registro, ws, [dato])
    acc.cerrar_hallazgos_ya_corregidos(registro, ws, correcciones=[], hoy="2026-09-10")

    acc.fusionar_hallazgos(registro, [acc._hallazgo("IMPUESTO_MENOR", dato, valor_actual=2)],
                           hoy="2026-09-11")
    assert registro["errores"][0]["estado"] == "abierto"


def test_no_toca_un_hallazgo_sin_columna():
    """Los duplicados y los DOC_SIN_DATOS no se arreglan escribiendo una celda:
    no hay celda que pueda dar evidencia de nada."""
    ws = _master([("BALF-001", "Alfa", "12345")])
    dato = {"proyecto": "Alfa", "archivo": "DOC_1.jpg", "n_documento": "12345"}
    registro = _registro_con([acc._hallazgo("DOC_SIN_DATOS", dato)])
    acc.anotar_n_ref_desde_master(registro, ws, [dato])

    assert acc.cerrar_hallazgos_ya_corregidos(registro, ws, correcciones=[]) == []
    assert registro["errores"][0]["estado"] == "abierto"


# ── sincronizar_registro_errores ────────────────────────────────────────────

def test_sincronizar_deja_el_hallazgo_resoluble(tmp_path):
    ws = _master([("BALF-001", "Alfa", "12345")])
    datos = [{"proyecto": "Alfa", "archivo": "DOC_1.jpg", "n_documento": "12345",
              "fecha": "15-07-2026", "tipo_documento": "Factura",
              "proveedor": "Demo SpA", "categoria": "Materiales", "iva": 1,
              "items": [{"nombre_item": "x", "cantidad": 1, "p_unitario_sin_iva": 100000}]}]
    registro, resumen = acc.sincronizar_registro_errores(
        ws, datos, pendientes=[], hoy="2026-09-10",
        ruta_registro=tmp_path / "errores_detectados.json")

    abiertos = acc.hallazgos_abiertos(registro)
    assert abiertos, "el descuadre de impuesto tiene que aparecer"
    assert all(e["n_ref"] == "BALF-001" for e in abiertos if e["columna"])
    assert resumen["enlazados"] >= 1


def test_sincronizar_reporta_los_archivos_sin_entrada_en_el_json(tmp_path):
    """`pendientes` se pasa explicito para poder reportar DOC_SIN_DATOS. Si se
    omitiera la busqueda, fusionar_hallazgos() daria por desaparecidos -- y
    cerraria -- los que siguen abiertos."""
    ws = _master([("BALF-001", "Alfa", "12345")])
    registro, _ = acc.sincronizar_registro_errores(
        ws, [], pendientes=[{"proyecto": "Alfa", "archivo": "sin_entrada.jpg"}],
        hoy="2026-09-10", ruta_registro=tmp_path / "errores_detectados.json")

    codigos = [e["codigo"] for e in acc.hallazgos_abiertos(registro)]
    assert codigos == ["DOC_SIN_DATOS"]


def test_sincronizar_persiste_el_registro(tmp_path):
    ruta = tmp_path / "errores_detectados.json"
    ws = _master([("BALF-001", "Alfa", "12345")])
    acc.sincronizar_registro_errores(
        ws, [], pendientes=[{"proyecto": "Alfa", "archivo": "sin_entrada.jpg"}],
        hoy="2026-09-10", ruta_registro=ruta)

    assert ruta.exists()
    assert acc.cargar_registro_errores(ruta)["errores"]


def test_sincronizar_puede_no_escribir(tmp_path):
    ruta = tmp_path / "errores_detectados.json"
    ws = _master([("BALF-001", "Alfa", "12345")])
    acc.sincronizar_registro_errores(
        ws, [], pendientes=[{"proyecto": "Alfa", "archivo": "sin_entrada.jpg"}],
        hoy="2026-09-10", guardar=False, ruta_registro=ruta)

    assert not ruta.exists()


# ── desempate por neto (tercer puente) ─────────────────────────────────────

def _detalle(filas):
    """Hoja Detalle minima. `filas`: lista de (n_ref, cantidad, p_unitario)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Detalle"
    for c, h in enumerate(acc.ENCABEZADOS_DETALLE, 1):
        ws.cell(row=1, column=c, value=h)
    for i, (n_ref, cantidad, precio) in enumerate(filas, start=2):
        ws.cell(row=i, column=1, value=n_ref)
        ws.cell(row=i, column=8, value=cantidad)
        ws.cell(row=i, column=9, value=precio)
    return ws


def _con_items(proyecto, archivo, n_documento, neto):
    return {"proyecto": proyecto, "archivo": archivo, "n_documento": n_documento,
            "items": [{"cantidad": 1, "p_unitario_sin_iva": neto}]}


def test_desempata_un_numero_repetido_por_el_neto():
    """Caso real: una guia de despacho registrada con el numero de SU factura.
    1913313 apuntaba a tres filas distintas ($65.930, $70.271 y $28.354), asi
    que el hallazgo quedaba huerfano justo cuando SI habia algo que corregir.
    El neto de Detalle no es una heuristica: son los items que se escribieron
    desde esa misma entrada del JSON."""
    ws_m = _master([("BALF-001", "Alfa", "1913313"), ("BALF-002", "Alfa", "1913313"),
                    ("BALF-003", "Alfa", "1913313")])
    ws_d = _detalle([("BALF-001", 1, 65930), ("BALF-002", 1, 70271), ("BALF-003", 1, 28354)])
    datos = [_con_items("Alfa", "guia_a.pdf", "1913313", 70271)]

    mapa = acc.mapa_documento_a_n_ref(ws_m, datos, correcciones=[], ws_detalle=ws_d)
    assert mapa == {"Alfa\\guia_a.pdf": "BALF-002"}


def test_no_desempata_si_dos_candidatas_tienen_el_mismo_neto():
    """El mismo documento cargado dos veces: el neto no distingue nada y
    adjudicar cualquiera de las dos seria inventar. Se deja huerfano a
    proposito -- ese caso se resuelve borrando una fila, no escribiendo una
    celda."""
    ws_m = _master([("BALF-001", "Alfa", "137096"), ("BALF-002", "Alfa", "137096")])
    ws_d = _detalle([("BALF-001", 1, 1279260), ("BALF-002", 1, 1279260)])
    datos = [_con_items("Alfa", "copia.pdf", "137096", 1279260)]

    assert acc.mapa_documento_a_n_ref(ws_m, datos, correcciones=[], ws_detalle=ws_d) == {}


def test_sin_hoja_detalle_no_desempata():
    ws_m = _master([("BALF-001", "Alfa", "999"), ("BALF-002", "Alfa", "999")])
    datos = [_con_items("Alfa", "x.pdf", "999", 100)]

    assert acc.mapa_documento_a_n_ref(ws_m, datos, correcciones=[]) == {}


def test_el_desempate_no_pisa_un_enlace_directo():
    ws_m = _master([("BALF-001", "Alfa", "111"), ("BALF-002", "Alfa", "222"),
                    ("BALF-003", "Alfa", "222")])
    ws_d = _detalle([("BALF-001", 1, 500), ("BALF-002", 1, 500), ("BALF-003", 1, 900)])
    datos = [_con_items("Alfa", "unico.pdf", "111", 500)]

    mapa = acc.mapa_documento_a_n_ref(ws_m, datos, correcciones=[], ws_detalle=ws_d)
    assert mapa == {"Alfa\\unico.pdf": "BALF-001"}


def test_el_duplicado_ambiguo_se_puede_corregir_por_celda():
    """La causa mas comun resulto ser un numero mal leido, no un documento de
    mas: la taxonomia declara la columna del N Documento para que 'resolver'
    pueda arreglarlo sin borrar una fila."""
    severidad, _titulo, _accion, columna = acc.TAXONOMIA_ERRORES["DUPLICADO_AMBIGUO"]
    assert severidad == "error"
    assert columna == 5
    assert columna in acc.COLUMNAS_CORREGIBLES


# ── cuadre contra el LIBRO, no contra el JSON ──────────────────────────────

def _libro_cuadre(n_documento="12345", tipo="Factura", categoria="Ferreteria",
                  iva=19000, cantidad=1, precio=100000):
    ws_m = _master([("BALF-001", "Alfa", n_documento)])
    ws_m.cell(row=2, column=6, value=tipo)
    ws_m.cell(row=2, column=9, value=categoria)
    ws_m.cell(row=2, column=12, value=iva)
    ws_d = _detalle([("BALF-001", cantidad, precio)])
    return ws_m, ws_d


def _abierto_de_cuadre(ws_m, ws_d, codigo, iva_json):
    dato = {"proyecto": "Alfa", "archivo": "DOC_1.jpg", "n_documento": "12345",
            "categoria": "Ferreteria", "tipo_documento": "Factura"}
    registro = _registro_con([acc._hallazgo(codigo, dato, valor_actual=iva_json)])
    acc.anotar_n_ref_desde_master(registro, ws_m, [dato], ws_detalle=ws_d)
    assert registro["errores"][0]["n_ref"] == "BALF-001"
    return registro


def test_cierra_el_cuadre_cuando_el_libro_ya_cuadra():
    """El neto vive en Detalle y el impuesto en Master, asi que la correccion
    puede entrar por cualquiera de los dos lados y ninguna celda concreta
    sirve de evidencia. Lo que vale es que el documento cuadre en el libro:
    fue el caso de HPIN-157, donde el arreglo fue una cantidad de Detalle."""
    ws_m, ws_d = _libro_cuadre(iva=19000, precio=100000)
    registro = _abierto_de_cuadre(ws_m, ws_d, "IMPUESTO_MENOR", iva_json=1000)

    cerrados = acc.cerrar_hallazgos_ya_corregidos(registro, ws_m, correcciones=[],
                                                  ws_detalle=ws_d, hoy="2026-09-10")
    assert len(cerrados) == 1
    assert registro["errores"][0]["estado"] == "resuelto"
    assert "ya cuadra" in registro["errores"][0]["resolucion"]


def test_no_cierra_el_cuadre_si_el_libro_sigue_descuadrado():
    ws_m, ws_d = _libro_cuadre(iva=1000, precio=100000)
    registro = _abierto_de_cuadre(ws_m, ws_d, "IMPUESTO_MENOR", iva_json=1000)

    assert acc.cerrar_hallazgos_ya_corregidos(registro, ws_m, correcciones=[],
                                              ws_detalle=ws_d) == []
    assert registro["errores"][0]["estado"] == "abierto"


def test_un_impuesto_estimado_no_se_cierra_por_cuadre():
    """El freno. Su celda de impuesto se escribio calculando el 19%, asi que
    el libro cuadra POR CONSTRUCCION: cerrarlo por "ya cuadra" seria taparlo,
    y el hallazgo justamente dice "esto se completo con un supuesto". Solo lo
    cierra una persona que mire el documento."""
    ws_m, ws_d = _libro_cuadre(categoria="Combustible", iva=19000, precio=100000)
    registro = _abierto_de_cuadre(ws_m, ws_d, "IMPUESTO_ESTIMADO", iva_json=None)

    assert acc.cerrar_hallazgos_ya_corregidos(registro, ws_m, correcciones=[],
                                              ws_detalle=ws_d) == []
    assert registro["errores"][0]["estado"] == "abierto"


def test_sin_hoja_detalle_no_se_cierra_por_cuadre():
    ws_m, ws_d = _libro_cuadre(iva=19000, precio=100000)
    registro = _abierto_de_cuadre(ws_m, ws_d, "IMPUESTO_MENOR", iva_json=1000)

    assert acc.cerrar_hallazgos_ya_corregidos(registro, ws_m, correcciones=[]) == []


def test_cuadre_en_el_libro_usa_el_neto_de_detalle_no_el_del_json():
    ws_m, ws_d = _libro_cuadre(iva=19000, cantidad=2, precio=50000)
    assert acc.cuadre_en_el_libro(ws_m, ws_d, "BALF-001") is None

    ws_m2, ws_d2 = _libro_cuadre(iva=19000, cantidad=2, precio=80000)
    assert acc.cuadre_en_el_libro(ws_m2, ws_d2, "BALF-001") == "error"


def test_cierra_el_item_agrupado_cuando_ya_no_queda_ninguno():
    """ITEM_AGRUPADO no declara columna (no se arregla escribiendo una celda
    de Master), asi que su evidencia tampoco puede ser una celda: es que en
    Detalle ya no quede ninguna fila 'varios' de ese documento."""
    ws_m, ws_d = _libro_cuadre()
    dato = {"proyecto": "Alfa", "archivo": "DOC_1.jpg", "n_documento": "12345"}
    registro = _registro_con([acc._hallazgo("ITEM_AGRUPADO", dato,
                                            valor_actual="Materiales varios")])
    acc.anotar_n_ref_desde_master(registro, ws_m, [dato], ws_detalle=ws_d)

    cerrados = acc.cerrar_hallazgos_ya_corregidos(registro, ws_m, correcciones=[],
                                                  ws_detalle=ws_d, hoy="2026-09-10")
    assert len(cerrados) == 1
    assert "varios" in registro["errores"][0]["resolucion"]


def test_no_cierra_el_item_agrupado_si_sigue_habiendo_uno():
    ws_m, ws_d = _libro_cuadre()
    ws_d.cell(row=2, column=5, value="Materiales varios")
    dato = {"proyecto": "Alfa", "archivo": "DOC_1.jpg", "n_documento": "12345"}
    registro = _registro_con([acc._hallazgo("ITEM_AGRUPADO", dato,
                                            valor_actual="Materiales varios")])
    acc.anotar_n_ref_desde_master(registro, ws_m, [dato], ws_detalle=ws_d)

    assert acc.cerrar_hallazgos_ya_corregidos(registro, ws_m, correcciones=[],
                                              ws_detalle=ws_d) == []
    assert registro["errores"][0]["estado"] == "abierto"


def test_una_correccion_de_detalle_no_cierra_un_hallazgo_de_master():
    """Las columnas de Detalle y Master numeran cosas distintas: la 9 es
    'P. Unitario' en Detalle y 'Categoria' en Master. Sin filtrar por hoja,
    corregir el precio de un item cerraria el hallazgo de categoria vacia del
    mismo documento -- 28 de las correcciones reales calzaban asi."""
    ws_m, ws_d = _libro_cuadre()
    dato = {"proyecto": "Alfa", "archivo": "DOC_1.jpg", "n_documento": "12345"}
    registro = _registro_con([acc._hallazgo("CATEGORIA_VACIA", dato)])
    acc.anotar_n_ref_desde_master(registro, ws_m, [dato], ws_detalle=ws_d)
    assert registro["errores"][0]["columna"] == 9

    correcciones = [{"n_ref": "BALF-001", "hoja": "Detalle", "columna": 9,
                     "campo": "Item 'X' (cantidad x precio unitario)",
                     "valor_anterior": "1 x 100", "valor_corregido": "1 x 200"}]
    assert acc.cerrar_hallazgos_ya_corregidos(registro, ws_m, correcciones=correcciones,
                                              ws_detalle=ws_d) == []
    assert registro["errores"][0]["estado"] == "abierto"


def test_una_correccion_de_master_si_cierra_su_hallazgo():
    ws_m, ws_d = _libro_cuadre()
    dato = {"proyecto": "Alfa", "archivo": "DOC_1.jpg", "n_documento": "12345"}
    registro = _registro_con([acc._hallazgo("CATEGORIA_VACIA", dato)])
    acc.anotar_n_ref_desde_master(registro, ws_m, [dato], ws_detalle=ws_d)

    correcciones = [{"n_ref": "BALF-001", "hoja": "Master", "columna": 9,
                     "campo": "Categoría", "valor_anterior": "", "valor_corregido": "Materiales"}]
    assert len(acc.cerrar_hallazgos_ya_corregidos(registro, ws_m, correcciones=correcciones,
                                                  ws_detalle=ws_d, hoy="2026-09-11")) == 1
