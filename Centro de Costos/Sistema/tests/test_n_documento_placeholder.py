"""Los marcadores de "sin numero" no son numeros de documento.

172 de las 660 entradas de datos_extraidos.json tienen n_documento 'N/A'
(peajes: Autopista de Los Andes, Concesionaria Litoral...) y 34 tienen
'S/N (archivo)'. El detector de duplicados comparaba el string crudo, asi que
cada peaje quedaba marcado como "posible duplicado" del peaje anterior: la
corrida del 2026-09-09 emitio 48 avisos y 47 eran peajes distintos.
"""

import openpyxl

import auditor_centro_costos as acc


# ── es_n_documento_real ────────────────────────────────────────────────────

def test_los_placeholders_no_son_numeros_reales():
    for valor in ("N/A", "n/a", "NA", "", "  ", None, "sin numero", "S/N (IMG_7533)", "S/N"):
        assert acc.es_n_documento_real(valor) is False, valor


def test_un_numero_de_verdad_si_lo_es():
    for valor in ("445566", 445566, "0000130020", "A-123"):
        assert acc.es_n_documento_real(valor) is True, valor


# ── leer_master no mete placeholders en docs_registrados ───────────────────

def _master(numeros):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master"
    for c, h in enumerate(acc.ENCABEZADOS_MASTER, 1):
        ws.cell(row=1, column=c, value=h)
    for i, n_doc in enumerate(numeros, start=2):
        ws.cell(row=i, column=1, value=f"UMAG-{i - 1:03d}")
        ws.cell(row=i, column=5, value=n_doc)
    return ws


def test_leer_master_ignora_placeholders_al_juntar_documentos_registrados():
    ws = _master(["N/A", "N/A", "S/N (IMG_1.jpg)", "445566"])

    _, _, docs = acc.leer_master(ws)

    assert "445566" in docs
    assert "N/A" not in docs
    assert "S/N (IMG_1.jpg)" not in docs


def test_leer_master_sigue_registrando_la_forma_sin_ceros():
    ws = _master(["0000130020"])

    _, _, docs = acc.leer_master(ws)

    assert {"0000130020", "130020"} <= docs
