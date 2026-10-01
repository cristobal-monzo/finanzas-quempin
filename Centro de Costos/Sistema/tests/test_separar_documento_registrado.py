# -*- coding: utf-8 -*-
"""
Separar documentos que quedaron registrados juntos en un solo N Ref, y
evitar que vuelva a pasar.

El libro llego a tener 26 N Ref que juntaban 2-4 facturas/boletas porque el
extractor las agrupaba "porque comparten el mismo archivo": impuesto sumado,
la fecha de una sola y el proveedor de todas en una celda. Se corrigio a mano
tres veces antes de que existiera esto. Los datos de aca son ficticios: este
repositorio es publico.
"""

import json

import openpyxl
import pytest

import auditor_centro_costos as acc


# ── Deteccion ────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("valor, esperado", [
    ("123456 y 0007891", ["123456", "7891"]),
    ("111222, 333444 y 555666", ["111222", "333444", "555666"]),
    ("123456/654321", ["123456", "654321"]),
    ("123456", []),
    ("N/A", []),
    ("S/N (Documento (18).pdf)", []),
    ("S/N (peajes Documento 10)", []),
    ("123456 y 123456", []),  # el mismo numero dos veces no son dos documentos
    (None, []),
])
def test_numeros_documento_combinados(valor, esperado):
    assert acc.numeros_documento_combinados(valor) == esperado


def _dato(**kwargs):
    base = {
        "archivo": "a.pdf", "proyecto": "Proyecto Test", "fecha": "15-07-2026",
        "n_documento": "5551", "tipo_documento": "Factura", "proveedor": "Prov SpA",
        "categoria": "Materiales", "iva": 1900,
        "items": [{"nombre_item": "Perno", "cantidad": 1, "p_unitario_sin_iva": 10000}],
    }
    base.update(kwargs)
    return base


def test_validar_marca_una_entrada_con_varios_documentos():
    codigos = [h["codigo"] for h in acc.validar_documento(_dato(n_documento="5551 y 7002"))]
    assert "DOCUMENTOS_MEZCLADOS" in codigos


def test_validar_no_marca_documentos_simples_ni_marcadores():
    for n_doc in ("5551", "N/A", "S/N (a.pdf)"):
        codigos = [h["codigo"] for h in acc.validar_documento(_dato(n_documento=n_doc))]
        assert "DOCUMENTOS_MEZCLADOS" not in codigos


# ── Sandbox con un pais ficticio (ver test_pipeline_errores.py) ─────────────

@pytest.fixture
def sandbox(tmp_path, monkeypatch):
    docs = tmp_path / "Facturas" / "Proyecto Test"
    docs.mkdir(parents=True)
    (tmp_path / "Excel").mkdir()
    (tmp_path / "Sitio").mkdir()
    cfg = dict(acc.PAISES["CL"])
    cfg.update({
        "ruta_docs": tmp_path / "Facturas",
        "ruta_excel": tmp_path / "Excel" / "Centro de Costos.xlsx",
        "ruta_backups": tmp_path / "Excel" / "Respaldos",
        "ruta_json": tmp_path / "datos_extraidos.json",
        "ruta_reconciliacion": tmp_path / "reconciliacion_archivos.json",
        "ruta_correcciones": tmp_path / "correcciones_manuales.json",
        "ruta_errores_md": tmp_path / "ERRORES.md",
        "ruta_logs": tmp_path / "logs",
        "ruta_visualizador_web": tmp_path / "Visualizador Web (inexistente)",
        "ruta_excel_sitio_comunicacion": tmp_path / "Sitio" / "Centro de Costos.xlsx",
        "prefijos_proyecto": {"Proyecto Test": "PTST"},
    })
    monkeypatch.setitem(acc.PAISES, "TEST", cfg)
    monkeypatch.setattr(acc, "RAIZ_ANALISIS_FINANCIERO", tmp_path / "AF (inexistente)")
    (tmp_path / "correcciones_manuales.json").write_text("[]", encoding="utf-8")
    return {"tmp": tmp_path, "docs": docs}


def _registrar(sandbox, documentos):
    for d in documentos:
        (sandbox["docs"] / d["archivo"]).write_bytes(b"pdf " + d["archivo"].encode())
    (sandbox["tmp"] / "datos_extraidos.json").write_text(
        json.dumps(documentos, ensure_ascii=False), encoding="utf-8")
    acc.main(pais="TEST")


def _combinado(**kwargs):
    """Lo que dejaba el extractor: dos boletas de dos emisores en una entrada.
    Se registra con un solo numero para poder llegar al libro (con dos
    numeros 'run' ya no lo escribe) -- es el caso de los peajes sin folio."""
    base = {
        "archivo": "Documento (1).pdf", "proyecto": "Proyecto Test", "tipo_proyecto": "Obra",
        "fecha": "10-07-2026", "n_documento": "5551", "tipo_documento": "Factura",
        "proveedor": "Ferreteria Uno SpA / Casino Dos SpA", "categoria": "Materiales",
        "estado": "Pagado", "iva": 3800 + 1900,
        "items": [
            {"nombre_item": "Perno", "descripcion": "Perno 1/2", "categoria_item": "Ferreteria",
             "cantidad": 2, "p_unitario_sin_iva": 10000},
            {"nombre_item": "Almuerzo", "descripcion": "Almuerzo x3 (dos boletas)",
             "categoria_item": "Alimentacion", "cantidad": 3, "p_unitario_sin_iva": 5000},
        ],
    }
    base.update(kwargs)
    return base


PEDIDO = {"PTST-001": [
    {"n_documento": "5551", "fecha": "10-07-2026", "tipo_documento": "Factura",
     "proveedor": "Ferreteria Uno SpA", "categoria": "Ferreteria", "iva": 3800, "items": [0]},
    {"n_documento": "7002", "fecha": "12-07-2026", "tipo_documento": "Boleta",
     "proveedor": "Casino Dos SpA", "categoria": "Alimentacion", "iva": 950,
     "items": [{"fila": 1, "cantidad": 1, "descripcion": "Almuerzo"}]},
    {"n_documento": "7003", "fecha": "13-07-2026", "tipo_documento": "Boleta",
     "proveedor": "Casino Dos SpA", "categoria": "Alimentacion", "iva": 1900,
     "items": [{"fila": 1, "cantidad": 2}], "nota": "Boleta impresa: neto 10.000, IVA 1.900"},
]}


def _libro(sandbox):
    return openpyxl.load_workbook(str(sandbox["tmp"] / "Excel" / "Centro de Costos.xlsx"))


def _filas_master(wb):
    ws = wb["Master"]
    return {ws.cell(row=r, column=1).value: r for r in range(2, acc.ultima_fila_datos(ws) + 1)}


def _items(wb, n_ref):
    ws = wb["Detalle"]
    return [[ws.cell(row=r, column=c).value for c in range(1, 12)]
            for r in range(2, acc.ultima_fila_datos(ws) + 1)
            if ws.cell(row=r, column=1).value == n_ref]


def test_run_no_registra_una_entrada_con_varios_documentos(sandbox, capsys):
    _registrar(sandbox, [_combinado(n_documento="5551 y 7002")])

    assert _filas_master(_libro(sandbox)) == {}
    salida = capsys.readouterr().out
    assert "No se registra (junta varios documentos)" in salida


def test_preview_no_escribe_nada(sandbox):
    _registrar(sandbox, [_combinado()])
    antes = (sandbox["tmp"] / "Excel" / "Centro de Costos.xlsx").read_bytes()
    archivos_antes = sorted(p.name for p in sandbox["docs"].iterdir())

    planes, resultados = acc.ejecutar_separacion_registrados(PEDIDO, aplicar=False)

    assert resultados == []
    assert [d["n_ref"] for d in planes[0]["documentos"]] == ["PTST-001", "PTST-002", "PTST-003"]
    assert (sandbox["tmp"] / "Excel" / "Centro de Costos.xlsx").read_bytes() == antes
    assert sorted(p.name for p in sandbox["docs"].iterdir()) == archivos_antes


def test_separa_en_un_n_ref_por_documento(sandbox):
    _registrar(sandbox, [_combinado()])

    acc.ejecutar_separacion_registrados(PEDIDO, aplicar=True)

    wb = _libro(sandbox)
    ws = wb["Master"]
    filas = _filas_master(wb)
    assert set(filas) == {"PTST-001", "PTST-002", "PTST-003"}
    fila = filas["PTST-002"]
    assert ws.cell(row=fila, column=5).value == "7002"
    assert ws.cell(row=fila, column=6).value == "Boleta"
    assert ws.cell(row=fila, column=8).value == "Casino Dos SpA"
    assert ws.cell(row=fila, column=12).value == 950
    assert ws.cell(row=fila, column=4).value.strftime("%d-%m-%Y") == "12-07-2026"
    assert ws.cell(row=fila, column=11).value.startswith("=SUMIF(")
    assert ws.cell(row=filas["PTST-001"], column=8).value == "Ferreteria Uno SpA"
    assert ws.cell(row=filas["PTST-001"], column=12).value == 3800


def test_los_items_se_mueven_y_la_fila_repartida_suma_la_original(sandbox):
    _registrar(sandbox, [_combinado()])

    acc.ejecutar_separacion_registrados(PEDIDO, aplicar=True)

    wb = _libro(sandbox)
    uno, dos, tres = _items(wb, "PTST-001"), _items(wb, "PTST-002"), _items(wb, "PTST-003")
    assert [(i[4], i[7], i[9]) for i in uno] == [("Perno", 2, 20000)]
    assert [(i[4], i[5], i[7], i[9]) for i in dos] == [("Almuerzo", "Almuerzo", 1, 5000)]
    assert [(i[7], i[9]) for i in tres] == [(2, 10000)]
    assert {i[3] for i in dos} == {"7002"}
    # Neto total intacto: separar no saca ni agrega costo.
    assert sum(i[9] for i in uno + dos + tres) == 2 * 10000 + 3 * 5000
    # Total con IVA con la tasa de CADA documento, no la del combinado.
    assert uno[0][10] == 23800 and dos[0][10] == 5950 and tres[0][10] == 11900


def test_deja_una_copia_del_archivo_por_documento(sandbox):
    _registrar(sandbox, [_combinado()])

    acc.ejecutar_separacion_registrados(PEDIDO, aplicar=True)

    wb = _libro(sandbox)
    ws = wb["Master"]
    archivos = {n_ref: ws.cell(row=f, column=15).value for n_ref, f in _filas_master(wb).items()}
    assert len(set(archivos.values())) == 3
    for n_ref, ruta in archivos.items():
        nombre = ruta.split("\\", 1)[1]
        assert nombre.startswith(f"{n_ref}_")
        assert (sandbox["docs"] / nombre).exists()
    assert len(list(sandbox["docs"].iterdir())) == 3


def test_la_corrida_siguiente_no_ve_pendientes_ni_correcciones_ni_renombra(sandbox, capsys):
    _registrar(sandbox, [_combinado()])
    acc.ejecutar_separacion_registrados(PEDIDO, aplicar=True)
    capsys.readouterr()

    acc.main(pais="TEST")

    salida = capsys.readouterr().out
    assert "Pendientes: 0" in salida
    assert "Sin correcciones manuales nuevas." in salida
    assert "[OK] 0 archivo(s) renombrado(s)/convertido(s)." in salida
    assert set(_filas_master(_libro(sandbox))) == {"PTST-001", "PTST-002", "PTST-003"}


def test_queda_en_la_bitacora(sandbox):
    _registrar(sandbox, [_combinado()])

    acc.ejecutar_separacion_registrados(PEDIDO, aplicar=True)

    bitacora = json.loads((sandbox["tmp"] / "correcciones_manuales.json").read_text(encoding="utf-8"))
    resumen = [c for c in bitacora if c["columna"] == 0]
    assert resumen and resumen[0]["n_ref"] == "PTST-001/PTST-002/PTST-003"
    iva_tres = [c for c in bitacora if c["n_ref"] == "PTST-003" and c["columna"] == 12]
    assert iva_tres[0]["valor_corregido"] == 1900
    assert iva_tres[0]["nota"] == "Boleta impresa: neto 10.000, IVA 1.900"


@pytest.mark.parametrize("documentos, mensaje", [
    ([{**PEDIDO["PTST-001"][0], "items": [0]}, {**PEDIDO["PTST-001"][1], "items": [0]}],
     "asigno completa"),
    ([PEDIDO["PTST-001"][0], PEDIDO["PTST-001"][1]], "no suman la original"),
    ([PEDIDO["PTST-001"][0], {**PEDIDO["PTST-001"][1], "items": []}], "faltan items"),
    ([{**PEDIDO["PTST-001"][0], "items": [0, 1]}], "2 o mas"),
    ([{**PEDIDO["PTST-001"][0], "fecha": "2026-07-10"}, PEDIDO["PTST-001"][1],
      PEDIDO["PTST-001"][2]], "DD-MM-AAAA"),
    ([PEDIDO["PTST-001"][0], {**PEDIDO["PTST-001"][1], "items": [5]}], "no existe"),
])
def test_rechaza_pedidos_que_no_cuadran(sandbox, documentos, mensaje):
    _registrar(sandbox, [_combinado()])
    antes = (sandbox["tmp"] / "Excel" / "Centro de Costos.xlsx").read_bytes()

    with pytest.raises(ValueError, match=mensaje):
        acc.ejecutar_separacion_registrados({"PTST-001": documentos}, aplicar=True)

    assert (sandbox["tmp"] / "Excel" / "Centro de Costos.xlsx").read_bytes() == antes


def test_con_el_excel_bloqueado_no_toca_archivos(sandbox, monkeypatch):
    _registrar(sandbox, [_combinado()])
    archivos_antes = sorted(p.name for p in sandbox["docs"].iterdir())
    monkeypatch.setattr(acc, "excel_esta_bloqueado", lambda ruta: True)

    with pytest.raises(PermissionError):
        acc.ejecutar_separacion_registrados(PEDIDO, aplicar=True)

    assert sorted(p.name for p in sandbox["docs"].iterdir()) == archivos_antes


# ── Cierre del hallazgo con evidencia del libro ──────────────────────────────

def test_el_hallazgo_se_cierra_cuando_cada_numero_tiene_su_fila(sandbox):
    """La entrada combinada del JSON no desaparece al separar (el JSON no se
    reescribe): la prueba de que se resolvio tiene que venir del libro."""
    _registrar(sandbox, [_combinado()])
    acc.ejecutar_separacion_registrados(PEDIDO, aplicar=True)
    wb = _libro(sandbox)
    registro = {"errores": []}
    hallazgos = [h for h in acc.validar_documento(_combinado(n_documento="5551, 7002 y 7003"))
                 if h["codigo"] == "DOCUMENTOS_MEZCLADOS"]
    acc.fusionar_hallazgos(registro, hallazgos)

    cerrados = acc.cerrar_hallazgos_ya_corregidos(registro, wb["Master"], correcciones=[])

    assert [e["codigo"] for e in cerrados] == ["DOCUMENTOS_MEZCLADOS"]
    assert "PTST-002 (7002)" in registro["errores"][0]["resolucion"]


def test_el_hallazgo_sigue_abierto_si_falta_un_numero(sandbox):
    _registrar(sandbox, [_combinado()])
    acc.ejecutar_separacion_registrados(PEDIDO, aplicar=True)
    wb = _libro(sandbox)
    registro = {"errores": []}
    hallazgos = [h for h in acc.validar_documento(_combinado(n_documento="5551 y 9999"))
                 if h["codigo"] == "DOCUMENTOS_MEZCLADOS"]
    acc.fusionar_hallazgos(registro, hallazgos)

    assert acc.cerrar_hallazgos_ya_corregidos(registro, wb["Master"], correcciones=[]) == []
