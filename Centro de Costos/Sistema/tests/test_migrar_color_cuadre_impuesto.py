"""migrar_color_cuadre_impuesto: repinta la columna de impuesto de Master.

El color se escribe una sola vez, cuando la fila se crea, y las filas de
datos no se vuelven a tocar -- asi que cambiar la regla de pintado no arregla
por si solo lo ya escrito. El libro real quedo con 55 celdas de IVA en rojo
bajo la regla vieja ("cualquier desvio del 19%"), casi todas facturas de
combustible correctas. Esa es la lista que recorre /Revision_de_Errores.
"""

import openpyxl

import auditor_centro_costos as acc


def _libro(filas):
    """filas: [(n_ref, tipo, categoria, iva, neto, roja_inicial)]"""
    wb = openpyxl.Workbook()
    ws_m = wb.active
    ws_m.title = "Master"
    for c, h in enumerate(acc.ENCABEZADOS_MASTER, 1):
        ws_m.cell(row=1, column=c, value=h)
    ws_d = wb.create_sheet("Detalle")
    for c, h in enumerate(acc.ENCABEZADOS_DETALLE, 1):
        ws_d.cell(row=1, column=c, value=h)

    cm = {h: i + 1 for i, h in enumerate(acc.ENCABEZADOS_MASTER)}
    cd = {h: i + 1 for i, h in enumerate(acc.ENCABEZADOS_DETALLE)}
    col_total_d = cd[f"Total sin {acc.NOMBRE_IMPUESTO_CORTO} ({acc.MONEDA})"]
    col_iva_m = cm[f"{acc.NOMBRE_IMPUESTO_PCT} ({acc.MONEDA})"]

    for i, (n_ref, tipo, categoria, iva, neto, roja) in enumerate(filas, start=2):
        ws_m.cell(row=i, column=cm["N° Ref."], value=n_ref)
        ws_m.cell(row=i, column=cm["Tipo Documento"], value=tipo)
        ws_m.cell(row=i, column=cm["Categoría"], value=categoria)
        celda = ws_m.cell(row=i, column=col_iva_m, value=iva)
        celda.font = acc.ROJO_FONT if roja else acc.NORMAL_FONT
        ws_d.cell(row=i, column=cd["N° Ref."], value=n_ref)
        ws_d.cell(row=i, column=col_total_d, value=neto)
    return ws_m, ws_d, col_iva_m


def test_libera_el_rojo_de_una_factura_de_combustible_correcta():
    ws_m, ws_d, col = _libro([("UMAG-1", "Factura", "Combustible", 10661, 19339, True)])

    limpiadas, marcadas = acc.migrar_color_cuadre_impuesto(ws_m, ws_d)

    assert (limpiadas, marcadas) == (1, 0)
    assert not acc._celda_es_roja(ws_m.cell(row=2, column=col))


def test_libera_el_rojo_de_un_deficit_de_combustible():
    """Actualizado el 2026-09-10: antes este caso mantenia el rojo, porque la
    regla daba 'error' a cualquier impuesto bajo el 19%. En combustible el
    FEPP/IEV puede ser negativo y dejarlo ahi legitimamente (JUNJ-238), asi
    que pasa a 'revisar' y la celda deja de pedir a gritos una correccion que
    no corresponde. El hallazgo sigue en el registro de errores."""
    ws_m, ws_d, col = _libro([("UMAG-1", "Factura", "Combustible", 1873, 38318, True)])

    limpiadas, marcadas = acc.migrar_color_cuadre_impuesto(ws_m, ws_d)

    assert (limpiadas, marcadas) == (1, 0)
    assert not acc._celda_es_roja(ws_m.cell(row=2, column=col))


def test_mantiene_el_rojo_en_un_deficit_sin_impuesto_especifico():
    """El contrapeso: fuera de combustible, un impuesto bajo el 19% sigue
    siendo algo seguro que corregir y la celda sigue roja."""
    ws_m, ws_d, col = _libro([("UMAG-1", "Factura", "Ferreteria", 1873, 38318, True)])

    limpiadas, marcadas = acc.migrar_color_cuadre_impuesto(ws_m, ws_d)

    assert (limpiadas, marcadas) == (0, 0)
    assert acc._celda_es_roja(ws_m.cell(row=2, column=col))


def test_marca_en_rojo_un_error_que_habia_quedado_sin_marcar():
    ws_m, ws_d, col = _libro([("UMAG-1", "Factura", "Ferreteria", 100, 10000, False)])

    limpiadas, marcadas = acc.migrar_color_cuadre_impuesto(ws_m, ws_d)

    assert (limpiadas, marcadas) == (0, 1)
    assert acc._celda_es_roja(ws_m.cell(row=2, column=col))


def test_guia_de_despacho_sin_tilde_tambien_se_evalua():
    """Antes el tipo se comparaba contra el literal con tilde, asi que estas
    filas quedaban fuera de la regla."""
    ws_m, ws_d, col = _libro([("UMAG-1", "Guia de Despacho", "Materiales", 100, 10000, False)])

    _, marcadas = acc.migrar_color_cuadre_impuesto(ws_m, ws_d)

    assert marcadas == 1
    assert acc._celda_es_roja(ws_m.cell(row=2, column=col))


def test_suma_varios_items_de_detalle_para_el_neto():
    ws_m, ws_d, col = _libro([("UMAG-1", "Factura", "Ferreteria", 19000, 60000, True)])
    cd = {h: i + 1 for i, h in enumerate(acc.ENCABEZADOS_DETALLE)}
    col_total_d = cd[f"Total sin {acc.NOMBRE_IMPUESTO_CORTO} ({acc.MONEDA})"]
    ws_d.cell(row=3, column=cd["N° Ref."], value="UMAG-1")
    ws_d.cell(row=3, column=col_total_d, value=40000)   # neto total = 100000

    limpiadas, _ = acc.migrar_color_cuadre_impuesto(ws_m, ws_d)

    assert limpiadas == 1  # 19% de 100000 = 19000, cuadra
    assert not acc._celda_es_roja(ws_m.cell(row=2, column=col))


def test_es_idempotente():
    ws_m, ws_d, _ = _libro([
        ("UMAG-1", "Factura", "Combustible", 10661, 19339, True),
        ("UMAG-2", "Factura", "Ferreteria", 100, 10000, False),
    ])

    primera = acc.migrar_color_cuadre_impuesto(ws_m, ws_d)
    segunda = acc.migrar_color_cuadre_impuesto(ws_m, ws_d)

    assert primera == (1, 1)
    assert segunda == (0, 0)


def test_ignora_filas_sin_neto_en_detalle():
    ws_m, ws_d, col = _libro([("UMAG-1", "Factura", "Ferreteria", 19000, 100000, True)])
    ws_d.delete_rows(2)   # queda Master sin su contraparte en Detalle

    assert acc.migrar_color_cuadre_impuesto(ws_m, ws_d) == (0, 0)
    assert acc._celda_es_roja(ws_m.cell(row=2, column=col))


def test_respeta_una_celda_ya_corregida_a_mano():
    """Azul marino = el valor lo adjudico una persona mirando el documento
    (ver CLAUDE.md del modulo). Repintarlo de rojo borra esa marca y vuelve a
    pedir una correccion que ya se hizo.

    Caso real que lo destapo (2026-09-14): UMAG-005 y UMAG-020 son facturas
    de Zona Franca de Punta Arenas, exentas de IVA. Se corrigieron a mano el
    2026-07-17 y quedaron azules; una corrida posterior las volvio a pintar
    de rojo porque su impuesto 0 nunca va a ser el 19% del neto, y en la
    revision del 2026-09-14 reaparecieron como si nadie las hubiera visto.
    """
    ws_m, ws_d, col = _libro([("UMAG-1", "Factura", "Ferreteria", 0, 20800, False)])
    ws_m.cell(row=2, column=col).font = acc.AZUL_MARINO_FONT

    limpiadas, marcadas = acc.migrar_color_cuadre_impuesto(ws_m, ws_d)

    assert (limpiadas, marcadas) == (0, 0)
    assert acc._celda_es_azul_marino(ws_m.cell(row=2, column=col))
    assert not acc._celda_es_roja(ws_m.cell(row=2, column=col))
