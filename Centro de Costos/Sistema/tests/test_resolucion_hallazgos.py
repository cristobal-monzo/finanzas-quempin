# -*- coding: utf-8 -*-
"""
Tests del canal unico de resolucion de hallazgos (corregir_hallazgos /
descartar_hallazgo).

Antes de la auditoria del 2026-09-10 el unico camino auditado para corregir un
dato eran las dos columnas que el script pinta de rojo. Estos tests cubren:
  * que ahora se puede resolver cualquier hallazgo cuya clase declare una
    columna de Master (fecha, tipo de documento, proveedor, categoria...);
  * que el lote usa UNA sola apertura/respaldo/guardado del libro;
  * que la puerta de entrada sigue siendo "hay un hallazgo abierto que dice que
    esta celda esta mal", no "escribeme cualquier celda";
  * que la correccion queda en las tres bitacoras (Excel en azul marino,
    correcciones_manuales.json y el registro de errores).
"""

from datetime import datetime

import openpyxl
import pytest

import auditor_centro_costos as acc


ERRORES_MD = """# Errores

## Correcciones manuales pendientes de recolorear

| Fecha | Hoja | N° Ref. | Campo / Columna | Valor anterior (rojo) | Valor corregido | Estado | Nota |
|---|---|---|---|---|---|---|---|
| *(sin entradas todavía)* | | | | | | | |

## Fin
"""


@pytest.fixture
def libro(tmp_path):
    """Un libro minimo con una fila en Master y sus dos items en Detalle, mas
    las rutas del resto de las bitacoras."""
    wb = openpyxl.Workbook()
    ws_m = wb.active
    ws_m.title = "Master"
    for c, h in enumerate(acc.ENCABEZADOS_MASTER, 1):
        ws_m.cell(row=1, column=c, value=h)
    valores = {1: "BALF-001", 2: "Proyecto Alfa", 4: datetime(2026, 7, 15),
               5: "12345", 6: "Factura", 7: "Demo", 8: "Proveedor Demo SpA",
               9: "Materiales", 12: 19000, 15: "Proyecto Alfa\\DOC_0001.jpg"}
    for c, v in valores.items():
        ws_m.cell(row=2, column=c, value=v)

    ws_d = wb.create_sheet("Detalle")
    for c, h in enumerate(acc.ENCABEZADOS_DETALLE, 1):
        ws_d.cell(row=1, column=c, value=h)
    for fila in (2, 3):
        ws_d.cell(row=fila, column=1, value="BALF-001")
        ws_d.cell(row=fila, column=4, value="12345")

    ruta_excel = tmp_path / "Centro de Costos.xlsx"
    wb.save(str(ruta_excel))

    ruta_correcciones = tmp_path / "correcciones_manuales.json"
    ruta_correcciones.write_text("[]", encoding="utf-8")
    ruta_errores_md = tmp_path / "ERRORES.md"
    ruta_errores_md.write_text(ERRORES_MD, encoding="utf-8")

    return {
        "ruta_excel": ruta_excel,
        "ruta_correcciones": ruta_correcciones,
        "ruta_errores": ruta_errores_md,
        "ruta_backups": tmp_path / "Respaldos",
        "ruta_registro": tmp_path / "errores_detectados.json",
    }


def sembrar(libro, *hallazgos):
    """Deja hallazgos abiertos en el registro, ya con su N Ref pegado."""
    registro = {"version": 1, "errores": []}
    acc.fusionar_hallazgos(registro, list(hallazgos), hoy="2026-09-10")
    acc.anotar_n_ref(registro, "Proyecto Alfa\\DOC_0001.jpg", "BALF-001")
    acc.guardar_registro_errores(registro, libro["ruta_registro"])
    return registro


def documento(**kwargs):
    base = {
        "archivo": "DOC_0001.jpg", "proyecto": "Proyecto Alfa",
        "fecha": "15-07-2026", "n_documento": "12345", "tipo_documento": "Factura",
        "proveedor": "Proveedor Demo SpA", "rut_proveedor": "76123456-7",
        "categoria": "Materiales", "iva": 19000,
        "items": [{"nombre_item": "Perfil", "cantidad": 2, "p_unitario_sin_iva": 50000}],
    }
    base.update(kwargs)
    return base


def aplicar(libro, pedidos):
    return acc.corregir_hallazgos(
        pedidos, ruta_excel=libro["ruta_excel"],
        ruta_correcciones=libro["ruta_correcciones"], ruta_errores=libro["ruta_errores"],
        ruta_backups=libro["ruta_backups"], ruta_registro=libro["ruta_registro"])


def celda(libro, columna):
    wb = openpyxl.load_workbook(str(libro["ruta_excel"]))
    try:
        return wb["Master"].cell(row=2, column=columna)
    finally:
        wb.close()


# ── Columnas que antes no tenian ningun canal de correccion ────────────────

def test_categoria_vacia_se_puede_resolver():
    """Antes solo se podian corregir las columnas 5 y 12 (las que se pintan de
    rojo); una categoria en blanco no tenia camino auditado."""
    assert 9 in acc.COLUMNAS_CORREGIBLES


def test_corregir_categoria_escribe_valor_y_lo_deja_en_azul_marino(libro):
    h = [x for x in acc.validar_documento(documento(categoria=""))
         if x["codigo"] == "CATEGORIA_VACIA"][0]
    sembrar(libro, h)

    aplicadas, rechazadas = aplicar(libro, [(h["id"], "Ferreteria")])

    assert rechazadas == [] and len(aplicadas) == 1
    c = celda(libro, 9)
    assert c.value == "Ferreteria"
    assert c.font.color.rgb.endswith(acc.NAVY_OSCURO)


def test_corregir_fecha_guarda_datetime_con_formato_ddmmaaaa(libro):
    h = [x for x in acc.validar_documento(documento(fecha="32-13-2026"))
         if x["codigo"] == "FECHA_INVALIDA"][0]
    sembrar(libro, h)

    aplicar(libro, [(h["id"], "03-08-2026")])

    c = celda(libro, 4)
    assert c.value == datetime(2026, 8, 3)
    assert c.number_format == acc.DATE_FORMAT


def test_corregir_razon_social_regenera_el_tag_corto(libro):
    """El tag de la columna 7 es 100% derivado: corregir uno sin el otro deja
    Master mostrando el proveedor viejo con la razon social nueva."""
    h = [x for x in acc.validar_documento(documento(proveedor=""))
         if x["codigo"] == "PROVEEDOR_VACIO"][0]
    sembrar(libro, h)

    aplicar(libro, [(h["id"], "Comercial Nueva Limitada")])

    assert celda(libro, 8).value == "Comercial Nueva Limitada"
    assert celda(libro, 7).value == acc.generar_tag_proveedor("Comercial Nueva Limitada")


def test_corregir_n_documento_propaga_a_detalle(libro):
    h = [x for x in acc.validar_documento(documento(n_documento="S/N (DOC_0001.jpg)"))
         if x["codigo"] == "N_DOC_ILEGIBLE"][0]
    sembrar(libro, h)

    aplicar(libro, [(h["id"], "0000987")])

    wb = openpyxl.load_workbook(str(libro["ruta_excel"]))
    try:
        assert wb["Master"].cell(row=2, column=5).value == "987"  # sin ceros a la izquierda
        assert [wb["Detalle"].cell(row=f, column=4).value for f in (2, 3)] == ["987", "987"]
    finally:
        wb.close()


# ── El lote: una apertura, un respaldo, un guardado ────────────────────────

def test_lote_hace_un_solo_respaldo_para_varias_correcciones(libro):
    d = documento(categoria="", proveedor="")
    hs = {x["codigo"]: x for x in acc.validar_documento(d)}
    sembrar(libro, hs["CATEGORIA_VACIA"], hs["PROVEEDOR_VACIO"])

    aplicadas, rechazadas = aplicar(libro, [
        (hs["CATEGORIA_VACIA"]["id"], "Ferreteria"),
        (hs["PROVEEDOR_VACIO"]["id"], "Comercial Nueva Limitada"),
    ])

    assert len(aplicadas) == 2 and rechazadas == []
    assert len(list(libro["ruta_backups"].rglob("*.xlsx"))) == 1
    assert celda(libro, 9).value == "Ferreteria"
    assert celda(libro, 8).value == "Comercial Nueva Limitada"


# ── La puerta de entrada sigue cerrada para lo que no corresponde ──────────

def test_rechaza_id_inexistente(libro):
    sembrar(libro)
    aplicadas, rechazadas = aplicar(libro, [("noexiste123", "x")])
    assert aplicadas == [] and rechazadas[0][0] == "noexiste123"


def test_rechaza_hallazgo_ya_cerrado(libro):
    h = [x for x in acc.validar_documento(documento(categoria=""))
         if x["codigo"] == "CATEGORIA_VACIA"][0]
    sembrar(libro, h)
    aplicar(libro, [(h["id"], "Ferreteria")])

    aplicadas, rechazadas = aplicar(libro, [(h["id"], "Otra")])

    assert aplicadas == []
    assert "ya esta resuelto" in rechazadas[0][1]
    assert celda(libro, 9).value == "Ferreteria"  # no se piso


def test_rechaza_clase_que_no_se_arregla_escribiendo_una_celda(libro):
    h = [x for x in acc.validar_documento(documento(items=[]))
         if x["codigo"] == "DOC_SIN_ITEMS"][0]
    sembrar(libro, h)

    aplicadas, rechazadas = aplicar(libro, [(h["id"], "loquesea")])

    assert aplicadas == []
    assert "no se arregla escribiendo" in rechazadas[0][1]


def test_rechaza_hallazgo_sin_n_ref(libro):
    """Un documento que todavia no llego a Master no tiene celda que corregir."""
    h = [x for x in acc.validar_documento(documento(categoria=""))
         if x["codigo"] == "CATEGORIA_VACIA"][0]
    registro = {"version": 1, "errores": []}
    acc.fusionar_hallazgos(registro, [h], hoy="2026-09-10")  # sin anotar_n_ref
    acc.guardar_registro_errores(registro, libro["ruta_registro"])

    aplicadas, rechazadas = aplicar(libro, [(h["id"], "Ferreteria")])

    assert aplicadas == []
    assert "todavia no tiene fila en Master" in rechazadas[0][1]


def test_un_rechazo_no_impide_aplicar_el_resto_del_lote(libro):
    d = documento(categoria="")
    h = [x for x in acc.validar_documento(d) if x["codigo"] == "CATEGORIA_VACIA"][0]
    sembrar(libro, h)

    aplicadas, rechazadas = aplicar(libro, [("noexiste", "x"), (h["id"], "Ferreteria")])

    assert len(aplicadas) == 1 and len(rechazadas) == 1


# ── Trazabilidad: la correccion queda en las tres bitacoras ────────────────

def test_la_correccion_queda_en_correcciones_manuales_y_en_errores_md(libro):
    h = [x for x in acc.validar_documento(documento(categoria=""))
         if x["codigo"] == "CATEGORIA_VACIA"][0]
    sembrar(libro, h)

    aplicar(libro, [(h["id"], "Ferreteria", "confirmado con la foto")])

    correcciones = acc.cargar_correcciones_manuales(libro["ruta_correcciones"])
    assert len(correcciones) == 1
    assert correcciones[0]["estado"] == "Aplicado"
    assert correcciones[0]["valor_corregido"] == "Ferreteria"
    assert correcciones[0]["id_hallazgo"] == h["id"]
    assert correcciones[0]["nota"] == "confirmado con la foto"
    assert "Ferreteria" in libro["ruta_errores"].read_text(encoding="utf-8")


def test_la_nota_queda_como_comentario_en_la_celda(libro):
    h = [x for x in acc.validar_documento(documento(categoria=""))
         if x["codigo"] == "CATEGORIA_VACIA"][0]
    sembrar(libro, h)
    aplicar(libro, [(h["id"], "Ferreteria", "confirmado con la foto")])
    assert "confirmado con la foto" in celda(libro, 9).comment.text


def test_el_hallazgo_queda_resuelto_con_el_detalle_del_cambio(libro):
    h = [x for x in acc.validar_documento(documento(categoria=""))
         if x["codigo"] == "CATEGORIA_VACIA"][0]
    sembrar(libro, h)
    aplicar(libro, [(h["id"], "Ferreteria")])

    entrada = acc.cargar_registro_errores(libro["ruta_registro"])["errores"][0]
    assert entrada["estado"] == "resuelto"
    assert "Ferreteria" in entrada["resolucion"]
    assert entrada["fecha_cierre"]


# ── Reapertura: solo si el dato de origen cambia ───────────────────────────

def test_hallazgo_resuelto_no_reabre_si_el_origen_sigue_igual(libro):
    """datos_extraidos.json es entrada del pipeline y no se reescribe: si se
    reabriera en cada corrida, los mismos descuadres se reimprimirian para
    siempre (que es justo lo que pasaba antes)."""
    d = documento(categoria="")
    h = [x for x in acc.validar_documento(d) if x["codigo"] == "CATEGORIA_VACIA"][0]
    sembrar(libro, h)
    aplicar(libro, [(h["id"], "Ferreteria")])

    registro = acc.cargar_registro_errores(libro["ruta_registro"])
    _, reabiertos, _ = acc.fusionar_hallazgos(registro, [h], hoy="2026-09-11")

    assert reabiertos == []
    assert registro["errores"][0]["estado"] == "resuelto"
    assert registro["errores"][0]["origen_sin_corregir"] is True


def test_hallazgo_resuelto_reabre_si_el_origen_cambia_a_otro_valor(libro):
    d = documento(proveedor="")
    h = [x for x in acc.validar_documento(d) if x["codigo"] == "PROVEEDOR_VACIO"][0]
    sembrar(libro, h)
    aplicar(libro, [(h["id"], "Comercial Nueva Limitada")])

    registro = acc.cargar_registro_errores(libro["ruta_registro"])
    otro = dict(h, valor_actual="   ")  # el JSON ahora dice otra cosa, igual de mala
    _, reabiertos, _ = acc.fusionar_hallazgos(registro, [otro], hoy="2026-09-11")

    assert len(reabiertos) == 1
    assert registro["errores"][0]["estado"] == "abierto"


# ── Descartar es una decision documentada, no un borrado ───────────────────

def test_descartar_deja_constancia_del_motivo(libro):
    h = [x for x in acc.validar_documento(documento(categoria=""))
         if x["codigo"] == "CATEGORIA_VACIA"][0]
    sembrar(libro, h)

    entrada = acc.descartar_hallazgo(
        h["id"], "el documento no trae categoria y asi corresponde",
        ruta_registro=libro["ruta_registro"])

    assert entrada["estado"] == "descartado"
    guardado = acc.cargar_registro_errores(libro["ruta_registro"])["errores"][0]
    assert guardado["estado"] == "descartado"
    assert "asi corresponde" in guardado["resolucion"]


def test_descartar_sin_motivo_falla(libro):
    h = [x for x in acc.validar_documento(documento(categoria=""))
         if x["codigo"] == "CATEGORIA_VACIA"][0]
    sembrar(libro, h)
    with pytest.raises(ValueError):
        acc.descartar_hallazgo(h["id"], "", ruta_registro=libro["ruta_registro"])


def test_descartar_id_inexistente_devuelve_none(libro):
    sembrar(libro)
    assert acc.descartar_hallazgo("nada", "motivo",
                                  ruta_registro=libro["ruta_registro"]) is None


# ── Coercion de tipos por columna ──────────────────────────────────────────

@pytest.mark.parametrize("columna,entrada,esperado", [
    (5, "0000130020", "130020"),
    (12, "8799", 8799),
    (12, "8799.5", 8799.5),
    (12, "no-numerico", "no-numerico"),
    (4, "03-08-2026", datetime(2026, 8, 3)),
    (4, "no-fecha", "no-fecha"),
    (9, "Ferreteria", "Ferreteria"),
])
def test_coaccionar_valor_columna(columna, entrada, esperado):
    assert acc.coaccionar_valor_columna(columna, entrada) == esperado
