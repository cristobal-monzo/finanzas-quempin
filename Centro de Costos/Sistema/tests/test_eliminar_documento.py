"""eliminar_documento: sacar del libro un documento registrado por error.

Es la unica operacion del modulo que borra una fila de datos ya escrita. La
motiva un caso encontrado en la auditoria del 2026-09-09: un mismo documento
tributario quedo registrado dos veces, desde dos escaneos distintos del mismo
original, sumando su costo dos veces al proyecto. Los identificadores del
caso real (proveedor, folio, monto) no van aca: este repositorio es publico.
"""

import openpyxl
import pytest

import auditor_centro_costos as acc


def _libro(tmp_path, documentos):
    """documentos: [(n_ref, proyecto, archivo_origen, [(nombre_item, total)])]"""
    wb = openpyxl.Workbook()
    ws_m = wb.active
    ws_m.title = "Master"
    for c, h in enumerate(acc.ENCABEZADOS_MASTER, 1):
        ws_m.cell(row=1, column=c, value=h)
    ws_d = wb.create_sheet("Detalle")
    for c, h in enumerate(acc.ENCABEZADOS_DETALLE, 1):
        ws_d.cell(row=1, column=c, value=h)

    fila_d = 2
    for fila_m, (n_ref, proyecto, archivo, items) in enumerate(documentos, start=2):
        ws_m.cell(row=fila_m, column=1, value=n_ref)
        ws_m.cell(row=fila_m, column=2, value=proyecto)
        ws_m.cell(row=fila_m, column=15, value=archivo)
        for nombre, total in items:
            ws_d.cell(row=fila_d, column=1, value=n_ref)
            ws_d.cell(row=fila_d, column=5, value=nombre)
            ws_d.cell(row=fila_d, column=10, value=total)
            fila_d += 1
    return ws_m, ws_d


DOCS = [
    ("PROY-001", "Proyecto Ejemplo", "Proyecto Ejemplo\\a.pdf", [("Item duplicado", 500000)]),
    ("PROY-002", "Proyecto Ejemplo", "Proyecto Ejemplo\\b.pdf", [("Item duplicado", 500000)]),
    ("UMAG-001", "UMAG", "UMAG\\c.pdf", [("Perno", 100), ("Tuerca", 50)]),
]


def test_borra_la_fila_de_master_y_sus_items_de_detalle(tmp_path):
    ws_m, ws_d = _libro(tmp_path, DOCS)

    resumen = acc.eliminar_documento("PROY-002", ws_m, ws_d, raiz_docs=tmp_path,
                                     ruta_backups=tmp_path / "resp")

    assert resumen["items_borrados"] == 1
    refs = {ws_m.cell(row=r, column=1).value for r in range(2, acc.ultima_fila_datos(ws_m) + 1)}
    assert refs == {"PROY-001", "UMAG-001"}
    refs_d = {ws_d.cell(row=r, column=1).value for r in range(2, acc.ultima_fila_datos(ws_d) + 1)}
    assert "PROY-002" not in refs_d


def test_no_toca_los_demas_documentos(tmp_path):
    ws_m, ws_d = _libro(tmp_path, DOCS)

    acc.eliminar_documento("PROY-002", ws_m, ws_d, raiz_docs=tmp_path,
                           ruta_backups=tmp_path / "resp")

    items_umag = [ws_d.cell(row=r, column=5).value
                  for r in range(2, acc.ultima_fila_datos(ws_d) + 1)
                  if ws_d.cell(row=r, column=1).value == "UMAG-001"]
    assert items_umag == ["Perno", "Tuerca"]


def test_borra_todos_los_items_de_un_documento_con_varios(tmp_path):
    ws_m, ws_d = _libro(tmp_path, DOCS)

    resumen = acc.eliminar_documento("UMAG-001", ws_m, ws_d, raiz_docs=tmp_path,
                                     ruta_backups=tmp_path / "resp")

    assert resumen["items_borrados"] == 2
    assert all(ws_d.cell(row=r, column=1).value != "UMAG-001"
               for r in range(2, acc.ultima_fila_datos(ws_d) + 1))


def test_archiva_el_archivo_fuente_y_lo_saca_de_la_carpeta(tmp_path):
    """Si el archivo se queda, el proximo 'run' lo ve como pendiente y repone
    el duplicado con un N Ref nuevo."""
    ws_m, ws_d = _libro(tmp_path, DOCS)
    carpeta = tmp_path / "Proyecto Ejemplo"
    carpeta.mkdir()
    fuente = carpeta / "b.pdf"
    fuente.write_bytes(b"pdf falso")

    resumen = acc.eliminar_documento("PROY-002", ws_m, ws_d, raiz_docs=tmp_path,
                                     ruta_backups=tmp_path / "resp")

    assert not fuente.exists()
    assert resumen["archivo_archivado"].exists()
    assert resumen["archivo_archivado"].read_bytes() == b"pdf falso"
    assert "PROY-002" in resumen["archivo_archivado"].name


def test_no_falla_si_el_archivo_fuente_ya_no_esta(tmp_path):
    ws_m, ws_d = _libro(tmp_path, DOCS)

    resumen = acc.eliminar_documento("PROY-002", ws_m, ws_d, raiz_docs=tmp_path,
                                     ruta_backups=tmp_path / "resp")

    assert resumen["archivo_archivado"] is None


def test_puede_dejar_el_archivo_donde_esta(tmp_path):
    ws_m, ws_d = _libro(tmp_path, DOCS)
    carpeta = tmp_path / "Proyecto Ejemplo"
    carpeta.mkdir()
    (carpeta / "b.pdf").write_bytes(b"pdf falso")

    acc.eliminar_documento("PROY-002", ws_m, ws_d, raiz_docs=tmp_path,
                           ruta_backups=tmp_path / "resp", archivar_fuente=False)

    assert (carpeta / "b.pdf").exists()


def test_falla_si_el_n_ref_no_existe(tmp_path):
    ws_m, ws_d = _libro(tmp_path, DOCS)

    with pytest.raises(ValueError, match="NOEXISTE-999"):
        acc.eliminar_documento("NOEXISTE-999", ws_m, ws_d, raiz_docs=tmp_path,
                               ruta_backups=tmp_path / "resp")


def test_no_renumera_los_n_ref_que_quedan(tmp_path):
    """El numero es historico: puede estar referenciado en correcciones,
    notas o respaldos. Queda un hueco en la secuencia, que es lo correcto."""
    ws_m, ws_d = _libro(tmp_path, DOCS)

    acc.eliminar_documento("PROY-001", ws_m, ws_d, raiz_docs=tmp_path,
                           ruta_backups=tmp_path / "resp")

    refs = [ws_m.cell(row=r, column=1).value for r in range(2, acc.ultima_fila_datos(ws_m) + 1)]
    assert refs == ["PROY-002", "UMAG-001"]
